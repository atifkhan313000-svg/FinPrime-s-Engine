"""
Finprime's Engine  -  v3.1  (design overhaul)
=============================================
A Windows desktop automation hub for saving, editing, testing, scheduling and
running Python scripts and VBA macros (built with customtkinter).

REQUIRED : customtkinter, pandas, openpyxl, Pillow
OPTIONAL : pywin32 (VBA runs)        send2trash (recycle-bin deletes)
           keyring (saved AI key)    pystray (system tray)
           plyer (desktop toasts)    tkinterdnd2 (drag & drop files)
           ruff / pyflakes / black   (lint + format in the editor)
Every optional package degrades gracefully - the app never crashes without it.

Script conventions (all optional)
---------------------------------
  # @param report_month: str = January -- Month to process
  # @param input_file: file            (types: str int float bool file folder)
Header params arrive as env vars  FP_PARAM_<NAME>.  argparse arguments are
detected automatically and passed on the command line.
Always available to scripts:  FP_WORKSPACE, FP_OUTPUT_DIR, FP_TRIGGER_FILE.
"""
from __future__ import annotations

import ast
import bisect
import difflib
import fnmatch
import hashlib
import hmac
import importlib
import importlib.metadata
import importlib.util
import json
import os
import queue
import random
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, ttk

import customtkinter as ctk

# Optional drag & drop support (must be imported before the root window exists)
try:  # pragma: no cover - optional dependency
    from tkinterdnd2 import TkinterDnD, DND_FILES
    HAS_DND = True
except Exception:  # noqa: BLE001
    TkinterDnD = None
    DND_FILES = None
    HAS_DND = False

APP_NAME = "Finprime's Engine"
APP_VERSION = "3.1"
IS_WIN = os.name == "nt"
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0) if IS_WIN else 0
SANDBOX_KEY = "__sandbox__"
KEYRING_SERVICE = "FinprimesEngine"


# =============================================================================
#  CONFIG & PATHS
# =============================================================================
CONFIG_DIR = Path(os.environ.get("APPDATA") or (Path.home() / ".config")) / "FinprimesEngine"
CONFIG_DIR.mkdir(parents=True, exist_ok=True)
CONFIG_FILE = CONFIG_DIR / "config.json"
DB_FILE = CONFIG_DIR / "engine.db"
AUTH_FILE = CONFIG_DIR / "auth.json"
DRAFT_FILE = CONFIG_DIR / "scratch_draft.json"
VERSIONS_DIR = CONFIG_DIR / "versions"
LEGACY_AUTH_NAME = ".finprime_engine_auth.json"
LEGACY_WORKSPACE = r"D:\FinPrime's pys"


def _default_workspace() -> str:
    if IS_WIN and os.path.isdir(LEGACY_WORKSPACE):
        return LEGACY_WORKSPACE
    return str(Path.home() / "FinPrime's pys")


class Config:
    DEFAULTS = {
        "workspace": "",
        "theme": "Matrix Blue",
        "matrix_rain": True,
        "themed_text": True,
        "sound": False,
        "auto_lock_minutes": 10,
        "run_timeout_s": 0,
        "notify_after_s": 10,
        "minimize_to_tray": False,
        "python_path": "",
        "glitch": True,
        "glitch_level": "Normal",
        "ai_model": "claude-sonnet-5-5",
    }

    def __init__(self):
        self.data = dict(self.DEFAULTS)
        try:
            if CONFIG_FILE.exists():
                self.data.update(json.loads(CONFIG_FILE.read_text(encoding="utf-8")))
        except Exception as exc:  # noqa: BLE001
            print(f"[CONFIG] could not read config: {exc}")
        if not self.data.get("workspace"):
            self.data["workspace"] = _default_workspace()

    def __getitem__(self, key):
        return self.data.get(key, self.DEFAULTS.get(key))

    def __setitem__(self, key, value):
        self.data[key] = value

    def save(self):
        try:
            CONFIG_FILE.write_text(json.dumps(self.data, indent=2), encoding="utf-8")
        except Exception as exc:  # noqa: BLE001
            print(f"[CONFIG] could not save config: {exc}")


CFG = Config()


def workspace_dir() -> Path:
    p = Path(CFG["workspace"])
    p.mkdir(parents=True, exist_ok=True)
    return p


def output_dir() -> Path:
    p = workspace_dir() / "outputs"
    p.mkdir(parents=True, exist_ok=True)
    return p


# =============================================================================
#  NEXT-LEVEL DESIGN SYSTEM  (deep blue-black surfaces, neon accents, themes)
# =============================================================================
COLOR_BG_MAIN = "#0A0E17"
COLOR_SIDEBAR = "#0D1220"
COLOR_CARD_SURFACE = "#111827"
COLOR_CARD_HOVER = "#182036"
COLOR_BORDER = "#1F2A44"
COLOR_ACCENT = "#4C9AFF"
COLOR_ACCENT_HOVER = "#7DB6FF"
COLOR_ACCENT2 = "#22D3EE"
COLOR_SUCCESS = "#16A34A"
COLOR_SUCCESS_HOVER = "#22C55E"
COLOR_DANGER = "#DC2626"
COLOR_DANGER_HOVER = "#F0525A"
COLOR_WARN = "#F59E0B"
COLOR_TEXT_MAIN = "#E8EEF9"
COLOR_TEXT_MUTED = "#7C8AA5"
COLOR_CONSOLE_BG = "#070B14"
COLOR_CONSOLE_TEXT = "#3FE08F"
COLOR_RAIN_HEAD = "#4C9AFF"
COLOR_EDITOR_BG = "#0B1020"

FONT_FAMILY = "Segoe UI"          # body text      (resolved for real in init_fonts)
FONT_DISPLAY = "Segoe UI"         # headings / big numbers
FONT_MONO = "Consolas"            # code, console, keycaps

THEMES = {
    "Matrix Blue":   {"accent": "#4C9AFF", "hover": "#7DB6FF", "accent2": "#22D3EE", "console": "#3FE08F"},
    "Neon Glitch":   {"accent": "#FF2E88", "hover": "#FF6BAA", "accent2": "#00E5FF", "console": "#00E5FF"},
    "Imperial Gold": {"accent": "#F5B942", "hover": "#FFD37A", "accent2": "#FF7A45", "console": "#F5B942"},
    "Necron Green":  {"accent": "#34D399", "hover": "#6EE7B7", "accent2": "#22D3EE", "console": "#34D399"},
    "Void Purple":   {"accent": "#A78BFA", "hover": "#C4B5FD", "accent2": "#F472B6", "console": "#C4B5FD"},
}


def apply_theme(name: str) -> None:
    global COLOR_ACCENT, COLOR_ACCENT_HOVER, COLOR_ACCENT2, COLOR_RAIN_HEAD, COLOR_CONSOLE_TEXT
    t = THEMES.get(name, THEMES["Matrix Blue"])
    COLOR_ACCENT, COLOR_ACCENT_HOVER, COLOR_ACCENT2 = t["accent"], t["hover"], t["accent2"]
    COLOR_RAIN_HEAD, COLOR_CONSOLE_TEXT = t["accent"], t["console"]


def init_fonts() -> None:
    """Pick the best installed fonts (must run after the Tk root exists)."""
    global FONT_FAMILY, FONT_DISPLAY, FONT_MONO
    from tkinter import font as tkfont
    have = {f.lower(): f for f in tkfont.families()}

    def pick(cands, fallback):
        for c in cands:
            if c.lower() in have:
                return have[c.lower()]
        return fallback

    FONT_FAMILY = pick(["Segoe UI Variable Text", "Segoe UI Variable", "Segoe UI", "Inter", "Helvetica Neue",
                        "DejaVu Sans"], "TkDefaultFont")
    FONT_DISPLAY = pick(["Bahnschrift", "Segoe UI Variable Display", "Segoe UI Semibold", "Segoe UI", "Poppins",
                         "DejaVu Sans"], FONT_FAMILY)
    FONT_MONO = pick(["JetBrains Mono", "Cascadia Code", "Cascadia Mono", "Fira Code", "Consolas",
                      "DejaVu Sans Mono", "Courier New"], "TkFixedFont")


apply_theme(CFG["theme"])
ctk.set_appearance_mode("Dark")

# status key -> (plain label, themed label).  A coloured dot replaces the old emoji.
STATUS_LABELS = {
    "idle":    ("Idle",        "Idle"),
    "running": ("Running",     "Purging data"),
    "success": ("Success",     "Litany complete"),
    "failed":  ("Failed",      "Rite failed"),
    "stopped": ("Stopped",     "Halted"),
    "timeout": ("Timed out",   "Timed out"),
}


def status_style(status: str):
    """(label, text_color, bg_color) for a status badge - evaluated lazily so themes apply."""
    idx = 1 if CFG["themed_text"] else 0
    label = "●  " + STATUS_LABELS.get(status, STATUS_LABELS["idle"])[idx]
    table = {
        "idle":    (COLOR_TEXT_MUTED, "#151C2E"),
        "running": (COLOR_ACCENT2, blend(COLOR_BG_MAIN, COLOR_ACCENT2, 0.18)),
        "success": ("#4ADE80", "#0F2E1D"),
        "failed":  ("#FB7185", "#3A1420"),
        "stopped": ("#FBBF24", "#33260B"),
        "timeout": ("#FBBF24", "#33260B"),
    }
    fg, bg = table.get(status, table["idle"])
    return label, fg, bg


# =============================================================================
#  SMALL HELPERS
# =============================================================================
def now_iso() -> str:
    return datetime.now().isoformat(sep=" ", timespec="seconds")


def parse_iso(value):
    try:
        return datetime.fromisoformat(value) if value else None
    except Exception:  # noqa: BLE001
        return None


def fmt_duration(sec) -> str:
    if sec is None:
        return "—"
    if sec < 1:
        return f"{sec * 1000:.0f} ms"
    if sec < 60:
        return f"{sec:.1f}s"
    m, s = divmod(int(sec), 60)
    return f"{m}m {s:02d}s"


def time_ago(iso_str) -> str:
    dt = parse_iso(iso_str)
    if not dt:
        return "never"
    secs = (datetime.now() - dt).total_seconds()
    if secs < 60:
        return "just now"
    if secs < 3600:
        return f"{int(secs // 60)}m ago"
    if secs < 86400:
        return f"{int(secs // 3600)}h ago"
    if secs < 86400 * 7:
        return f"{int(secs // 86400)}d ago"
    return dt.strftime("%d %b %Y")


def safe_name(name: str) -> str:
    """Strip anything that could escape the workspace or break a Windows filename."""
    name = re.sub(r"[^\w\- .]", "_", name.strip())
    name = name.strip(" .")
    return name[:80]


def open_path(path) -> None:
    try:
        if IS_WIN:
            os.startfile(str(path))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception as exc:  # noqa: BLE001
        print(f"[OPEN] {exc}")


def fuzzy_score(query: str, text: str):
    """Subsequence fuzzy match. Returns a score (higher = better) or None."""
    q, t = query.lower().strip(), text.lower()
    if not q:
        return 0.0
    score, ti, last = 0.0, 0, -2
    for ch in q:
        idx = t.find(ch, ti)
        if idx < 0:
            return None
        score += 10 if idx == last + 1 else 1
        if idx == 0 or t[idx - 1] in " _-:/":
            score += 5
        last, ti = idx, idx + 1
    return score - len(t) * 0.01


def try_import(name: str):
    try:
        return importlib.import_module(name)
    except Exception:  # noqa: BLE001
        return None


def classify_line(line: str, stream: str) -> str:
    low = line.lower()
    if re.search(r"traceback|error|exception|fatal|\bfailed\b", low):
        return "error"
    if "warn" in low:
        return "warning"
    if "✅" in line or re.search(r"\b(success|completed|done)\b", low):
        return "success"
    return "stderr" if stream == "stderr" else "out"


# =============================================================================
#  SECURITY UTILITIES
# =============================================================================
PBKDF2_ITERATIONS = 600_000          # OWASP guidance for PBKDF2-HMAC-SHA256
LEGACY_ITERATIONS = 100_000          # what v2.0 used


def hash_password(password: str, salt: bytes | None = None, iterations: int = PBKDF2_ITERATIONS) -> tuple:
    if salt is None:
        salt = os.urandom(16)
    hashed = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, iterations)
    return hashed.hex(), salt.hex()


def verify_password(stored_hash: str, stored_salt: str, attempt: str, iterations: int = LEGACY_ITERATIONS) -> bool:
    salt = bytes.fromhex(stored_salt)
    attempt_hash, _ = hash_password(attempt, salt, iterations)
    return hmac.compare_digest(attempt_hash, stored_hash)      # constant-time compare


def migrate_legacy_auth() -> None:
    """Reuse the password created by v2.0 (it lived inside the workspace folder)."""
    if AUTH_FILE.exists():
        return
    for folder in {str(CFG["workspace"]), LEGACY_WORKSPACE}:
        legacy = Path(folder) / LEGACY_AUTH_NAME
        try:
            if legacy.exists():
                shutil.copy2(legacy, AUTH_FILE)
                return
        except Exception:  # noqa: BLE001
            pass


def load_auth():
    try:
        return json.loads(AUTH_FILE.read_text(encoding="utf-8")) if AUTH_FILE.exists() else None
    except Exception:  # noqa: BLE001
        return None


def save_auth(password: str) -> None:
    h, s = hash_password(password)
    AUTH_FILE.write_text(json.dumps({"hash": h, "salt": s, "iterations": PBKDF2_ITERATIONS}), encoding="utf-8")


def check_password(password: str) -> bool:
    """Verify against the stored credential, transparently upgrading old hashes."""
    creds = load_auth()
    if not creds:
        return False
    iters = int(creds.get("iterations", LEGACY_ITERATIONS))
    ok = verify_password(creds["hash"], creds["salt"], password, iters)
    if not ok and "iterations" not in creds and password != password.strip():
        # v2.0 stripped whitespace before hashing
        ok = verify_password(creds["hash"], creds["salt"], password.strip(), iters)
        password = password.strip()
    if ok and iters < PBKDF2_ITERATIONS:
        try:
            save_auth(password)
        except Exception:  # noqa: BLE001
            pass
    return ok


# =============================================================================
#  DATABASE  (run history, favourites/tags, schedules, pipelines, watchers)
# =============================================================================
SCHEMA = """
CREATE TABLE IF NOT EXISTS runs(
    id INTEGER PRIMARY KEY AUTOINCREMENT, script TEXT NOT NULL, kind TEXT,
    started TEXT NOT NULL, duration REAL, exit_code INTEGER, status TEXT,
    stdout TEXT, stderr TEXT, args TEXT, source TEXT, outputs TEXT);
CREATE INDEX IF NOT EXISTS idx_runs_script  ON runs(script);
CREATE INDEX IF NOT EXISTS idx_runs_started ON runs(started);
CREATE TABLE IF NOT EXISTS script_meta(
    name TEXT PRIMARY KEY, favorite INTEGER DEFAULT 0, tags TEXT DEFAULT '',
    last_params TEXT DEFAULT '{}');
CREATE TABLE IF NOT EXISTS schedules(
    id INTEGER PRIMARY KEY AUTOINCREMENT, target TEXT NOT NULL, kind TEXT NOT NULL,
    at TEXT DEFAULT '09:00', weekdays TEXT DEFAULT '', interval_min INTEGER DEFAULT 60,
    enabled INTEGER DEFAULT 1, last_fire TEXT, created TEXT);
CREATE TABLE IF NOT EXISTS pipelines(
    name TEXT PRIMARY KEY, steps TEXT NOT NULL, stop_on_fail INTEGER DEFAULT 1);
CREATE TABLE IF NOT EXISTS watchers(
    id INTEGER PRIMARY KEY AUTOINCREMENT, folder TEXT NOT NULL, pattern TEXT DEFAULT '*.xlsx',
    target TEXT NOT NULL, enabled INTEGER DEFAULT 1, fired INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS kv(key TEXT PRIMARY KEY, value TEXT);
"""

MAX_STORED_OUTPUT = 200_000


class Database:
    """Thin thread-safe SQLite wrapper (one short-lived connection per call)."""

    def __init__(self, path: Path):
        self.path = str(path)
        self._lock = threading.Lock()
        with self._lock:
            con = sqlite3.connect(self.path, timeout=15)
            con.executescript(SCHEMA)
            # retention: keep the newest 5000 runs
            con.execute("DELETE FROM runs WHERE id NOT IN (SELECT id FROM runs ORDER BY id DESC LIMIT 5000)")
            con.commit()
            con.close()

    def q(self, sql: str, params=()) -> list:
        with self._lock:
            con = sqlite3.connect(self.path, timeout=15)
            con.row_factory = sqlite3.Row
            try:
                return [dict(r) for r in con.execute(sql, params).fetchall()]
            finally:
                con.close()

    def x(self, sql: str, params=()) -> int:
        with self._lock:
            con = sqlite3.connect(self.path, timeout=15)
            try:
                cur = con.execute(sql, params)
                con.commit()
                return cur.lastrowid or 0
            finally:
                con.close()

    def q1(self, sql: str, params=()):
        rows = self.q(sql, params)
        return rows[0] if rows else None

    # ---- runs ---------------------------------------------------------------
    def add_run(self, res: "RunResult") -> int:
        return self.x(
            "INSERT INTO runs(script,kind,started,duration,exit_code,status,stdout,stderr,args,source,outputs)"
            " VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            (res.script, res.kind, res.started, res.duration, res.exit_code, res.status,
             res.stdout[-MAX_STORED_OUTPUT:], res.stderr[-MAX_STORED_OUTPUT:],
             json.dumps(res.args), res.source, json.dumps(res.outputs)))

    def get_run(self, run_id: int):
        return self.q1("SELECT * FROM runs WHERE id=?", (run_id,))

    def list_runs(self, script: str | None = None, status: str | None = None, limit: int = 500) -> list:
        sql, params = "SELECT id,script,kind,started,duration,exit_code,status,source,outputs FROM runs WHERE 1=1", []
        if script and script != "All scripts":
            sql += " AND script=?"
            params.append(script)
        if status and status != "Any status":
            sql += " AND status=?"
            params.append(status)
        sql += " ORDER BY id DESC LIMIT ?"
        params.append(limit)
        return self.q(sql, params)

    def clear_runs(self, script: str | None = None) -> None:
        if script:
            self.x("DELETE FROM runs WHERE script=?", (script,))
        else:
            self.x("DELETE FROM runs")

    def last_runs_by_script(self) -> dict:
        rows = self.q(
            "SELECT r.script, r.started, r.duration, r.status, r.outputs, c.n FROM runs r "
            "JOIN (SELECT script, MAX(id) mid, COUNT(*) n FROM runs GROUP BY script) c ON r.id=c.mid")
        return {r["script"]: r for r in rows}

    # ---- stats for the dashboard -------------------------------------------
    def overview(self) -> dict:
        today = datetime.now().strftime("%Y-%m-%d")
        week = (datetime.now() - timedelta(days=6)).strftime("%Y-%m-%d")
        runs_today = self.q1("SELECT COUNT(*) c FROM runs WHERE substr(started,1,10)=?", (today,))["c"]
        w = self.q1("SELECT COUNT(*) n, SUM(status='success') ok, AVG(duration) avg_d "
                    "FROM runs WHERE substr(started,1,10)>=?", (week,))
        fail = self.q1("SELECT script, started FROM runs WHERE status IN ('failed','timeout') ORDER BY id DESC LIMIT 1")
        n = w["n"] or 0
        return {"runs_today": runs_today, "week_runs": n,
                "success_rate": (100.0 * (w["ok"] or 0) / n) if n else None,
                "avg_duration": w["avg_d"], "last_failure": fail}

    def daily_counts(self, days: int = 7) -> list:
        out = []
        for i in range(days - 1, -1, -1):
            d = (datetime.now() - timedelta(days=i)).strftime("%Y-%m-%d")
            r = self.q1("SELECT SUM(status='success') ok, SUM(status!='success') bad "
                        "FROM runs WHERE substr(started,1,10)=?", (d,))
            out.append((d, r["ok"] or 0, r["bad"] or 0))
        return out

    def recent_runs(self, limit: int = 8) -> list:
        return self.q("SELECT id,script,started,duration,status FROM runs ORDER BY id DESC LIMIT ?", (limit,))

    def achievement_stats(self) -> dict:
        tot = self.q1("SELECT COUNT(*) n, SUM(status='success') ok FROM runs")
        days = [r["d"] for r in self.q("SELECT DISTINCT substr(started,1,10) d FROM runs "
                                       "WHERE status='success' ORDER BY d DESC")]
        streak, cursor = 0, datetime.now().date()
        if days and days[0] != cursor.strftime("%Y-%m-%d"):
            cursor -= timedelta(days=1)          # streak stays alive until end of the next day
        for d in days:
            if d == cursor.strftime("%Y-%m-%d"):
                streak += 1
                cursor -= timedelta(days=1)
            else:
                break
        last10 = self.q("SELECT status FROM runs ORDER BY id DESC LIMIT 10")
        night = self.q1("SELECT COUNT(*) c FROM runs WHERE CAST(substr(started,12,2) AS INTEGER) < 4")["c"]
        kinds = {r["kind"] for r in self.q("SELECT DISTINCT kind FROM runs WHERE status='success'")}
        return {"total": tot["n"] or 0, "ok": tot["ok"] or 0, "streak": streak,
                "flawless": len(last10) == 10 and all(r["status"] == "success" for r in last10),
                "night": night > 0, "polyglot": {"python", "vba"} <= kinds,
                "pipelines": int(self.kv_get("pipeline_runs", "0"))}

    # ---- script meta --------------------------------------------------------
    def meta_all(self) -> dict:
        out = {}
        for r in self.q("SELECT * FROM script_meta"):
            r["tags"] = [t for t in (r["tags"] or "").split(",") if t]
            out[r["name"]] = r
        return out

    def _ensure_meta(self, name: str) -> None:
        self.x("INSERT OR IGNORE INTO script_meta(name) VALUES(?)", (name,))

    def set_favorite(self, name: str, fav: bool) -> None:
        self._ensure_meta(name)
        self.x("UPDATE script_meta SET favorite=? WHERE name=?", (1 if fav else 0, name))

    def set_tags(self, name: str, tags: list) -> None:
        self._ensure_meta(name)
        self.x("UPDATE script_meta SET tags=? WHERE name=?", (",".join(tags), name))

    def get_last_params(self, name: str) -> dict:
        r = self.q1("SELECT last_params FROM script_meta WHERE name=?", (name,))
        try:
            return json.loads(r["last_params"]) if r and r["last_params"] else {}
        except Exception:  # noqa: BLE001
            return {}

    def set_last_params(self, name: str, values: dict) -> None:
        self._ensure_meta(name)
        self.x("UPDATE script_meta SET last_params=? WHERE name=?", (json.dumps(values), name))

    def rename_meta(self, old: str, new: str) -> None:
        self.x("UPDATE script_meta SET name=? WHERE name=?", (new, old))

    # ---- schedules / pipelines / watchers -----------------------------------
    def schedules(self) -> list:
        return self.q("SELECT * FROM schedules ORDER BY id")

    def add_schedule(self, d: dict) -> int:
        return self.x("INSERT INTO schedules(target,kind,at,weekdays,interval_min,enabled,created) VALUES(?,?,?,?,?,1,?)",
                      (d["target"], d["kind"], d.get("at", "09:00"), d.get("weekdays", ""),
                       int(d.get("interval_min", 60)), now_iso()))

    def update_schedule(self, sid: int, d: dict) -> None:
        self.x("UPDATE schedules SET target=?,kind=?,at=?,weekdays=?,interval_min=? WHERE id=?",
               (d["target"], d["kind"], d.get("at", "09:00"), d.get("weekdays", ""),
                int(d.get("interval_min", 60)), sid))

    def toggle_schedule(self, sid: int, enabled: bool) -> None:
        self.x("UPDATE schedules SET enabled=? WHERE id=?", (1 if enabled else 0, sid))

    def mark_fired(self, sid: int) -> None:
        self.x("UPDATE schedules SET last_fire=? WHERE id=?", (now_iso(), sid))

    def delete_schedule(self, sid: int) -> None:
        self.x("DELETE FROM schedules WHERE id=?", (sid,))

    def pipelines(self) -> list:
        rows = self.q("SELECT * FROM pipelines ORDER BY name")
        for r in rows:
            try:
                r["steps"] = json.loads(r["steps"])
            except Exception:  # noqa: BLE001
                r["steps"] = []
        return rows

    def get_pipeline(self, name: str):
        for p in self.pipelines():
            if p["name"] == name:
                return p
        return None

    def save_pipeline(self, name: str, steps: list, stop_on_fail: bool, old_name: str | None = None) -> None:
        if old_name and old_name != name:
            self.x("DELETE FROM pipelines WHERE name=?", (old_name,))
        self.x("INSERT OR REPLACE INTO pipelines(name,steps,stop_on_fail) VALUES(?,?,?)",
               (name, json.dumps(steps), 1 if stop_on_fail else 0))

    def delete_pipeline(self, name: str) -> None:
        self.x("DELETE FROM pipelines WHERE name=?", (name,))

    def watchers(self) -> list:
        return self.q("SELECT * FROM watchers ORDER BY id")

    def add_watcher(self, d: dict) -> None:
        self.x("INSERT INTO watchers(folder,pattern,target,enabled) VALUES(?,?,?,1)",
               (d["folder"], d.get("pattern", "*.xlsx"), d["target"]))

    def toggle_watcher(self, wid: int, enabled: bool) -> None:
        self.x("UPDATE watchers SET enabled=? WHERE id=?", (1 if enabled else 0, wid))

    def bump_watcher(self, wid: int) -> None:
        self.x("UPDATE watchers SET fired=fired+1 WHERE id=?", (wid,))

    def delete_watcher(self, wid: int) -> None:
        self.x("DELETE FROM watchers WHERE id=?", (wid,))

    # ---- key/value ----------------------------------------------------------
    def kv_get(self, key: str, default: str = "") -> str:
        r = self.q1("SELECT value FROM kv WHERE key=?", (key,))
        return r["value"] if r else default

    def kv_set(self, key: str, value: str) -> None:
        self.x("INSERT OR REPLACE INTO kv(key,value) VALUES(?,?)", (key, value))


# =============================================================================
#  SCRIPT ANALYSIS  (parameters, imports, lint, format, VBA, outputs, pip)
# =============================================================================
def get_python() -> str:
    """Interpreter used to run user scripts (safe when packaged with PyInstaller)."""
    custom = str(CFG["python_path"]).strip()
    if custom and Path(custom).exists():
        return custom
    if getattr(sys, "frozen", False):
        return shutil.which("python") or shutil.which("py") or "python"
    return sys.executable


def own_interpreter() -> bool:
    return get_python() == sys.executable


@dataclass
class Param:
    name: str
    type: str = "str"
    default: str = ""
    help: str = ""
    source: str = "header"          # "header" (env var) or "argparse" (command line)
    flag: str | None = None
    required: bool = False
    choices: list | None = None
    action: str = ""


HEADER_PARAM_RE = re.compile(
    r"^\s*#\s*@param\s+(?P<name>\w+)(?:\s*:\s*(?P<type>\w+))?"
    r"(?:\s*=\s*(?P<default>.*?))?(?:\s+--\s*(?P<help>.*))?\s*$")


def _literal(node):
    if node is None:
        return None
    try:
        return ast.literal_eval(node)
    except Exception:  # noqa: BLE001
        return None


def parse_params(code: str) -> list:
    """Detect runtime parameters: argparse.add_argument(...) calls and '# @param' header lines."""
    params: list[Param] = []
    try:
        tree = ast.parse(code)
    except SyntaxError:
        tree = None
    if tree is not None:
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
                    and node.func.attr == "add_argument"):
                continue
            strs = [a.value for a in node.args if isinstance(a, ast.Constant) and isinstance(a.value, str)]
            if not strs or "-h" in strs or "--help" in strs:
                continue
            kw = {k.arg: k.value for k in node.keywords if k.arg}
            longs = [s for s in strs if s.startswith("--")]
            shorts = [s for s in strs if s.startswith("-") and not s.startswith("--")]
            if longs:
                flag, name = longs[0], longs[0].lstrip("-").replace("-", "_")
            elif shorts:
                flag, name = shorts[0], shorts[0].lstrip("-")
            else:
                flag, name = None, strs[0]
            action = _literal(kw.get("action")) or ""
            ptype = "str"
            t = kw.get("type")
            if isinstance(t, ast.Name) and t.id in ("int", "float"):
                ptype = t.id
            if action in ("store_true", "store_false"):
                ptype = "bool"
            lname = name.lower()
            if ptype == "str":
                if any(k in lname for k in ("dir", "folder")):
                    ptype = "folder"
                elif any(k in lname for k in ("file", "input", "workbook", "excel", "path")):
                    ptype = "file"
            default = _literal(kw.get("default"))
            choices = _literal(kw.get("choices"))
            params.append(Param(
                name=name, type=ptype, default="" if default is None else str(default),
                help=_literal(kw.get("help")) or "", source="argparse", flag=flag,
                required=bool(_literal(kw.get("required"))) or flag is None,
                choices=list(choices) if isinstance(choices, (list, tuple)) else None, action=action))
    seen = {p.name for p in params}
    for line in code.splitlines()[:80]:
        m = HEADER_PARAM_RE.match(line)
        if m and m.group("name") not in seen:
            ptype = (m.group("type") or "str").lower()
            if ptype not in ("str", "int", "float", "bool", "file", "folder"):
                ptype = "str"
            params.append(Param(name=m.group("name"), type=ptype,
                                default=(m.group("default") or "").strip().strip("\"'"),
                                help=(m.group("help") or "").strip(), source="header"))
            seen.add(m.group("name"))
    return params


def build_cli(params: list, values: dict) -> list:
    positional, args = [], []
    for p in params:
        if p.source != "argparse":
            continue
        v = values.get(p.name, p.default)
        if p.type == "bool":
            truthy = str(v).lower() in ("1", "true", "yes", "on")
            if p.flag and ((p.action == "store_true" and truthy) or (p.action == "store_false" and not truthy)):
                args.append(p.flag)
            continue
        if v in (None, ""):
            continue
        if p.flag:
            args += [p.flag, str(v)]
        else:
            positional.append(str(v))
    return positional + args


def build_param_env(params: list, values: dict) -> dict:
    env = {}
    for p in params:
        if p.source == "header":
            env[f"FP_PARAM_{p.name.upper()}"] = str(values.get(p.name, p.default))
    return env


# ---- imports & dependencies -------------------------------------------------
_FALLBACK_STDLIB = {
    "os", "sys", "re", "json", "math", "time", "datetime", "pathlib", "csv", "shutil", "subprocess",
    "threading", "collections", "itertools", "functools", "typing", "logging", "argparse", "random",
    "sqlite3", "tempfile", "glob", "io", "copy", "string", "statistics", "decimal", "fractions",
    "zipfile", "hashlib", "urllib", "http", "email", "smtplib", "tkinter", "unittest", "dataclasses",
    "enum", "abc", "contextlib", "traceback", "warnings", "queue", "socket", "ssl", "struct", "uuid",
    "calendar", "textwrap", "pprint", "operator", "heapq", "bisect", "base64", "getpass", "platform",
    "ast", "inspect", "importlib", "configparser", "xml", "html", "asyncio", "concurrent", "multiprocessing",
}
STDLIB = set(getattr(sys, "stdlib_module_names", _FALLBACK_STDLIB)) | {"__future__"}

IMPORT_TO_PIP = {
    "cv2": "opencv-python", "PIL": "Pillow", "yaml": "PyYAML", "sklearn": "scikit-learn",
    "win32com": "pywin32", "win32api": "pywin32", "win32con": "pywin32", "win32gui": "pywin32",
    "pythoncom": "pywin32", "pywintypes": "pywin32", "bs4": "beautifulsoup4", "dateutil": "python-dateutil",
    "docx": "python-docx", "pptx": "python-pptx", "fitz": "PyMuPDF", "serial": "pyserial",
    "dotenv": "python-dotenv", "xlsxwriter": "XlsxWriter", "attr": "attrs", "Crypto": "pycryptodome",
    "jwt": "PyJWT", "skimage": "scikit-image", "git": "GitPython", "PyPDF2": "PyPDF2", "pypdf": "pypdf",
    "customtkinter": "customtkinter", "tkinterdnd2": "tkinterdnd2", "xlwings": "xlwings",
}


def scan_imports(code: str, local_dir: Path | None = None) -> set:
    try:
        tree = ast.parse(code)
    except SyntaxError:
        return set()
    mods = set()
    for n in ast.walk(tree):
        if isinstance(n, ast.Import):
            mods.update(a.name.split(".")[0] for a in n.names)
        elif isinstance(n, ast.ImportFrom) and n.level == 0 and n.module:
            mods.add(n.module.split(".")[0])
    mods -= STDLIB
    if local_dir:
        mods = {m for m in mods if not ((local_dir / f"{m}.py").exists() or (local_dir / m).is_dir())}
    return mods


def pip_name(module: str) -> str:
    return IMPORT_TO_PIP.get(module, module)


def check_modules(mods) -> dict:
    """{module: installed?}  Checked in the interpreter that will actually run the scripts."""
    mods = sorted(set(mods))
    if not mods:
        return {}
    if own_interpreter():
        out = {}
        for m in mods:
            try:
                out[m] = importlib.util.find_spec(m) is not None
            except Exception:  # noqa: BLE001
                out[m] = False
        return out
    code = ("import importlib.util,json,sys;"
            "print(json.dumps({m: importlib.util.find_spec(m) is not None for m in sys.argv[1:]}))")
    try:
        p = subprocess.run([get_python(), "-c", code, *mods], capture_output=True, text=True,
                           timeout=30, creationflags=NO_WINDOW)
        return json.loads(p.stdout.strip().splitlines()[-1])
    except Exception:  # noqa: BLE001
        return {m: False for m in mods}


def package_version(dist: str):
    try:
        return importlib.metadata.version(dist)
    except Exception:  # noqa: BLE001
        return None


def pip_install(packages: list, emit_line) -> int:
    """Blocking pip install that streams output through emit_line(str). Run in a thread."""
    cmd = [get_python(), "-m", "pip", "install", "--disable-pip-version-check", *packages]
    emit_line("$ " + " ".join(cmd) + "\n")
    try:
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
                                encoding="utf-8", errors="replace", creationflags=NO_WINDOW)
        for line in proc.stdout:
            emit_line(line)
        return proc.wait()
    except Exception as exc:  # noqa: BLE001
        emit_line(f"pip failed to start: {exc}\n")
        return 1


# ---- lint / format ----------------------------------------------------------
def _run_tool(args: list, stdin: str, timeout: int = 40):
    return subprocess.run([get_python(), *args], input=stdin, capture_output=True, text=True,
                          encoding="utf-8", errors="replace", timeout=timeout, creationflags=NO_WINDOW)


def builtin_lint(code: str) -> list:
    issues = []
    try:
        tree = ast.parse(code)
    except SyntaxError as e:
        return [f"{e.lineno}:{e.offset or 0}: SyntaxError: {e.msg}"]
    imported = {}
    for n in ast.walk(tree):
        if isinstance(n, ast.Import):
            for a in n.names:
                imported[(a.asname or a.name).split(".")[0]] = n.lineno
        elif isinstance(n, ast.ImportFrom):
            for a in n.names:
                if a.name != "*":
                    imported[a.asname or a.name] = n.lineno
        elif isinstance(n, ast.ExceptHandler) and n.type is None:
            issues.append(f"{n.lineno}:1: bare 'except:' hides errors - catch Exception instead")
        elif isinstance(n, ast.Compare):
            for op, comp in zip(n.ops, n.comparators):
                if isinstance(op, (ast.Eq, ast.NotEq)) and isinstance(comp, ast.Constant) and comp.value is None:
                    issues.append(f"{n.lineno}:1: comparison to None should use 'is' / 'is not'")
        elif isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)):
            for d in n.args.defaults:
                if isinstance(d, (ast.List, ast.Dict, ast.Set)):
                    issues.append(f"{n.lineno}:1: mutable default argument in '{n.name}'")
    used = {n.id for n in ast.walk(tree) if isinstance(n, ast.Name)}
    used |= {n.value.id for n in ast.walk(tree) if isinstance(n, ast.Attribute) and isinstance(n.value, ast.Name)}
    for name, line in sorted(imported.items(), key=lambda kv: kv[1]):
        if name not in used:
            issues.append(f"{line}:1: '{name}' imported but unused")
    return issues


def run_lint(code: str):
    """Returns (tool_name, [issue lines]). Prefers ruff, then pyflakes, then a built-in checker."""
    attempts = [
        ("ruff", ["-m", "ruff", "check", "--stdin-filename", "script.py", "--output-format", "concise", "-"]),
        ("pyflakes", ["-m", "pyflakes"]),
    ]
    for tool, args in attempts:
        try:
            p = _run_tool(args, code)
        except Exception:  # noqa: BLE001
            continue
        if "No module named" in (p.stderr or ""):
            continue
        lines = (p.stdout + "\n" + p.stderr).strip().splitlines()
        issues = [ln for ln in lines if ln.strip()
                  and not ln.startswith(("All checks passed", "Found ", "warning:"))]
        return tool, issues
    return "built-in checker", builtin_lint(code)


def format_code(code: str):
    """Returns (tool_name, formatted_code). Raises RuntimeError when no formatter is installed."""
    for tool, args in (("black", ["-m", "black", "-q", "-"]),
                       ("ruff format", ["-m", "ruff", "format", "--stdin-filename", "script.py", "-"])):
        try:
            p = _run_tool(args, code)
        except Exception:  # noqa: BLE001
            continue
        if "No module named" in (p.stderr or ""):
            continue
        if p.returncode != 0:
            raise RuntimeError((p.stderr or "formatter failed").strip()[:400])
        return tool, p.stdout
    raise RuntimeError("No formatter found. Install one from Env Health Check (black or ruff).")


# ---- VBA helpers ------------------------------------------------------------
VBA_SUB_RE = re.compile(r"^\s*(?P<vis>Public\s+|Private\s+|Friend\s+)?Sub\s+(?P<name>[A-Za-z_]\w*)\s*\((?P<args>[^)]*)\)",
                        re.I | re.M)


def vba_subs(code: str) -> list:
    """Public, parameter-less Subs - the ones Excel can run directly."""
    out = []
    for m in VBA_SUB_RE.finditer(code):
        if (m.group("vis") or "").strip().lower() == "private":
            continue
        if m.group("args").strip():
            continue
        out.append(m.group("name"))
    return out


_VBA_OPEN = [("Sub", r"^\s*(?:(?:Public|Private|Friend|Static)\s+)*Sub\b", r"^\s*End\s+Sub\b"),
             ("Function", r"^\s*(?:(?:Public|Private|Friend|Static)\s+)*Function\b", r"^\s*End\s+Function\b"),
             ("If", r"^\s*If\b.*\bThen\s*$", r"^\s*End\s+If\b"),
             ("For", r"^\s*For\b", r"^\s*Next\b"),
             ("Do", r"^\s*Do\b", r"^\s*Loop\b"),
             ("While", r"^\s*While\b", r"^\s*Wend\b"),
             ("Select", r"^\s*Select\s+Case\b", r"^\s*End\s+Select\b"),
             ("With", r"^\s*With\b", r"^\s*End\s+With\b")]


def check_vba(code: str) -> list:
    """Lightweight structural check (block pairing). Real compilation only happens inside Excel."""
    issues, stack = [], []
    for no, raw in enumerate(code.splitlines(), 1):
        line = re.sub(r'"(?:""|[^"\n])*"', '""', raw)
        line = re.sub(r"'.*$", "", line)
        line = re.sub(r"^\s*Rem\b.*$", "", line, flags=re.I)
        if not line.strip():
            continue
        closed = False
        for kind, _op, cl in _VBA_OPEN:
            if re.match(cl, line, re.I):
                closed = True
                if stack and stack[-1][0] == kind:
                    stack.pop()
                else:
                    issues.append(f"{no}: '{line.strip()}' closes a block that is not open")
                break
        if closed:
            continue
        for kind, op, _cl in _VBA_OPEN:
            if re.match(op, line, re.I):
                stack.append((kind, no))
                break
    for kind, no in stack:
        issues.append(f"{no}: '{kind}' block is never closed")
    if not vba_subs(code):
        issues.append("0: no public parameter-less Sub found - nothing for Excel to run")
    if not re.search(r"^\s*Option\s+Explicit", code, re.I | re.M):
        issues.append("1: tip - add 'Option Explicit' at the top to catch typos in variable names")
    return issues


# ---- output-file detection (feeds the preview panel) ------------------------
OUTPUT_EXTS = {".xlsx", ".xlsm", ".xls", ".csv"}
_PATH_RE = re.compile(r"""[A-Za-z]:\\[^\r\n"'<>|*?]+?\.(?:xlsx|xlsm|xls|csv)|/[^\s"'<>|*?]+\.(?:xlsx|xlsm|xls|csv)""", re.I)


def _scan_dir(folder: Path, depth: int, acc: dict) -> None:
    try:
        with os.scandir(folder) as it:
            for e in it:
                if e.name.startswith(("~$", ".")):
                    continue
                if e.is_file() and Path(e.name).suffix.lower() in OUTPUT_EXTS:
                    acc[e.path] = e.stat().st_mtime
                elif e.is_dir() and depth > 0 and not e.name.startswith("_"):
                    _scan_dir(Path(e.path), depth - 1, acc)
    except OSError:
        pass


def snapshot_outputs() -> dict:
    acc: dict = {}
    _scan_dir(workspace_dir(), 0, acc)
    _scan_dir(output_dir(), 2, acc)
    return acc


def detect_outputs(before: dict, text: str, started_ts: float) -> list:
    after = snapshot_outputs()
    found = [p for p, m in after.items() if p not in before or m > before[p]]
    for m in _PATH_RE.finditer(text or ""):
        p = m.group(0).strip()
        try:
            if os.path.exists(p) and os.path.getmtime(p) >= started_ts - 2 and p not in found:
                found.append(p)
        except OSError:
            pass
    return found[:12]


# =============================================================================
#  EXECUTION ENGINE  (no UI code in here - everything talks through emit())
# =============================================================================
@dataclass
class RunResult:
    key: str
    script: str
    kind: str
    status: str
    exit_code: int | None
    duration: float
    stdout: str
    stderr: str
    started: str
    outputs: list = field(default_factory=list)
    args: list = field(default_factory=list)
    source: str = "manual"
    message: str = ""


def _kill_tree(proc: subprocess.Popen) -> None:
    try:
        if IS_WIN:
            subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"], capture_output=True,
                           creationflags=NO_WINDOW)
        else:
            proc.kill()
    except Exception:  # noqa: BLE001
        try:
            proc.kill()
        except Exception:  # noqa: BLE001
            pass


class ScriptRunner:
    """Runs scripts/macros on worker threads and streams output as events.

    Events pushed through emit():
        ("status",   key, status)
        ("log",      key, text, tag)
        ("finished", key, RunResult)
    """

    def __init__(self, emit, db: Database):
        self.emit = emit
        self.db = db
        self._lock = threading.Lock()
        self.active: set = set()
        self.procs: dict = {}
        self.flags: dict = {}

    def is_running(self, key: str) -> bool:
        return key in self.active

    def running_keys(self) -> list:
        return list(self.active)

    # ---- control ------------------------------------------------------------
    def stop(self, key: str) -> bool:
        proc = self.procs.get(key)
        if proc and proc.poll() is None:
            self.flags[key] = "stopped"
            _kill_tree(proc)
            return True
        return False

    def stop_all(self) -> None:
        for k in list(self.procs):
            self.stop(k)

    def _timeout_kill(self, key: str) -> None:
        proc = self.procs.get(key)
        if proc and proc.poll() is None:
            self.flags[key] = "timeout"
            _kill_tree(proc)

    def _log(self, key: str, text: str, tag: str = "info") -> None:
        self.emit(("log", key, text, tag))

    # ---- main entry ---------------------------------------------------------
    def execute(self, key: str, script: str, path: Path, kind: str, args=None, env_extra=None,
                timeout: int = 0, source: str = "manual") -> RunResult | None:
        args = list(args or [])
        with self._lock:
            if key in self.active:
                return None
            self.active.add(key)
            self.flags.pop(key, None)
        started_dt = datetime.now()
        t0 = time.monotonic()
        before = snapshot_outputs()
        self.emit(("status", key, "running"))
        stamp = started_dt.strftime("%H:%M:%S")
        self._log(key, f"\n[{stamp}] ▶ EXECUTING {kind.upper()}: {script}"
                       + (f"  args: {' '.join(args)}" if args else "") + "\n" + "─" * 60 + "\n")
        try:
            if kind == "vba":
                status, code, out, err = self._run_vba(key, path)
            else:
                status, code, out, err = self._run_python(key, path, args, env_extra or {}, timeout)
        except Exception as exc:  # noqa: BLE001
            status, code, out, err = "failed", -1, "", f"{type(exc).__name__}: {exc}\n"
            self._log(key, err, "error")
        finally:
            with self._lock:
                self.active.discard(key)
                self.procs.pop(key, None)
        duration = time.monotonic() - t0
        outputs = detect_outputs(before, out + "\n" + err, started_dt.timestamp())
        res = RunResult(key, script, kind, status, code, duration, out, err,
                        started_dt.isoformat(sep=" ", timespec="seconds"), outputs, args, source)
        if key != SANDBOX_KEY:
            try:
                self.db.add_run(res)
            except Exception as exc:  # noqa: BLE001
                self._log(key, f"[history] could not save run: {exc}\n", "warning")
        verdict = {"success": "✅ finished OK", "failed": "❌ failed", "stopped": "⏹ stopped by user",
                   "timeout": "⌛ killed - timeout reached"}.get(status, status)
        self._log(key, "─" * 60 + f"\n[{datetime.now():%H:%M:%S}] {verdict}  "
                       f"(exit {code}, {fmt_duration(duration)})\n",
                  "success" if status == "success" else "error" if status in ("failed", "timeout") else "warning")
        if outputs:
            self._log(key, "📊 Output files: " + "; ".join(Path(o).name for o in outputs) + "\n", "info")
        self.emit(("status", key, status))
        self.emit(("finished", key, res))
        return res

    # ---- python -------------------------------------------------------------
    def _run_python(self, key, path, args, env_extra, timeout):
        env = os.environ.copy()
        env.update({"PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1", "PYTHONUNBUFFERED": "1",
                    "FP_WORKSPACE": str(workspace_dir()), "FP_OUTPUT_DIR": str(output_dir())})
        env.update(env_extra)
        cmd = [get_python(), "-u", str(path), *args]
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, stdin=subprocess.DEVNULL,
                                text=True, encoding="utf-8", errors="replace", cwd=str(workspace_dir()),
                                env=env, creationflags=NO_WINDOW)
        self.procs[key] = proc
        out_lines: list = []
        err_lines: list = []

        def pump(stream, name, sink):
            try:
                for line in iter(stream.readline, ""):
                    sink.append(line)
                    self._log(key, line, classify_line(line, name))
            finally:
                stream.close()

        threads = [threading.Thread(target=pump, args=(proc.stdout, "stdout", out_lines), daemon=True),
                   threading.Thread(target=pump, args=(proc.stderr, "stderr", err_lines), daemon=True)]
        for t in threads:
            t.start()
        timer = None
        if timeout and timeout > 0:
            timer = threading.Timer(timeout, self._timeout_kill, (key,))
            timer.daemon = True
            timer.start()
        rc = proc.wait()
        for t in threads:
            t.join(timeout=5)
        if timer:
            timer.cancel()
        flag = self.flags.pop(key, None)
        status = flag or ("success" if rc == 0 else "failed")
        return status, rc, "".join(out_lines), "".join(err_lines)

    # ---- vba ----------------------------------------------------------------
    def _run_vba(self, key, path):
        if not IS_WIN:
            msg = "VBA macros can only run on Windows with Microsoft Excel installed.\n"
            self._log(key, msg, "error")
            return "failed", 1, "", msg
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore
        except ImportError:
            msg = "pywin32 is not installed.  Open Env Health Check and install 'pywin32'.\n"
            self._log(key, msg, "error")
            return "failed", 1, "", msg
        code = Path(path).read_text(encoding="utf-8", errors="replace")
        subs = vba_subs(code)
        if not subs:
            msg = "No public parameter-less Sub found in this module - nothing to run.\n"
            self._log(key, msg, "error")
            return "failed", 1, "", msg
        module_name = Path(path).stem
        sub = next((s for s in subs if s.lower() == module_name.lower()), subs[0])
        pythoncom.CoInitialize()
        try:
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = True
            wb = excel.Workbooks.Add()
            try:
                wb.VBProject.VBComponents.Import(str(path))
            except Exception as exc:  # noqa: BLE001
                msg = ("Excel refused to import the macro.\n"
                       "→ In Excel: File ▸ Options ▸ Trust Center ▸ Trust Center Settings ▸ Macro Settings ▸ "
                       "tick 'Trust access to the VBA project object model'.\n"
                       f"Details: {exc}\n")
                self._log(key, msg, "error")
                return "failed", 1, "", msg
            self._log(key, f"Running Sub '{sub}' (available: {', '.join(subs)})\n", "info")
            try:
                excel.Run(f"'{wb.Name}'!{sub}")
            except Exception as exc:  # noqa: BLE001
                msg = f"Macro '{sub}' raised an error: {exc}\n"
                self._log(key, msg, "error")
                return "failed", 1, "", msg
            out = f"Macro '{sub}' completed in workbook '{wb.Name}' (left open in Excel).\n"
            self._log(key, out, "success")
            return "success", 0, out, ""
        finally:
            pythoncom.CoUninitialize()


# =============================================================================
#  SCHEDULING LOGIC  (pure functions - easy to test)
# =============================================================================
def _hhmm(at: str):
    try:
        h, m = at.split(":")
        return max(0, min(23, int(h))), max(0, min(59, int(m)))
    except Exception:  # noqa: BLE001
        return 9, 0


def _weekday_set(s: dict) -> set:
    return {int(x) for x in (s.get("weekdays") or "").split(",") if x.strip().isdigit()}


def is_due(s: dict, now: datetime) -> bool:
    if not s.get("enabled"):
        return False
    last, created = parse_iso(s.get("last_fire")), parse_iso(s.get("created"))
    if s["kind"] == "interval":
        base = last or created or now
        return (now - base) >= timedelta(minutes=max(1, int(s.get("interval_min") or 60)))
    h, m = _hhmm(s.get("at") or "09:00")
    occ = now.replace(hour=h, minute=m, second=0, microsecond=0)
    if s["kind"] == "weekly" and now.weekday() not in _weekday_set(s):
        return False
    if now < occ or (now - occ) > timedelta(minutes=15):      # don't fire hours late
        return False
    if created and created > occ:
        return False
    return last is None or last < occ


def next_run(s: dict, now: datetime | None = None):
    now = now or datetime.now()
    if not s.get("enabled"):
        return None
    last, created = parse_iso(s.get("last_fire")), parse_iso(s.get("created"))
    if s["kind"] == "interval":
        base = last or created or now
        return max(now, base + timedelta(minutes=max(1, int(s.get("interval_min") or 60))))
    h, m = _hhmm(s.get("at") or "09:00")
    days = _weekday_set(s) if s["kind"] == "weekly" else set(range(7))
    for offset in range(0, 9):
        cand = (now + timedelta(days=offset)).replace(hour=h, minute=m, second=0, microsecond=0)
        if cand.weekday() not in days:
            continue
        if cand <= now or (last and last >= cand):
            continue
        return cand
    return None


def describe_schedule(s: dict) -> str:
    if s["kind"] == "interval":
        return f"every {s.get('interval_min')} min"
    at = s.get("at") or "09:00"
    if s["kind"] == "weekly":
        names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
        days = ",".join(names[d] for d in sorted(_weekday_set(s)) if d < 7) or "—"
        return f"{days} at {at}"
    return f"daily at {at}"


class FolderWatcher:
    """Polling watcher (no external deps). Fires once per NEW, fully-written file."""

    def __init__(self):
        self.seen: dict = {}
        self.pending: dict = {}

    @staticmethod
    def _matching(folder: str, pattern: str) -> dict:
        found = {}
        try:
            with os.scandir(folder) as it:
                for e in it:
                    if e.is_file() and not e.name.startswith(("~$", ".")) and fnmatch.fnmatch(e.name.lower(), pattern.lower()):
                        found[e.path] = e.stat().st_size
        except OSError:
            pass
        return found

    def poll(self, watchers: list) -> list:
        hits, now = [], time.monotonic()
        for w in watchers:
            if not w.get("enabled"):
                continue
            wid = w["id"]
            current = self._matching(w["folder"], w.get("pattern") or "*")
            if wid not in self.seen:                       # first poll = baseline, ignore existing files
                self.seen[wid] = set(current)
                continue
            for path, size in current.items():
                if path in self.seen[wid]:
                    continue
                prev = self.pending.get((wid, path))
                if prev is None or prev[0] != size:
                    self.pending[(wid, path)] = (size, now)
                elif now - prev[1] >= 2:                  # size stable for 2s -> fully written
                    self.seen[wid].add(path)
                    self.pending.pop((wid, path), None)
                    hits.append((w, path))
        return hits


# =============================================================================
#  ACHIEVEMENTS
# =============================================================================
ACHIEVEMENTS = [
    ("first",    "First Rite",         "Complete your first successful run",     lambda s: s["ok"] >= 1),
    ("ten",      "Apprentice",         "10 successful runs",                     lambda s: s["ok"] >= 10),
    ("hundred",  "Adept",              "100 successful runs",                    lambda s: s["ok"] >= 100),
    ("fivehund", "Techmarine",         "500 successful runs",                    lambda s: s["ok"] >= 500),
    ("streak3",  "Three-Day Vigil",    "Successful runs 3 days in a row",        lambda s: s["streak"] >= 3),
    ("streak7",  "Seven-Day Vigil",    "Successful runs 7 days in a row",        lambda s: s["streak"] >= 7),
    ("streak30", "Iron Discipline",    "Successful runs 30 days in a row",       lambda s: s["streak"] >= 30),
    ("flawless", "Flawless",           "Last 10 runs all succeeded",             lambda s: s["flawless"]),
    ("night",    "Night Owl",          "Run something between midnight and 4am", lambda s: s["night"]),
    ("polyglot", "Polyglot",           "Succeed with both Python and VBA",       lambda s: s["polyglot"]),
    ("pipeline", "Pipeline Master",    "Run a pipeline",                         lambda s: s["pipelines"] >= 1),
]


# =============================================================================
#  GRAPHICS ENGINE  -  hand-drawn colour icons, gradient tiles, wordmark,
#                      glitch page transitions
# =============================================================================
import math
from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont

try:
    from PIL import ImageGrab, ImageTk
except Exception:  # noqa: BLE001  (ImageGrab needs a display; degrade to "no transitions")
    ImageGrab = None
    ImageTk = None


def hex_rgb(h: str) -> tuple:
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def blend(c1: str, c2: str, t: float) -> str:
    a, b = hex_rgb(c1), hex_rgb(c2)
    return "#" + "".join(f"{int(x + (y - x) * t):02x}" for x, y in zip(a, b))


# ---- fonts for PIL-rendered text ---------------------------------------------
def pil_font(px: int, bold: bool = True, mono: bool = False):
    files = (["JetBrainsMono-Bold.ttf", "CascadiaCode.ttf", "consolab.ttf", "DejaVuSansMono-Bold.ttf"] if mono else
             (["bahnschrift.ttf", "segoeuib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold
              else ["segoeui.ttf", "arial.ttf", "DejaVuSans.ttf"]))
    for f in files:
        try:
            return ImageFont.truetype(f, px)
        except Exception:  # noqa: BLE001
            continue
    try:
        return ImageFont.load_default(size=px)
    except TypeError:
        return ImageFont.load_default()


# ---- vector pen (draws into an "L" mask: 255 = ink, 0 = cut-out) ---------------
class _Pen:
    def __init__(self, n: int, inset: float):
        self.n = n
        self.img = Image.new("L", (n, n), 0)
        self.d = ImageDraw.Draw(self.img)
        self.k = n * (1 - 2 * inset) / 100.0
        self.o = n * inset
        self.sw = 8
        self.ink = 255

    def P(self, x, y):
        return (self.o + x * self.k, self.o + y * self.k)

    def cut(self):
        self.ink = 0
        return self

    def on(self):
        self.ink = 255
        return self

    def _w(self, w):
        return max(1, int(round((w or self.sw) * self.k)))

    def line(self, pts, w=None, closed=False):
        w = w or self.sw
        pp = [self.P(*p) for p in pts]
        if closed:
            pp.append(pp[0])
        self.d.line(pp, fill=self.ink, width=self._w(w), joint="curve")
        r = w * self.k / 2
        for x, y in pp:
            self.d.ellipse((x - r, y - r, x + r, y + r), fill=self.ink)

    def poly(self, pts):
        self.d.polygon([self.P(*p) for p in pts], fill=self.ink)

    def spoly(self, pts, w=6):                 # filled polygon with rounded corners
        self.poly(pts)
        self.line(pts, w, closed=True)

    def circle(self, cx, cy, r, fill=False, w=None):
        x, y = self.P(cx, cy)
        rr = r * self.k
        if fill:
            self.d.ellipse((x - rr, y - rr, x + rr, y + rr), fill=self.ink)
        else:
            self.d.ellipse((x - rr, y - rr, x + rr, y + rr), outline=self.ink, width=self._w(w))

    def rrect(self, x0, y0, x1, y1, r=8, fill=False, w=None):
        a, b = self.P(x0, y0), self.P(x1, y1)
        if fill:
            self.d.rounded_rectangle((*a, *b), radius=r * self.k, fill=self.ink)
        else:
            self.d.rounded_rectangle((*a, *b), radius=r * self.k, outline=self.ink, width=self._w(w))

    def arc(self, x0, y0, x1, y1, start, end, w=None):
        a, b = self.P(x0, y0), self.P(x1, y1)
        self.d.arc((*a, *b), start, end, fill=self.ink, width=self._w(w))
        cx, cy = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        rx, ry = (b[0] - a[0]) / 2, (b[1] - a[1]) / 2
        r = (w or self.sw) * self.k / 2
        for ang in (start, end):
            x = cx + rx * math.cos(math.radians(ang))
            y = cy + ry * math.sin(math.radians(ang))
            self.d.ellipse((x - r, y - r, x + r, y + r), fill=self.ink)


def _polar(cx, cy, r, deg):
    return (cx + r * math.cos(math.radians(deg)), cy + r * math.sin(math.radians(deg)))


def _bez(p0, p1, p2, n=24):
    out = []
    for i in range(n + 1):
        t = i / n
        out.append(((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
                    (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1]))
    return out


def _star_pts(cx, cy, ro, ri, n=5):
    pts = []
    for i in range(n * 2):
        pts.append(_polar(cx, cy, ro if i % 2 == 0 else ri, -90 + i * 180 / n))
    return pts


# ---- the glyph library (100 x 100 design grid) ---------------------------------
def g_bolt(p): p.spoly([(58, 6), (22, 56), (46, 56), (38, 94), (78, 40), (54, 40)], 5)
def g_home(p):
    p.sw = 8
    p.line([(10, 50), (50, 14), (90, 50)])
    p.line([(20, 44), (20, 88), (80, 88), (80, 44)])
    p.line([(42, 88), (42, 62), (58, 62), (58, 88)], 7)
def g_rocket(p):
    p.spoly([(50, 6), (66, 26), (70, 52), (70, 66), (30, 66), (30, 52), (34, 26)], 4)
    p.spoly([(30, 50), (14, 72), (14, 80), (30, 72)], 3)
    p.spoly([(70, 50), (86, 72), (86, 80), (70, 72)], 3)
    p.spoly([(40, 72), (50, 94), (60, 72)], 3)
    p.cut().circle(50, 40, 8, fill=True)
def g_pencil(p):
    p.spoly([(38, 8), (62, 8), (62, 64), (50, 92), (38, 64)], 3)
    p.cut().line([(38, 22), (62, 22)], 4)
    p.line([(38, 64), (62, 64)], 3)
def g_code(p):
    p.sw = 9
    p.line([(34, 26), (8, 50), (34, 74)])
    p.line([(66, 26), (92, 50), (66, 74)])
    p.line([(57, 18), (43, 82)], 8)
def g_flask(p):
    p.sw = 7
    p.line([(40, 10), (60, 10)])
    p.line([(43, 10), (43, 42), (14, 84), (14, 88), (86, 88), (86, 84), (57, 42), (57, 10)])
    p.poly([(25, 66), (75, 66), (86, 84), (86, 88), (14, 88), (14, 84)])
    p.cut().circle(40, 78, 4, fill=True)
    p.circle(58, 74, 3, fill=True)
def g_clock(p):
    p.sw = 8
    p.circle(50, 50, 38)
    p.line([(50, 26), (50, 50), (68, 62)])
def g_timer(p):
    p.sw = 8
    p.circle(50, 58, 31)
    p.line([(50, 58), (64, 43)])
    p.line([(40, 9), (60, 9)])
    p.line([(50, 9), (50, 26)])
    p.line([(79, 27), (86, 20)], 7)
def g_pulse(p):
    p.sw = 8
    p.line([(4, 54), (24, 54), (36, 26), (52, 82), (64, 38), (72, 54), (96, 54)])
def g_gear(p):
    for i in range(8):
        a = i * 45
        p.poly([_polar(50, 50, 28, a - 14), _polar(50, 50, 45, a - 9), _polar(50, 50, 45, a + 9), _polar(50, 50, 28, a + 14)])
    p.circle(50, 50, 33, fill=True)
    p.cut().circle(50, 50, 14, fill=True)
def g_play(p): p.spoly([(28, 16), (28, 84), (86, 50)], 8)
def g_stop(p): p.rrect(20, 20, 80, 80, 10, fill=True)
def g_trash(p):
    p.sw = 7
    p.line([(14, 26), (86, 26)])
    p.line([(36, 26), (38, 12), (62, 12), (64, 26)])
    p.line([(24, 26), (30, 88), (70, 88), (76, 26)])
    p.line([(42, 44), (43, 72)], 6)
    p.line([(58, 44), (57, 72)], 6)
def g_star(p): p.spoly(_star_pts(50, 54, 46, 20), 6)
def g_star_o(p):
    p.sw = 7
    p.line(_star_pts(50, 54, 44, 19), closed=True)
def g_tag(p):
    p.sw = 7
    p.line([(10, 12), (50, 12), (90, 52), (52, 90), (10, 50)], closed=True)
    p.circle(28, 30, 7, fill=True)
def g_chart(p):
    p.rrect(10, 54, 32, 90, 5, fill=True)
    p.rrect(39, 30, 61, 90, 5, fill=True)
    p.rrect(68, 8, 90, 90, 5, fill=True)
def g_folder(p):
    p.rrect(8, 32, 92, 88, 9, fill=True)
    p.spoly([(8, 36), (8, 20), (38, 20), (48, 34)], 4)
    p.cut().line([(12, 38), (88, 38)], 3)
def g_lock(p):
    p.rrect(20, 44, 80, 90, 9, fill=True)
    p.sw = 8
    p.arc(30, 12, 70, 56, 180, 360)
    p.line([(30, 34), (30, 48)])
    p.line([(70, 34), (70, 48)])
    p.cut().circle(50, 64, 6, fill=True)
    p.line([(50, 66), (50, 78)], 5)
def g_search(p):
    p.sw = 9
    p.circle(42, 42, 28)
    p.line([(63, 63), (88, 88)], 11)
def g_copy(p):
    p.sw = 7
    p.rrect(32, 32, 88, 90, 9)
    p.cut().rrect(10, 10, 70, 72, 9, fill=True)
    p.on().rrect(12, 12, 68, 70, 9)
def g_save(p):
    p.rrect(12, 12, 88, 88, 9, fill=True)
    p.cut().rrect(28, 12, 70, 38, 2, fill=True)
    p.on().rrect(56, 16, 64, 34, 1, fill=True)
    p.cut().rrect(24, 54, 76, 80, 3, fill=True)
def g_plus(p):
    p.sw = 11
    p.line([(50, 16), (50, 84)])
    p.line([(16, 50), (84, 50)])
def g_refresh(p):
    p.sw = 8
    p.arc(16, 16, 84, 84, 40, 330)
    ex, ey = _polar(50, 50, 34, 330)
    tx, ty = -math.sin(math.radians(330)), math.cos(math.radians(330))
    nx, ny = -ty, tx
    tip = (ex + tx * 13, ey + ty * 13)
    p.poly([tip, (ex - tx * 3 + nx * 13, ey - ty * 3 + ny * 13), (ex - tx * 3 - nx * 13, ey - ty * 3 - ny * 13)])
def g_terminal(p):
    p.sw = 7
    p.rrect(6, 16, 94, 84, 11)
    p.line([(22, 38), (40, 50), (22, 62)], 8)
    p.line([(48, 66), (70, 66)], 8)
def g_sparkle(p):
    p.spoly([(42, 14), (51, 47), (84, 56), (51, 65), (42, 98), (33, 65), (0, 56), (33, 47)], 3)
    p.spoly([(80, 4), (84, 17), (97, 21), (84, 25), (80, 38), (76, 25), (63, 21), (76, 17)], 2)
def g_download(p):
    p.sw = 9
    p.line([(50, 10), (50, 62)])
    p.line([(28, 42), (50, 64), (72, 42)])
    p.line([(16, 84), (84, 84)])
def g_eye(p):
    p.sw = 7
    up, lo = _bez((6, 50), (50, 6), (94, 50)), _bez((6, 50), (50, 94), (94, 50))
    p.line(up + lo[::-1], closed=True)
    p.circle(50, 50, 13, fill=True)
def g_link(p):
    p.sw = 8
    p.rrect(4, 32, 58, 68, 18)
    p.rrect(42, 32, 96, 68, 18)
def g_close(p):
    p.sw = 11
    p.line([(22, 22), (78, 78)])
    p.line([(78, 22), (22, 78)])
def g_check(p):
    p.sw = 11
    p.line([(18, 52), (40, 74), (84, 28)])
def g_warn(p):
    p.sw = 8
    p.line([(50, 10), (92, 86), (8, 86)], closed=True)
    p.line([(50, 38), (50, 60)], 8)
    p.circle(50, 73, 4.5, fill=True)
def g_info(p):
    p.sw = 8
    p.circle(50, 50, 40)
    p.line([(50, 44), (50, 70)], 9)
    p.circle(50, 29, 5, fill=True)
def g_palette(p):
    p.sw = 7
    p.circle(50, 50, 40)
    for x, y in ((32, 42), (48, 28), (66, 34), (70, 54)):
        p.circle(x, y, 6, fill=True)
def g_chip(p):
    p.sw = 7
    p.rrect(24, 24, 76, 76, 8)
    p.rrect(38, 38, 62, 62, 3, fill=True)
    for v in (38, 50, 62):
        p.line([(v, 8), (v, 24)], 6)
        p.line([(v, 76), (v, 92)], 6)
        p.line([(8, v), (24, v)], 6)
        p.line([(76, v), (92, v)], 6)
def g_trophy(p):
    p.spoly([(28, 10), (72, 10), (70, 44), (60, 58), (40, 58), (30, 44)], 4)
    p.sw = 6
    p.arc(4, 16, 34, 46, 90, 270)
    p.arc(66, 16, 96, 46, -90, 90)
    p.rrect(44, 56, 56, 76, 2, fill=True)
    p.rrect(30, 74, 70, 92, 4, fill=True)


GLYPHS = {n[2:]: f for n, f in list(globals().items()) if n.startswith("g_") and callable(f)}
ROTATE = {"rocket": -45, "pencil": -45, "link": -45}

# name -> (light colour, deep colour) : used for gradient tiles and tinted glyphs
ICON_STYLES = {
    "bolt": ("#38BDF8", "#6366F1"), "home": ("#38BDF8", "#6366F1"), "rocket": ("#FB923C", "#F43F5E"),
    "code": ("#A78BFA", "#6366F1"), "flask": ("#34D399", "#0D9488"), "clock": ("#FBBF24", "#F97316"),
    "timer": ("#C084FC", "#EC4899"), "pulse": ("#FB7185", "#E11D48"), "gear": ("#94A3B8", "#3B82F6"),
    "play": ("#4ADE80", "#16A34A"), "stop": ("#F87171", "#DC2626"), "trash": ("#FB7185", "#DC2626"),
    "pencil": ("#FBBF24", "#F59E0B"), "star": ("#FDE047", "#F59E0B"), "star_o": ("#94A3B8", "#64748B"),
    "tag": ("#D8B4FE", "#A855F7"), "chart": ("#22D3EE", "#3B82F6"), "folder": ("#FBBF24", "#F59E0B"),
    "lock": ("#2DD4BF", "#0891B2"), "search": ("#60A5FA", "#3B82F6"), "copy": ("#94A3B8", "#64748B"),
    "save": ("#4ADE80", "#16A34A"), "plus": ("#4ADE80", "#16A34A"), "refresh": ("#22D3EE", "#0EA5E9"),
    "terminal": ("#4ADE80", "#0D9488"), "sparkle": ("#F0ABFC", "#8B5CF6"), "download": ("#60A5FA", "#6366F1"),
    "eye": ("#38BDF8", "#818CF8"), "link": ("#FB923C", "#EC4899"), "close": ("#F87171", "#DC2626"),
    "check": ("#4ADE80", "#16A34A"), "warn": ("#FBBF24", "#F97316"), "info": ("#60A5FA", "#3B82F6"),
    "palette": ("#F472B6", "#A78BFA"), "chip": ("#94A3B8", "#3B82F6"), "trophy": ("#FDE047", "#F59E0B"),
}

# text symbols used across the UI  ->  icon name
EMOJI_ICON = {
    "▶": "play", "⏹": "stop", "✕": "close", "🗑": "trash", "🧹": "trash", "✏": "pencil", "✎": "pencil",
    "🏷": "tag", "🕘": "clock", "📊": "chart", "💾": "save", "🔍": "search", "🔎": "search", "📋": "copy",
    "⬇": "download", "＋": "plus", "🔄": "refresh", "✨": "sparkle", "📁": "folder", "📂": "folder",
    "⏯": "check", "🔒": "lock", "💬": "code", "⌨": "search", "⚡": "bolt", "🩺": "pulse", "🚀": "rocket",
    "📝": "code", "🧪": "flask", "⏱": "timer", "⚙": "gear", "🏠": "home", "🔗": "link", "👁": "eye",
    "🛠": "gear", "🎨": "palette", "🧰": "chip", "📦": "chip", "🏆": "trophy", "🐍": "code", "📈": "chart",
}


def split_icon(text: str):
    """'▶ Run' -> ('play', 'Run').  Returns (None, text) when no known leading symbol."""
    t = text.lstrip()
    for sym, name in EMOJI_ICON.items():
        if t.startswith(sym):
            rest = t[len(sym):].lstrip(" \u00a0\ufe0f")
            return name, rest
    return None, text


def _glyph_mask(name: str, final: int, inset: float) -> Image.Image:
    n = final * 2
    pen = _Pen(n, inset)
    GLYPHS[name](pen)
    m = pen.img
    if name in ROTATE:
        m = m.rotate(ROTATE[name], resample=Image.BICUBIC)
    return m.resize((final, final), Image.LANCZOS)


def _diag_gradient(size: int, c1: str, c2: str) -> Image.Image:
    base, top = Image.new("RGB", (size, size), c1), Image.new("RGB", (size, size), c2)
    v = Image.linear_gradient("L").resize((size, size))
    h = v.transpose(Image.Transpose.ROTATE_90)
    return Image.composite(top, base, ImageChops.add(h, v, scale=2.0))


def _rounded_mask(size: int, radius_frac: float = 0.27) -> Image.Image:
    big = Image.new("L", (size * 4, size * 4), 0)
    ImageDraw.Draw(big).rounded_rectangle((0, 0, size * 4 - 1, size * 4 - 1), radius=size * 4 * radius_frac, fill=255)
    return big.resize((size, size), Image.LANCZOS)


def render_glyph(name: str, size: int, color: str) -> Image.Image:
    f = size * 3
    out = Image.new("RGBA", (f, f), (0, 0, 0, 0))
    mask = _glyph_mask(name, f, 0.08)
    out.paste(Image.new("RGBA", (f, f), hex_rgb(color) + (255,)), (0, 0), mask)
    return out


def render_tile(name: str, size: int, c1: str, c2: str, glyph: str = "#FFFFFF", glow: bool = True) -> Image.Image:
    f = size * 3
    tile = _diag_gradient(f, c1, c2).convert("RGBA")
    if glow:                                                    # soft top highlight for a glassy look
        hl = Image.linear_gradient("L").resize((f, f)).transpose(Image.Transpose.FLIP_TOP_BOTTOM)
        hl = hl.point(lambda v: int(v * 0.20))
        tile.paste((255, 255, 255, 255), (0, 0), hl)
    if name:
        mask = _glyph_mask(name, f, 0.23)
        # subtle drop shadow under the glyph
        shadow = mask.filter(ImageFilter.GaussianBlur(f * 0.02)).point(lambda v: int(v * 0.35))
        tile.paste((0, 0, 0, 255), (0, int(f * 0.025)), shadow)
        tile.paste(Image.new("RGBA", (f, f), hex_rgb(glyph) + (255,)), (0, 0), mask)
    out = Image.new("RGBA", (f, f), (0, 0, 0, 0))
    out.paste(tile, (0, 0), _rounded_mask(f))
    return out


_ICON_CACHE: dict = {}


def _ctk_image(pil: Image.Image, size) -> "ctk.CTkImage":
    return ctk.CTkImage(light_image=pil, dark_image=pil, size=(size, size))


def icon(name: str, size: int = 18, color: str | None = None):
    """Coloured glyph (no tile). color=None uses the icon's own palette."""
    key = ("g", name, size, color)
    if key not in _ICON_CACHE:
        col = color or ICON_STYLES.get(name, ("#94A3B8", "#64748B"))[0]
        _ICON_CACHE[key] = _ctk_image(render_glyph(name, size, col), size)
    return _ICON_CACHE[key]


def icon_tile(name: str, size: int = 34, styles: tuple | None = None):
    """Gradient tile with a white glyph - the 'colourful icon' used for nav, headers, stat cards."""
    st = styles or ICON_STYLES.get(name, ("#94A3B8", "#64748B"))
    key = ("t", name, size, st)
    if key not in _ICON_CACHE:
        _ICON_CACHE[key] = _ctk_image(render_tile(name, size, st[0], st[1]), size)
    return _ICON_CACHE[key]


LANG_STYLES = {"py": ("#3B82F6", "#FBBF24", "Py"), "vba": ("#059669", "#34D399", "XL")}


def lang_tile(kind: str, size: int = 40):
    key = ("l", kind, size)
    if key not in _ICON_CACHE:
        c1, c2, label = LANG_STYLES[kind]
        f = size * 3
        img = render_tile("", size, c1, c2)
        d = ImageDraw.Draw(img)
        font = pil_font(int(f * 0.42), True)
        box = d.textbbox((0, 0), label, font=font)
        d.text(((f - (box[2] - box[0])) / 2 - box[0], (f - (box[3] - box[1])) / 2 - box[1]), label, font=font, fill="#FFFFFF")
        _ICON_CACHE[key] = _ctk_image(img, size)
    return _ICON_CACHE[key]


def gradient_wordmark(text: str, px: int, c1: str, c2: str, spacing: int = 0):
    """Gradient-filled text as a CTkImage (fallback: None so callers can use a plain label)."""
    try:
        f = px * 3
        font = pil_font(f, True)
        tmp = ImageDraw.Draw(Image.new("L", (10, 10)))
        widths = [tmp.textlength(ch, font=font) + spacing * 3 for ch in text]
        w, h = int(sum(widths)) + 8, int(f * 1.35)
        mask = Image.new("L", (w, h), 0)
        md = ImageDraw.Draw(mask)
        x = 4
        for ch, cw in zip(text, widths):
            md.text((x, int(f * 0.1)), ch, font=font, fill=255)
            x += cw
        grad = Image.new("RGB", (w, h), c1)
        ramp = Image.linear_gradient("L").resize((h, w)).transpose(Image.Transpose.ROTATE_90).resize((w, h))
        grad = Image.composite(Image.new("RGB", (w, h), c2), grad, ramp)
        out = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        out.paste(grad, (0, 0), mask)
        return ctk.CTkImage(light_image=out, dark_image=out, size=(w // 3, h // 3))
    except Exception as exc:  # noqa: BLE001
        print(f"[wordmark] {exc}")
        return None


def make_app_icon_file(path: Path) -> bool:
    try:
        big = render_tile("bolt", 256, COLOR_ACCENT, COLOR_ACCENT2).resize((256, 256), Image.LANCZOS)
        big.save(path, format="ICO", sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"[icon] {exc}")
        return False


# =============================================================================
#  GLITCH TRANSITION
#  Old page is sliced into strips that tear sideways, split into RGB channels,
#  flicker and dissolve away in random bands to reveal the new page underneath.
# =============================================================================
GLITCH_LEVELS = {"Subtle": (280, 22), "Normal": (430, 52), "Extreme": (620, 110)}


class GlitchTransition:
    def __init__(self, host, level: str = "Normal", caption: str = ""):
        self.host = host
        self.duration, self.amp = GLITCH_LEVELS.get(level, GLITCH_LEVELS["Normal"])
        self.caption = caption
        self.strips: list = []
        self.bars: list = []
        self.images: list = []
        self.active = False
        self.cap_label = None
        self._job = None

    @staticmethod
    def available() -> bool:
        return ImageGrab is not None and ImageTk is not None

    def start(self, switch_fn, done=None) -> bool:
        """Capture -> cover with strips -> run switch_fn underneath -> animate. False = nothing happened."""
        if not self.available():
            return False
        host = self.host
        try:
            host.update_idletasks()
            w, h = host.winfo_width(), host.winfo_height()
            x, y = host.winfo_rootx(), host.winfo_rooty()
            if w < 80 or h < 80:
                return False
            shot = ImageGrab.grab(bbox=(x, y, x + w, y + h), all_screens=True).convert("RGB")
            if shot.size != (w, h):
                shot = shot.resize((w, h))
            r, g, b = shot.split()
            split = Image.merge("RGB", (ImageChops.offset(r, 7, 0), g, ImageChops.offset(b, -7, 0)))
            swapped = Image.merge("RGB", (b, g, r))
            self.images = [ImageTk.PhotoImage(shot, master=host), ImageTk.PhotoImage(split, master=host),
                           ImageTk.PhotoImage(swapped, master=host)]
        except Exception as exc:  # noqa: BLE001
            print(f"[glitch] capture failed: {exc}")
            return False
        self.done = done
        self.w, self.h = w, h
        y0 = 0
        while y0 < h:
            sh = min(h - y0, random.randint(8, 46))
            cv = tk.Canvas(host, width=w, height=sh, highlightthickness=0, bd=0, bg=COLOR_BG_MAIN)
            item = cv.create_image(0, -y0, image=self.images[0], anchor="nw")
            cv.place(x=0, y=y0, width=w, height=sh)
            self.strips.append({"cv": cv, "item": item, "y": y0, "h": sh, "dx": 0,
                                "vanish": random.uniform(0.28, 0.93), "alive": True, "blk": None})
            y0 += sh
        self.active = True
        try:
            switch_fn()                      # new page is built / raised underneath the strips
            host.update_idletasks()
        finally:
            for s in self.strips:
                tk.Misc.tkraise(s["cv"])
        if self.caption:
            self.cap_label = tk.Label(host, text="", bg=COLOR_BG_MAIN, fg=COLOR_ACCENT2, font=(FONT_MONO, 11, "bold"),
                                      padx=8, pady=3)
            self.cap_label.place(x=14, y=12)
            self.cap_label.lift()
        self.t0 = time.perf_counter()
        self._job = host.after(1, self._step)
        return True

    def abort(self):
        if self._job:
            try:
                self.host.after_cancel(self._job)
            except tk.TclError:
                pass
        self._cleanup()

    def _cleanup(self):
        self.active = False
        for s in self.strips:
            try:
                s["cv"].destroy()
            except tk.TclError:
                pass
        for b in self.bars:
            try:
                b.destroy()
            except tk.TclError:
                pass
        if self.cap_label is not None:
            try:
                self.cap_label.destroy()
            except tk.TclError:
                pass
        self.strips, self.bars, self.images = [], [], []
        if self.done:
            cb, self.done = self.done, None
            cb()

    def _step(self):
        if not self.active:
            return
        p = (time.perf_counter() - self.t0) / (self.duration / 1000.0)
        if p >= 1.0:
            self._cleanup()
            return
        amp = self.amp * (0.35 + 0.65 * math.sin(min(1.0, p / 0.6) * math.pi / 2))
        for s in self.strips:
            if not s["alive"]:
                continue
            cv = s["cv"]
            if p >= s["vanish"]:
                s["alive"] = False
                cv.place_forget()
                continue
            if s["vanish"] - p < 0.07 and random.random() < 0.45:        # flicker just before vanishing
                cv.place_forget()
                continue
            cv.place(x=0, y=s["y"], width=self.w, height=s["h"])
            if random.random() < 0.30 + 0.55 * p:
                s["dx"] = int(random.randint(-int(amp), int(amp)) * random.choice((0.4, 0.7, 1.0)))
            else:
                s["dx"] = int(s["dx"] * 0.4)
            cv.coords(s["item"], s["dx"], -s["y"])
            roll = random.random()
            cv.itemconfigure(s["item"], image=self.images[1 if roll < 0.35 else 2 if roll < 0.42 else 0])
            if s["blk"] is not None:
                cv.delete(s["blk"])
                s["blk"] = None
            if random.random() < 0.06 and s["h"] > 10:                   # neon block artefact
                bx = random.randint(0, max(1, self.w - 120))
                s["blk"] = cv.create_rectangle(bx, 0, bx + random.randint(40, 260), s["h"],
                                               fill=random.choice((COLOR_ACCENT2, COLOR_ACCENT, "#FF2E88", "#FFFFFF")),
                                               outline="", stipple=random.choice(("gray50", "gray25", "gray75")))
            tk.Misc.tkraise(cv)
        for b in self.bars:                                              # thin horizontal glitch bars
            try:
                b.destroy()
            except tk.TclError:
                pass
        self.bars = []
        if random.random() < 0.85 * (1 - p * 0.6):
            for _ in range(random.randint(1, 3)):
                bar = tk.Frame(self.host, bg=random.choice((COLOR_ACCENT2, COLOR_ACCENT, "#FF2E88", "#E8EEF9")))
                bw = random.randint(80, max(120, int(self.w * 0.7)))
                bar.place(x=random.randint(0, max(1, self.w - bw)), y=random.randint(0, self.h - 6),
                          width=bw, height=random.randint(2, 9))
                self.bars.append(bar)
        if self.cap_label is not None:
            txt = self.caption
            if random.random() < 0.5:
                i = random.randrange(len(txt)) if txt else 0
                txt = txt[:i] + random.choice("#%&$01<>/\\") + txt[i + 1:]
            self.cap_label.configure(text=txt if p < 0.7 else "")
            self.cap_label.lift()
        self._job = self.host.after(14, self._step)


# =============================================================================
#  UI TOOLKIT  (buttons, modals, dialogs)
# =============================================================================
def setup_ttk_style() -> None:
    s = ttk.Style()
    try:
        s.theme_use("clam")
    except tk.TclError:
        pass
    s.configure("Treeview", background=COLOR_BG_MAIN, foreground=COLOR_TEXT_MAIN,
                fieldbackground=COLOR_BG_MAIN, borderwidth=0, rowheight=30, font=(FONT_FAMILY, 11))
    s.configure("Treeview.Heading", background=COLOR_SIDEBAR, foreground=COLOR_TEXT_MAIN, relief="flat",
                font=(FONT_FAMILY, 11, "bold"))
    s.layout("Treeview", [("Treeview.treearea", {"sticky": "nswe"})])       # drop the light default border
    s.map("Treeview", background=[("selected", COLOR_CARD_HOVER)], foreground=[("selected", COLOR_ACCENT)])
    s.map("Treeview.Heading", background=[("active", COLOR_CARD_HOVER)])
    for orient in ("Vertical", "Horizontal"):
        s.configure(f"{orient}.TScrollbar", background=COLOR_BORDER, troughcolor=COLOR_BG_MAIN,
                    bordercolor=COLOR_BG_MAIN, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER,
                    arrowcolor=COLOR_TEXT_MUTED, relief="flat", gripcount=0)
        s.map(f"{orient}.TScrollbar", background=[("active", COLOR_TEXT_MUTED)])


KIND_TINT = {"success": "#FFFFFF", "danger": "#FFFFFF", "primary": "#0A0E17"}


def make_button(parent, text, command=None, kind="ghost", width=None, height=32, glyph=None, **kw):
    """Buttons pick up a colour icon automatically from a leading symbol in `text` (e.g. '▶ Run')."""
    name, rest = (glyph, text) if glyph else split_icon(text)
    styles = {
        "primary": dict(fg_color=COLOR_ACCENT, hover_color=COLOR_ACCENT_HOVER, text_color="#0A0E17"),
        "success": dict(fg_color=COLOR_SUCCESS, hover_color=COLOR_SUCCESS_HOVER, text_color="#FFFFFF"),
        "danger": dict(fg_color=COLOR_DANGER, hover_color=COLOR_DANGER_HOVER, text_color="#FFFFFF"),
        "ghost": dict(fg_color="transparent", border_width=1, border_color=COLOR_BORDER,
                      hover_color=COLOR_CARD_HOVER, text_color=COLOR_TEXT_MAIN),
        "outline": dict(fg_color="transparent", border_width=1, border_color=COLOR_ACCENT,
                        hover_color=COLOR_CARD_HOVER, text_color=COLOR_ACCENT),
        "flat": dict(fg_color="transparent", hover_color=COLOR_CARD_HOVER, text_color=COLOR_TEXT_MAIN),
    }
    opts = dict(font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"), corner_radius=8, height=height)
    if name:
        opts["image"] = icon(name, 16, KIND_TINT.get(kind))
        opts["compound"] = "left"
        text = rest
        if not rest and width is None:
            width = height + 4
    if width:
        opts["width"] = width
    opts.update(styles[kind])
    opts.update(kw)
    return ctk.CTkButton(parent, text=text, command=command, **opts)


def make_entry(parent, placeholder="", width=200, **kw):
    opts = dict(placeholder_text=placeholder, width=width, fg_color=COLOR_BG_MAIN, border_color=COLOR_BORDER,
                corner_radius=8, height=36, font=ctk.CTkFont(family=FONT_FAMILY, size=13))
    opts.update(kw)
    e = ctk.CTkEntry(parent, **opts)
    e.bind("<FocusIn>", lambda _e: e.configure(border_color=COLOR_ACCENT), add="+")
    e.bind("<FocusOut>", lambda _e: e.configure(border_color=COLOR_BORDER), add="+")
    return e


def make_label(parent, text, size=13, bold=False, muted=False, display=False, **kw):
    """Labels get a colour icon from a leading symbol. Big bold text uses the display font."""
    name, rest = split_icon(text) if text else (None, text)
    family = FONT_DISPLAY if (display or (bold and size >= 15)) else FONT_FAMILY
    opts = dict(text_color=COLOR_TEXT_MUTED if muted else COLOR_TEXT_MAIN,
                font=ctk.CTkFont(family=family, size=size, weight="bold" if bold else "normal"))
    if name and "image" not in kw:
        opts["image"] = icon(name, max(14, int(size * 1.25) + 2))
        opts["compound"] = "left"
        text = "  " + rest if rest else ""
    opts.update(kw)
    return ctk.CTkLabel(parent, text=text, **opts)


def tk_text(box):
    """The real tk.Text behind a CTkTextbox (falls back to the widget itself)."""
    return getattr(box, "_textbox", box)


def run_bg(widget, fn, callback):
    """Run fn() on a worker thread, then callback(kind, value) on the UI thread ('ok' | 'err')."""
    q: queue.Queue = queue.Queue()

    def work():
        try:
            q.put(("ok", fn()))
        except Exception as exc:  # noqa: BLE001
            q.put(("err", exc))

    threading.Thread(target=work, daemon=True).start()

    def poll():
        try:
            kind, val = q.get_nowait()
        except queue.Empty:
            try:
                if widget.winfo_exists():
                    widget.after(80, poll)
            except tk.TclError:
                pass
            return
        callback(kind, val)

    widget.after(80, poll)


class Modal(ctk.CTkToplevel):
    def __init__(self, parent, title, width=520, height=360, modal=True):
        super().__init__(parent)
        self.result = None
        self._modal = modal
        self.title(title)
        self.configure(fg_color=COLOR_BG_MAIN)
        root = parent.winfo_toplevel()
        x = root.winfo_rootx() + max(0, (root.winfo_width() - width) // 2)
        y = root.winfo_rooty() + max(0, (root.winfo_height() - height) // 3)
        self.geometry(f"{width}x{height}+{x}+{y}")
        self.transient(root)
        self.bind("<Escape>", lambda e: self.close())
        self.after(150, self._grab)

    def _grab(self):
        try:
            if self._modal:
                self.grab_set()
            self.focus_force()
        except tk.TclError:
            pass

    def finish(self, result=True):
        self.result = result
        self.close()

    def close(self):
        try:
            self.grab_release()
        except tk.TclError:
            pass
        self.destroy()


def ask_confirm(parent, title, message, yes_text="Confirm", danger=False) -> bool:
    dlg = Modal(parent, title, 460, 210)
    make_label(dlg, title, 16, True).pack(anchor="w", padx=24, pady=(20, 6))
    make_label(dlg, message, 12, muted=True, wraplength=410, justify="left").pack(anchor="w", padx=24)
    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(side="bottom", fill="x", padx=24, pady=18)
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, yes_text, lambda: dlg.finish(True), "danger" if danger else "primary", 120).pack(side="right")
    parent.wait_window(dlg)
    return bool(dlg.result)


def ask_text(parent, title, label, initial="", width=420):
    dlg = Modal(parent, title, width + 60, 200)
    make_label(dlg, label, 12, True).pack(anchor="w", padx=24, pady=(22, 6))
    entry = make_entry(dlg, width=width)
    entry.pack(padx=24, anchor="w")
    entry.insert(0, initial)
    entry.focus_set()
    entry.bind("<Return>", lambda e: dlg.finish(entry.get()))
    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(side="bottom", fill="x", padx=24, pady=18)
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, "OK", lambda: dlg.finish(entry.get()), "primary", 100).pack(side="right")
    parent.wait_window(dlg)
    return dlg.result


def render_diff(box, old: str, new: str, old_label="current", new_label="new") -> int:
    tw = tk_text(box)
    tw.configure(state="normal")
    tw.delete("1.0", "end")
    for tag, color in (("add", "#56d364"), ("del", "#ff7b72"), ("hunk", COLOR_ACCENT), ("ctx", COLOR_TEXT_MUTED)):
        tw.tag_configure(tag, foreground=color)
    changed = 0
    lines = list(difflib.unified_diff(old.splitlines(), new.splitlines(), old_label, new_label, lineterm="", n=2))
    if not lines:
        tw.insert("end", "(no differences)\n", "ctx")
    for ln in lines:
        if ln.startswith("+") and not ln.startswith("+++"):
            tw.insert("end", ln + "\n", "add"); changed += 1
        elif ln.startswith("-") and not ln.startswith("---"):
            tw.insert("end", ln + "\n", "del"); changed += 1
        elif ln.startswith("@@") or ln.startswith(("---", "+++")):
            tw.insert("end", ln + "\n", "hunk")
        else:
            tw.insert("end", ln + "\n", "ctx")
    tw.configure(state="disabled")
    return changed


def ask_diff(parent, title, old, new, ok_text="Save", old_label="on disk", new_label="editor") -> bool:
    dlg = Modal(parent, title, 860, 560)
    make_label(dlg, title, 16, True).pack(anchor="w", padx=24, pady=(18, 4))
    make_label(dlg, "Review what will change before it is written.", 12, muted=True).pack(anchor="w", padx=24)
    box = ctk.CTkTextbox(dlg, fg_color=COLOR_CONSOLE_BG, font=(FONT_MONO, 12), border_width=1,
                         border_color=COLOR_BORDER, wrap="none")
    box.pack(fill="both", expand=True, padx=24, pady=12)
    render_diff(box, old, new, old_label, new_label)
    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(fill="x", padx=24, pady=(0, 16))
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, ok_text, lambda: dlg.finish(True), "success", 130).pack(side="right")
    parent.wait_window(dlg)
    return bool(dlg.result)


# ---- script parameter form --------------------------------------------------
def ask_params(parent, script: str, params: list, saved: dict, trigger_file: str | None = None):
    height = min(640, 200 + 82 * len(params))
    dlg = Modal(parent, f"Parameters · {script}", 660, height)
    make_label(dlg, f"▶ Run {script}", 16, True).pack(anchor="w", padx=24, pady=(18, 0))
    make_label(dlg, "Values are remembered for next time.", 11, muted=True).pack(anchor="w", padx=24, pady=(0, 6))
    scroll = ctk.CTkScrollableFrame(dlg, fg_color="transparent")
    scroll.pack(fill="both", expand=True, padx=14)
    vars_: dict = {}
    assigned_trigger = False
    for p in params:
        row = ctk.CTkFrame(scroll, fg_color=COLOR_CARD_SURFACE, corner_radius=8, border_width=1, border_color=COLOR_BORDER)
        row.pack(fill="x", padx=6, pady=5)
        make_label(row, p.name + ("  *" if p.required else ""), 13, True).pack(anchor="w", padx=12, pady=(8, 0))
        if p.help:
            make_label(row, p.help, 11, muted=True).pack(anchor="w", padx=12)
        value = saved.get(p.name, p.default)
        if trigger_file and p.type == "file" and not assigned_trigger:
            value, assigned_trigger = trigger_file, True
        line = ctk.CTkFrame(row, fg_color="transparent")
        line.pack(fill="x", padx=12, pady=(4, 10))
        if p.type == "bool":
            var = tk.BooleanVar(value=str(value).lower() in ("1", "true", "yes", "on"))
            ctk.CTkSwitch(line, text="", variable=var).pack(anchor="w")
        elif p.choices:
            var = tk.StringVar(value=str(value) if str(value) else str(p.choices[0]))
            ctk.CTkOptionMenu(line, values=[str(c) for c in p.choices], variable=var, fg_color=COLOR_BG_MAIN,
                              button_color=COLOR_BORDER).pack(anchor="w")
        else:
            var = tk.StringVar(value=str(value))
            make_entry(line, width=400, textvariable=var).pack(side="left", fill="x", expand=True)
            if p.type in ("file", "folder"):
                def browse(v=var, kind=p.type):
                    path = filedialog.askopenfilename() if kind == "file" else filedialog.askdirectory()
                    if path:
                        v.set(path)
                make_button(line, "Browse…", browse, "ghost", 90).pack(side="left", padx=(8, 0))
        vars_[p.name] = var
    err = make_label(dlg, "", 11)
    err.configure(text_color=COLOR_DANGER_HOVER)
    err.pack(anchor="w", padx=24)

    def submit():
        out = {}
        for p in params:
            v = vars_[p.name].get()
            if isinstance(v, bool):
                out[p.name] = "true" if v else "false"
                continue
            v = str(v).strip()
            if p.required and not v:
                err.configure(text=f"'{p.name}' is required.")
                return
            try:
                if v and p.type == "int":
                    int(v)
                elif v and p.type == "float":
                    float(v)
            except ValueError:
                err.configure(text=f"'{p.name}' must be a {p.type}.")
                return
            out[p.name] = v
        dlg.finish(out)

    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(fill="x", padx=24, pady=(4, 16))
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, "▶ Run", submit, "success", 120).pack(side="right")
    parent.wait_window(dlg)
    return dlg.result


# ---- schedule / pipeline / watcher dialogs -------------------------
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def ask_schedule(parent, targets: list, existing: dict | None = None):
    """targets: [(label, 'script:x.py' | 'pipe:Name'), ...]"""
    if not targets:
        return None
    dlg = Modal(parent, "Schedule", 560, 470)
    make_label(dlg, "⏱ Schedule a run", 16, True).pack(anchor="w", padx=24, pady=(18, 10))
    labels = [t[0] for t in targets]
    lookup = dict(targets)
    cur = existing or {}
    cur_label = next((l for l, t in targets if t == cur.get("target")), labels[0])
    target_var = tk.StringVar(value=cur_label)
    make_label(dlg, "What to run", 12, True).pack(anchor="w", padx=24)
    ctk.CTkOptionMenu(dlg, values=labels, variable=target_var, width=380, fg_color=COLOR_BG_MAIN,
                      button_color=COLOR_BORDER).pack(anchor="w", padx=24, pady=(4, 12))
    kind_map = {"Daily": "daily", "Weekly": "weekly", "Every N minutes": "interval"}
    inv = {v: k for k, v in kind_map.items()}
    kind_var = tk.StringVar(value=inv.get(cur.get("kind", "daily"), "Daily"))
    make_label(dlg, "Repeat", 12, True).pack(anchor="w", padx=24)
    ctk.CTkSegmentedButton(dlg, values=list(kind_map), variable=kind_var, selected_color=COLOR_ACCENT,
                           selected_hover_color=COLOR_ACCENT_HOVER, text_color=COLOR_TEXT_MAIN
                           ).pack(anchor="w", padx=24, pady=(4, 12))
    line = ctk.CTkFrame(dlg, fg_color="transparent")
    line.pack(anchor="w", padx=24)
    make_label(line, "Time (HH:MM)", 12, True).grid(row=0, column=0, sticky="w")
    make_label(line, "Interval (minutes)", 12, True).grid(row=0, column=1, sticky="w", padx=(24, 0))
    time_e = make_entry(line, width=120)
    time_e.grid(row=1, column=0, pady=4)
    time_e.insert(0, cur.get("at", "09:00"))
    int_e = make_entry(line, width=120)
    int_e.grid(row=1, column=1, pady=4, padx=(24, 0))
    int_e.insert(0, str(cur.get("interval_min", 60)))
    make_label(dlg, "Days (weekly)", 12, True).pack(anchor="w", padx=24, pady=(12, 0))
    days_row = ctk.CTkFrame(dlg, fg_color="transparent")
    days_row.pack(anchor="w", padx=24)
    picked = {int(x) for x in (cur.get("weekdays") or "0,1,2,3,4").split(",") if x.strip().isdigit()}
    day_vars = []
    for i, name in enumerate(WEEKDAYS):
        v = tk.BooleanVar(value=i in picked)
        ctk.CTkCheckBox(days_row, text=name, variable=v, width=60, checkbox_width=18, checkbox_height=18).pack(side="left")
        day_vars.append(v)
    make_label(dlg, "Runs only while Finprime's Engine is open (it can live in the system tray).", 11, muted=True
               ).pack(anchor="w", padx=24, pady=(12, 0))
    err = make_label(dlg, "", 11)
    err.configure(text_color=COLOR_DANGER_HOVER)
    err.pack(anchor="w", padx=24)

    def save():
        kind = kind_map[kind_var.get()]
        at = time_e.get().strip()
        if not re.fullmatch(r"([01]?\d|2[0-3]):[0-5]\d", at):
            err.configure(text="Time must look like 09:30")
            return
        try:
            minutes = int(int_e.get())
            assert minutes >= 1
        except Exception:  # noqa: BLE001
            err.configure(text="Interval must be a whole number of minutes (1+)")
            return
        wd = ",".join(str(i) for i, v in enumerate(day_vars) if v.get())
        if kind == "weekly" and not wd:
            err.configure(text="Pick at least one weekday")
            return
        h, m = at.split(":")
        dlg.finish({"target": lookup[target_var.get()], "kind": kind, "at": f"{int(h):02d}:{m}",
                    "weekdays": wd, "interval_min": minutes})

    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(side="bottom", fill="x", padx=24, pady=16)
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, "Save schedule", save, "success", 140).pack(side="right")
    parent.wait_window(dlg)
    return dlg.result


def ask_pipeline(parent, scripts: list, existing: dict | None = None):
    dlg = Modal(parent, "Pipeline", 600, 600)
    cur = existing or {}
    make_label(dlg, "🔗 Pipeline builder", 16, True).pack(anchor="w", padx=24, pady=(18, 6))
    make_label(dlg, "Name", 12, True).pack(anchor="w", padx=24)
    name_e = make_entry(dlg, width=360)
    name_e.pack(anchor="w", padx=24, pady=(4, 10))
    name_e.insert(0, cur.get("name", ""))
    make_label(dlg, "Steps (run top to bottom)", 12, True).pack(anchor="w", padx=24)
    steps: list = list(cur.get("steps", []))
    holder = ctk.CTkScrollableFrame(dlg, fg_color=COLOR_CARD_SURFACE, height=230, border_width=1, border_color=COLOR_BORDER)
    holder.pack(fill="x", padx=24, pady=(4, 8))

    def redraw():
        for w in holder.winfo_children():
            w.destroy()
        if not steps:
            make_label(holder, "No steps yet - add scripts below.", 12, muted=True).pack(pady=20)
        for i, s in enumerate(steps):
            r = ctk.CTkFrame(holder, fg_color="transparent")
            r.pack(fill="x", pady=2)
            make_label(r, f"{i + 1}.  {s}", 13).pack(side="left", padx=8)
            make_button(r, "✕", lambda i=i: (steps.pop(i), redraw()), "flat", 30).pack(side="right")
            make_button(r, "▼", lambda i=i: (steps.insert(i + 1, steps.pop(i)), redraw()) if i < len(steps) - 1 else None, "flat", 30).pack(side="right")
            make_button(r, "▲", lambda i=i: (steps.insert(i - 1, steps.pop(i)), redraw()) if i > 0 else None, "flat", 30).pack(side="right")

    redraw()
    add_row = ctk.CTkFrame(dlg, fg_color="transparent")
    add_row.pack(anchor="w", padx=24)
    pick = tk.StringVar(value=scripts[0] if scripts else "")
    if scripts:
        ctk.CTkOptionMenu(add_row, values=scripts, variable=pick, width=300, fg_color=COLOR_BG_MAIN,
                          button_color=COLOR_BORDER).pack(side="left")
        make_button(add_row, "＋ Add step", lambda: (steps.append(pick.get()), redraw()), "outline", 120).pack(side="left", padx=8)
    stop_var = tk.BooleanVar(value=bool(cur.get("stop_on_fail", 1)))
    ctk.CTkSwitch(dlg, text="Stop the pipeline if a step fails", variable=stop_var).pack(anchor="w", padx=24, pady=12)
    err = make_label(dlg, "", 11)
    err.configure(text_color=COLOR_DANGER_HOVER)
    err.pack(anchor="w", padx=24)

    def save():
        name = safe_name(name_e.get())
        if not name:
            err.configure(text="Give the pipeline a name.")
            return
        if not steps:
            err.configure(text="Add at least one step.")
            return
        dlg.finish({"name": name, "steps": steps, "stop_on_fail": stop_var.get()})

    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(side="bottom", fill="x", padx=24, pady=16)
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, "Save pipeline", save, "success", 140).pack(side="right")
    parent.wait_window(dlg)
    return dlg.result


def ask_watcher(parent, targets: list):
    if not targets:
        return None
    dlg = Modal(parent, "Folder watcher", 600, 400)
    make_label(dlg, "👁 Watch a folder", 16, True).pack(anchor="w", padx=24, pady=(18, 4))
    make_label(dlg, "When a new matching file finishes arriving, the target runs. The file path is passed as "
                    "FP_TRIGGER_FILE (and fills the first 'file' parameter).", 11, muted=True,
               wraplength=540, justify="left").pack(anchor="w", padx=24, pady=(0, 10))
    make_label(dlg, "Folder", 12, True).pack(anchor="w", padx=24)
    frow = ctk.CTkFrame(dlg, fg_color="transparent")
    frow.pack(anchor="w", padx=24, pady=(4, 10))
    folder_e = make_entry(frow, width=400)
    folder_e.pack(side="left")
    make_button(frow, "Browse…", lambda: (lambda p: (folder_e.delete(0, "end"), folder_e.insert(0, p)) if p else None)(filedialog.askdirectory()), "ghost", 90).pack(side="left", padx=8)
    make_label(dlg, "File pattern", 12, True).pack(anchor="w", padx=24)
    pat_e = make_entry(dlg, width=200)
    pat_e.pack(anchor="w", padx=24, pady=(4, 10))
    pat_e.insert(0, "*.xlsx")
    make_label(dlg, "Run", 12, True).pack(anchor="w", padx=24)
    lookup = dict(targets)
    tv = tk.StringVar(value=targets[0][0])
    ctk.CTkOptionMenu(dlg, values=[t[0] for t in targets], variable=tv, width=380, fg_color=COLOR_BG_MAIN,
                      button_color=COLOR_BORDER).pack(anchor="w", padx=24, pady=(4, 6))
    err = make_label(dlg, "", 11)
    err.configure(text_color=COLOR_DANGER_HOVER)
    err.pack(anchor="w", padx=24)

    def save():
        folder = folder_e.get().strip()
        if not folder or not os.path.isdir(folder):
            err.configure(text="Pick an existing folder.")
            return
        dlg.finish({"folder": folder, "pattern": pat_e.get().strip() or "*", "target": lookup[tv.get()]})

    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(side="bottom", fill="x", padx=24, pady=16)
    make_button(row, "Cancel", dlg.close, "ghost", 100).pack(side="right", padx=(8, 0))
    make_button(row, "Start watching", save, "success", 150).pack(side="right")
    parent.wait_window(dlg)
    return dlg.result


# ---- Excel / CSV output preview --------------------------------------------
class PreviewWindow(Modal):
    def __init__(self, parent, paths: list):
        super().__init__(parent, "Output preview", 1100, 660, modal=False)
        self.paths = [p for p in paths if os.path.exists(p)]
        self.df_cache = None
        top = ctk.CTkFrame(self, fg_color="transparent")
        top.pack(fill="x", padx=16, pady=(14, 6))
        make_label(top, "📊 Output preview", 16, True).pack(side="left")
        self.file_var = tk.StringVar(value=Path(self.paths[0]).name if self.paths else "")
        self.sheet_var = tk.StringVar(value="")
        self.file_menu = ctk.CTkOptionMenu(top, values=[Path(p).name for p in self.paths] or ["—"],
                                           variable=self.file_var, command=lambda _v: self.load(), fg_color=COLOR_SIDEBAR,
                                           button_color=COLOR_BORDER, width=260)
        self.file_menu.pack(side="left", padx=14)
        self.sheet_menu = ctk.CTkOptionMenu(top, values=["—"], variable=self.sheet_var, command=lambda v: self.load(v),
                                            fg_color=COLOR_SIDEBAR, button_color=COLOR_BORDER, width=160)
        self.sheet_menu.pack(side="left")
        make_button(top, "Open file", lambda: open_path(self.current_path()), "ghost", 90).pack(side="right")
        make_button(top, "Open folder", lambda: open_path(Path(self.current_path()).parent), "ghost", 100).pack(side="right", padx=6)
        make_button(top, "Copy path", self.copy_path, "ghost", 90).pack(side="right")
        self.info = make_label(self, "", 11, muted=True)
        self.info.pack(anchor="w", padx=18)
        holder = ctk.CTkFrame(self, fg_color=COLOR_BG_MAIN, border_width=1, border_color=COLOR_BORDER)
        holder.pack(fill="both", expand=True, padx=16, pady=(6, 14))
        holder.grid_rowconfigure(0, weight=1)
        holder.grid_columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(holder, show="headings")
        ys = ttk.Scrollbar(holder, orient="vertical", command=self.tree.yview)
        xs = ttk.Scrollbar(holder, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ys.grid(row=0, column=1, sticky="ns")
        xs.grid(row=1, column=0, sticky="ew")
        self.after(200, self.load)

    def current_path(self) -> str:
        for p in self.paths:
            if Path(p).name == self.file_var.get():
                return p
        return self.paths[0] if self.paths else ""

    def copy_path(self):
        self.clipboard_clear()
        self.clipboard_append(self.current_path())

    def load(self, sheet=None):
        pd = try_import("pandas")
        path = self.current_path()
        if not path:
            self.info.configure(text="No output files to preview.")
            return
        if pd is None:
            self.info.configure(text="pandas is not installed - install it from Env Health Check.")
            return
        try:
            if path.lower().endswith(".csv"):
                df = pd.read_csv(path, nrows=1000, encoding_errors="replace")
                self.sheet_menu.configure(values=["—"])
                self.sheet_var.set("—")
            else:
                xl = pd.ExcelFile(path)
                self.sheet_menu.configure(values=xl.sheet_names)
                if not sheet or sheet not in xl.sheet_names:
                    sheet = xl.sheet_names[0]
                self.sheet_var.set(sheet)
                df = xl.parse(sheet, nrows=1000)
        except Exception as exc:  # noqa: BLE001
            self.info.configure(text=f"Could not read file: {exc}")
            return
        self.tree.delete(*self.tree.get_children())
        cols = [str(c) for c in df.columns]
        self.tree.configure(columns=cols)
        for c in cols:
            self.tree.heading(c, text=c)
            self.tree.column(c, width=max(80, min(260, len(c) * 11 + 24)), anchor="w")
        for row in df.head(500).itertuples(index=False):
            self.tree.insert("", "end", values=["" if v != v else v for v in row])
        self.info.configure(text=f"{Path(path).name}  ·  showing {min(len(df), 500)} of {len(df)}+ rows  ·  {len(cols)} columns")


# ---- version history ---------------------------------------------------------
def list_versions(filename: str) -> list:
    folder = VERSIONS_DIR / filename
    return sorted(folder.glob("*.txt"), reverse=True) if folder.exists() else []


def save_version(filename: str, content: str, keep: int = 30) -> None:
    folder = VERSIONS_DIR / filename
    folder.mkdir(parents=True, exist_ok=True)
    (folder / f"{datetime.now():%Y%m%d_%H%M%S}.txt").write_text(content, encoding="utf-8")
    for old in list_versions(filename)[keep:]:
        try:
            old.unlink()
        except OSError:
            pass


def show_versions(parent, filename: str, current: str, on_restore) -> None:
    versions = list_versions(filename)
    dlg = Modal(parent, f"Version history · {filename}", 1000, 600, modal=False)
    make_label(dlg, f"🕘 Version history · {filename}", 16, True).pack(anchor="w", padx=20, pady=(16, 4))
    body = ctk.CTkFrame(dlg, fg_color="transparent")
    body.pack(fill="both", expand=True, padx=20, pady=8)
    body.grid_columnconfigure(1, weight=1)
    body.grid_rowconfigure(0, weight=1)
    left = ctk.CTkScrollableFrame(body, width=220, fg_color=COLOR_CARD_SURFACE, border_width=1, border_color=COLOR_BORDER)
    left.grid(row=0, column=0, sticky="ns", padx=(0, 12))
    box = ctk.CTkTextbox(body, fg_color=COLOR_CONSOLE_BG, font=(FONT_MONO, 12), border_width=1,
                         border_color=COLOR_BORDER, wrap="none")
    box.grid(row=0, column=1, sticky="nsew")
    state = {"text": None}
    info = make_label(dlg, "Pick a version to compare it with the editor.", 11, muted=True)
    info.pack(anchor="w", padx=20)

    def pick(path: Path):
        text = path.read_text(encoding="utf-8", errors="replace")
        state["text"] = text
        n = render_diff(box, text, current, "that version", "editor now")
        info.configure(text=f"{path.stem}: {n} changed lines vs. the editor")

    if not versions:
        make_label(left, "No saved versions yet.\nA snapshot is stored\nevery time you overwrite\nthis script.", 12, muted=True,
                   justify="left").pack(padx=8, pady=20)
    for v in versions:
        try:
            label = datetime.strptime(v.stem, "%Y%m%d_%H%M%S").strftime("%d %b %Y  %H:%M:%S")
        except ValueError:
            label = v.stem
        make_button(left, label, lambda v=v: pick(v), "flat", 200).pack(pady=2)
    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(fill="x", padx=20, pady=(4, 14))
    make_button(row, "Close", dlg.close, "ghost", 100).pack(side="right")

    def restore():
        if state["text"] is not None:
            on_restore(state["text"])
            dlg.close()

    make_button(row, "Restore into editor", restore, "primary", 170).pack(side="right", padx=8)


# ---- Claude API assistant -----------------------------------------------------
_SESSION_KEY = {"key": ""}
AI_KEYRING_NAME = "__anthropic_api_key__"


def get_api_key() -> str:
    if _SESSION_KEY["key"]:
        return _SESSION_KEY["key"]
    if os.environ.get("ANTHROPIC_API_KEY"):
        return os.environ["ANTHROPIC_API_KEY"]
    kr = try_import("keyring")
    if kr:
        try:
            return kr.get_password(KEYRING_SERVICE, AI_KEYRING_NAME) or ""
        except Exception:  # noqa: BLE001
            return ""
    return ""


def set_api_key(key: str, persist: bool) -> None:
    _SESSION_KEY["key"] = key
    if persist:
        kr = try_import("keyring")
        if kr:
            kr.set_password(KEYRING_SERVICE, AI_KEYRING_NAME, key)


def call_claude(prompt: str, system: str, timeout: int = 120) -> str:
    key = get_api_key()
    if not key:
        raise RuntimeError("No API key yet. Add one in Settings ▸ AI assistant.")
    body = json.dumps({"model": CFG["ai_model"], "max_tokens": 4096, "system": system,
                       "messages": [{"role": "user", "content": prompt}]}).encode("utf-8")
    req = urllib.request.Request("https://api.anthropic.com/v1/messages", data=body, method="POST", headers={
        "content-type": "application/json", "x-api-key": key, "anthropic-version": "2023-06-01"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", "replace")[:400]
        raise RuntimeError(f"API error {exc.code}: {detail}") from None
    return "".join(b.get("text", "") for b in data.get("content", []) if b.get("type") == "text")


AI_ACTIONS = {
    "Explain this code": "Explain clearly what this script does, step by step, and point out risks.",
    "Find & fix bugs": "Find bugs and fix them. Return the COMPLETE corrected script in one fenced code block, "
                       "then a short bullet list of what you changed.",
    "Generate from description": "Write a complete script that does what my instruction says. Return it in one fenced code block.",
    "Add comments & docstrings": "Add helpful comments and docstrings without changing behaviour. Return the full script in one fenced code block.",
    "Optimize (pandas / speed)": "Optimize this for speed and readability (vectorized pandas, no needless loops). "
                                 "Return the full script in one fenced code block and explain the gains briefly.",
}


def extract_code_block(text: str) -> str | None:
    m = re.search(r"```(?:python|py|vba|vb|basic)?\s*\n(.*?)```", text, re.S | re.I)
    return m.group(1).rstrip() + "\n" if m else None


def show_ai_assistant(parent, mode: str, get_code, get_last_error, on_replace, on_insert) -> None:
    dlg = Modal(parent, "AI Assist", 900, 680, modal=False)
    make_label(dlg, "✨ AI Assist (Claude)", 16, True).pack(anchor="w", padx=20, pady=(16, 0))
    make_label(dlg, "Your code is sent to Anthropic's API when you press Ask. Avoid pasting secrets or client data.",
               11, muted=True).pack(anchor="w", padx=20, pady=(0, 8))
    top = ctk.CTkFrame(dlg, fg_color="transparent")
    top.pack(fill="x", padx=20)
    action = tk.StringVar(value=list(AI_ACTIONS)[0])
    ctk.CTkOptionMenu(top, values=list(AI_ACTIONS), variable=action, width=240, fg_color=COLOR_SIDEBAR,
                      button_color=COLOR_BORDER).pack(side="left")
    inc_err = tk.BooleanVar(value=True)
    ctk.CTkCheckBox(top, text="Include last run error", variable=inc_err, checkbox_width=18, checkbox_height=18
                    ).pack(side="left", padx=16)
    instr = make_entry(dlg, "Extra instruction (required for 'Generate from description')", width=800)
    instr.pack(fill="x", padx=20, pady=10)
    out = ctk.CTkTextbox(dlg, fg_color=COLOR_CONSOLE_BG, font=(FONT_MONO, 12), border_width=1,
                         border_color=COLOR_BORDER, wrap="word")
    out.pack(fill="both", expand=True, padx=20)
    status = make_label(dlg, "", 11, muted=True)
    status.pack(anchor="w", padx=20, pady=4)
    result = {"text": ""}

    def set_out(text):
        result["text"] = text
        out.configure(state="normal")
        out.delete("1.0", "end")
        out.insert("1.0", text)

    def ask():
        act = action.get()
        code = get_code()
        if act != "Generate from description" and not code.strip():
            status.configure(text="The editor is empty.")
            return
        if act == "Generate from description" and not instr.get().strip():
            status.configure(text="Type what the script should do first.")
            return
        prompt = AI_ACTIONS[act] + "\n\n"
        if instr.get().strip():
            prompt += f"My instruction: {instr.get().strip()}\n\n"
        if code.strip() and act != "Generate from description":
            prompt += f"Language: {mode}\n\n```\n{code}\n```\n"
        err = get_last_error() if inc_err.get() else ""
        if err:
            prompt += f"\nLast run error output:\n```\n{err[-3000:]}\n```\n"
        system = ("You are a senior automation engineer helping a finance/Excel developer. Prefer pandas and openpyxl for "
                  f"Python. The language is {mode}. Be concise and correct.")
        status.configure(text="Thinking…")
        ask_btn.configure(state="disabled")

        def done(kind, val):
            ask_btn.configure(state="normal")
            if kind == "ok":
                set_out(val)
                status.configure(text="Done. Use the buttons below to apply the code block.")
            else:
                status.configure(text=f"Failed: {val}")

        run_bg(dlg, lambda: call_claude(prompt, system), done)

    row = ctk.CTkFrame(dlg, fg_color="transparent")
    row.pack(fill="x", padx=20, pady=(4, 14))
    ask_btn = make_button(row, "✨ Ask", ask, "primary", 100)
    ask_btn.pack(side="left")

    def apply(kind):
        code = extract_code_block(result["text"])
        if code is None:
            status.configure(text="No code block found in the answer.")
            return
        (on_replace if kind == "replace" else on_insert)(code)
        status.configure(text="Applied to the editor.")

    make_button(row, "Replace editor", lambda: apply("replace"), "outline", 130).pack(side="left", padx=8)
    make_button(row, "Insert at cursor", lambda: apply("insert"), "ghost", 130).pack(side="left")
    make_button(row, "Close", dlg.close, "ghost", 90).pack(side="right")
    make_button(row, "Copy answer", lambda: (dlg.clipboard_clear(), dlg.clipboard_append(result["text"])), "ghost", 110).pack(side="right", padx=8)


# ---- command palette ----------------------------------------------------------
def show_palette(parent, actions: list) -> None:
    """actions: [(label, callable), ...]"""
    dlg = Modal(parent, "Command palette", 640, 440)
    entry = make_entry(dlg, "Type a command or script name…", width=600)
    entry.pack(fill="x", padx=16, pady=(16, 8))
    lb = tk.Listbox(dlg, bg=COLOR_SIDEBAR, fg=COLOR_TEXT_MAIN, selectbackground=COLOR_ACCENT,
                    selectforeground="#0D1117", highlightthickness=0, borderwidth=0, font=(FONT_FAMILY, 12),
                    activestyle="none")
    lb.pack(fill="both", expand=True, padx=16, pady=(0, 16))
    shown: list = []

    def refresh(_e=None):
        q = entry.get()
        scored = []
        for label, cb in actions:
            s = fuzzy_score(q, label)
            if s is not None:
                scored.append((s, label, cb))
        scored.sort(key=lambda t: -t[0])
        shown[:] = [(l, c) for _s, l, c in scored[:60]]
        lb.delete(0, "end")
        for l, _c in shown:
            lb.insert("end", "  " + l)
        if shown:
            lb.selection_set(0)

    def move(delta):
        if not shown:
            return "break"
        cur = lb.curselection()
        idx = max(0, min(len(shown) - 1, (cur[0] if cur else 0) + delta))
        lb.selection_clear(0, "end")
        lb.selection_set(idx)
        lb.see(idx)
        return "break"

    def run(_e=None):
        cur = lb.curselection()
        if not shown:
            return
        cb = shown[cur[0] if cur else 0][1]
        dlg.close()
        parent.after(60, cb)

    entry.bind("<KeyRelease>", lambda e: refresh() if e.keysym not in ("Up", "Down", "Return") else None)
    entry.bind("<Down>", lambda e: move(1))
    entry.bind("<Up>", lambda e: move(-1))
    entry.bind("<Return>", run)
    lb.bind("<Double-Button-1>", run)
    refresh()
    dlg.after(200, entry.focus_set)


# =============================================================================
#  CODE EDITOR  (line numbers, syntax highlighting, auto-indent, comment toggle)
# =============================================================================
import builtins as _builtins
import keyword as _keyword

PY_KEYWORDS = set(_keyword.kwlist) | {"self", "cls"}
PY_BUILTINS = set(dir(_builtins))
PY_TOKEN = re.compile(
    r"""(?P<comment>\#[^\n]*)"""
    r"""|(?P<string>(?:[rRbBfFuU]{1,2})?(?:\"\"\"[\s\S]*?\"\"\"|'''[\s\S]*?'''|"(?:\\.|[^"\\\n])*"|'(?:\\.|[^'\\\n])*'))"""
    r"""|(?P<deco>^[ \t]*@[\w.]+)"""
    r"""|(?P<number>\b\d[\d_]*(?:\.\d+)?(?:[eE][+-]?\d+)?\b)"""
    r"""|(?P<word>[A-Za-z_]\w*)""", re.M)
VBA_KEYWORDS = {w.lower() for w in (
    "Sub Function End If Then Else ElseIf For Next To Step Each In Do Loop While Wend Until Select Case Dim As Set "
    "Let Call Const Public Private Static ByVal ByRef Optional With New Nothing True False And Or Not Xor Mod Exit "
    "Option Explicit On Error Resume GoTo Return Integer Long Double String Boolean Variant Object Date Byte Single "
    "Currency Type Property Get ReDim Preserve").split()}
VBA_TOKEN = re.compile(
    r"""(?P<comment>'[^\n]*|\bRem\b[^\n]*)|(?P<string>"(?:""|[^"\n])*")|(?P<number>\b\d+(?:\.\d+)?\b)|(?P<word>[A-Za-z_]\w*)""",
    re.M | re.I)
HL_TAGS = {"kw": "#BB9AF7", "string": "#9ECE6A", "comment": "#565F89", "number": "#FF9E64",
           "builtin": "#7DCFFF", "defname": "#7AA2F7", "deco": "#E0AF68"}


class CodeEditor(ctk.CTkFrame):
    def __init__(self, master, on_change=None, on_cursor=None, **kw):
        super().__init__(master, fg_color=COLOR_EDITOR_BG, corner_radius=12, border_width=1,
                         border_color=COLOR_BORDER, **kw)
        self.mode = "python"
        self.on_change, self.on_cursor = on_change, on_cursor
        self._hl_job = self._gutter_job = None
        self.box = ctk.CTkTextbox(self, fg_color=COLOR_EDITOR_BG, text_color="#C0CAF5", font=(FONT_MONO, 14),
                                  border_width=0, corner_radius=0, wrap="none", undo=True, autoseparators=True,
                                  maxundo=-1)
        self.tw = tk_text(self.box)
        self.gutter = tk.Canvas(self, width=56, bg=blend(COLOR_EDITOR_BG, "#000000", 0.25), highlightthickness=0)
        self.gutter.pack(side="left", fill="y", padx=(1, 0), pady=1)
        self.box.pack(side="left", fill="both", expand=True, padx=(0, 1), pady=1)
        for tag, color in HL_TAGS.items():
            self.tw.tag_configure(tag, foreground=color)
        self.tw.tag_configure("curline", background="#141C38")
        self.tw.tag_configure("errline", background="#3A1420")
        self.tw.tag_lower("curline")
        self.tw.tag_raise("errline")
        self.tw.tag_raise("sel")
        try:                                   # keep the gutter in sync with scrolling
            sb = self.box._y_scrollbar
            self.tw.configure(yscrollcommand=lambda a, b: (sb.set(a, b), self._schedule_gutter()))
        except Exception:  # noqa: BLE001
            pass
        self.tw.bind("<KeyRelease>", self._after_key, add="+")
        self.tw.bind("<<Paste>>", lambda e: self.after(20, self._after_edit), add="+")
        self.tw.bind("<<Cut>>", lambda e: self.after(20, self._after_edit), add="+")
        self.tw.bind("<ButtonRelease-1>", lambda e: self._report_cursor(), add="+")
        self.tw.bind("<MouseWheel>", lambda e: self._schedule_gutter(), add="+")
        self.tw.bind("<Configure>", lambda e: self._schedule_gutter(), add="+")
        self.tw.bind("<Tab>", self._on_tab)
        self.tw.bind("<Return>", self._on_return)
        self.tw.bind("<Control-slash>", self.toggle_comment)
        self.tw.bind("<Control-a>", lambda e: (self.tw.tag_add("sel", "1.0", "end-1c"), "break")[1])

    # ---- public API ---------------------------------------------------------
    def get(self) -> str:
        return self.tw.get("1.0", "end-1c")

    def set(self, text: str) -> None:
        self.tw.delete("1.0", "end")
        self.tw.insert("1.0", text)
        try:
            self.tw.edit_reset()
        except tk.TclError:
            pass
        self.highlight()
        self._schedule_gutter()

    def clear(self) -> None:
        self.set("")

    def set_mode(self, mode: str) -> None:
        self.mode = mode.lower()
        self.highlight()

    def insert_at_cursor(self, text: str) -> None:
        self.tw.insert("insert", text)
        self._after_edit()

    def goto_line(self, n: int) -> None:
        self.tw.see(f"{n}.0")
        self.tw.mark_set("insert", f"{n}.0")
        self.tw.focus_set()
        self._schedule_gutter()

    def mark_line(self, n: int) -> None:
        self.tw.tag_add("errline", f"{n}.0", f"{n}.end+1c")

    def clear_marks(self) -> None:
        self.tw.tag_remove("errline", "1.0", "end")

    # ---- events -------------------------------------------------------------
    def _after_key(self, event=None):
        self._after_edit()
        self._report_cursor()

    def _after_edit(self):
        if self._hl_job:
            self.after_cancel(self._hl_job)
        self._hl_job = self.after(250, self.highlight)
        self._schedule_gutter()
        if self.on_change:
            self.on_change()

    def _mark_curline(self):
        try:
            self.tw.tag_remove("curline", "1.0", "end")
            self.tw.tag_add("curline", "insert linestart", "insert lineend+1c")
        except tk.TclError:
            pass

    def _report_cursor(self):
        self._mark_curline()
        self._schedule_gutter()
        if self.on_cursor:
            line, col = str(self.tw.index("insert")).split(".")
            self.on_cursor(int(line), int(col) + 1)

    def _on_tab(self, event):
        if self.tw.tag_ranges("sel"):
            self.tw.delete("sel.first", "sel.last")
        self.tw.insert("insert", "    ")
        return "break"

    def _on_return(self, event):
        if self.tw.tag_ranges("sel"):
            self.tw.delete("sel.first", "sel.last")
        line = self.tw.get("insert linestart", "insert")
        indent = re.match(r"[ \t]*", line).group()
        extra = ""
        if self.mode == "python" and line.rstrip().endswith(":"):
            extra = "    "
        elif self.mode == "vba" and re.search(r"(\bThen|\bDo|\bElse)\s*$|^\s*(Sub|Function|For|With|Select Case|While)\b", line, re.I):
            extra = "    "
        self.tw.insert("insert", "\n" + indent + extra)
        self.tw.see("insert")
        self._after_edit()
        return "break"

    def toggle_comment(self, event=None):
        mark = "#" if self.mode == "python" else "'"
        try:
            first = int(str(self.tw.index("sel.first")).split(".")[0])
            last = int(str(self.tw.index("sel.last")).split(".")[0])
        except tk.TclError:
            first = last = int(str(self.tw.index("insert")).split(".")[0])
        lines = {n: self.tw.get(f"{n}.0", f"{n}.end") for n in range(first, last + 1)}
        active = [l for l in lines.values() if l.strip()]
        all_commented = bool(active) and all(l.lstrip().startswith(mark) for l in active)
        for n, l in lines.items():
            if not l.strip():
                continue
            if all_commented:
                new = re.sub(r"^(\s*)" + re.escape(mark) + r" ?", r"\1", l, count=1)
            else:
                new = re.sub(r"^(\s*)", lambda m: m.group(1) + mark + " ", l, count=1)
            self.tw.delete(f"{n}.0", f"{n}.end")
            self.tw.insert(f"{n}.0", new)
        self._after_edit()
        return "break"

    # ---- highlighting -------------------------------------------------------
    def highlight(self):
        self._hl_job = None
        try:
            self._highlight()
        except tk.TclError:          # widget was destroyed while the timer was pending
            pass

    def _highlight(self):
        text = self.get()
        for tag in HL_TAGS:
            self.tw.tag_remove(tag, "1.0", "end")
        starts = [0] + [m.end() for m in re.finditer("\n", text)]

        def pos(off):
            line = bisect.bisect_right(starts, off) - 1
            return f"{line + 1}.{off - starts[line]}"

        py = self.mode == "python"
        ranges: dict = {t: [] for t in HL_TAGS}
        prev = ""
        for m in (PY_TOKEN if py else VBA_TOKEN).finditer(text):
            kind, tag = m.lastgroup, None
            if kind == "word":
                w = m.group()
                if py:
                    tag = "kw" if w in PY_KEYWORDS else "defname" if prev in ("def", "class") else "builtin" if w in PY_BUILTINS else None
                else:
                    tag = "kw" if w.lower() in VBA_KEYWORDS else None
                prev = w
            else:
                tag = kind
                prev = ""
            if tag:
                ranges[tag] += [pos(m.start()), pos(m.end())]
        for tag, pairs in ranges.items():
            if pairs:
                self.tw.tag_add(tag, *pairs)

    # ---- gutter -------------------------------------------------------------
    def _schedule_gutter(self):
        if self._gutter_job is None:
            self._gutter_job = self.after_idle(self._draw_gutter)

    def _draw_gutter(self):
        self._gutter_job = None
        try:
            self.gutter.delete("all")
            font = self.tw.cget("font")
            idx, guard = self.tw.index("@0,0"), 0
            cur_line = str(self.tw.index("insert")).split(".")[0]
            while guard < 400:
                info = self.tw.dlineinfo(idx)
                if info is None:
                    break
                ln = str(idx).split(".")[0]
                self.gutter.create_text(48, info[1], anchor="ne", text=ln, font=font,
                                        fill=COLOR_ACCENT2 if ln == cur_line else "#4A5680")
                nxt = self.tw.index(f"{idx}+1line")
                if nxt == idx:
                    break
                idx, guard = nxt, guard + 1
        except tk.TclError:
            pass


# =============================================================================
#  CONSOLE PANEL  (streaming output, per-script tabs, search, export)
# =============================================================================
def style_console_tags(tw) -> None:
    for tag, color in (("out", COLOR_CONSOLE_TEXT), ("stderr", "#FF9E64"), ("error", "#F7768E"),
                       ("warning", "#E0AF68"), ("success", "#9ECE6A"), ("info", COLOR_ACCENT2),
                       ("dim", "#5A6A8A")):
        tw.tag_configure(tag, foreground=color)
    tw.tag_configure("find", background=COLOR_ACCENT, foreground="#0A0E17")
    tw.tag_raise("find")


def make_console_box(parent):
    box = ctk.CTkTextbox(parent, fg_color=COLOR_CONSOLE_BG, text_color=COLOR_CONSOLE_TEXT, font=(FONT_MONO, 13),
                         border_width=1, border_color=COLOR_BORDER, corner_radius=10, wrap="word")
    style_console_tags(tk_text(box))
    box.configure(state="disabled")
    return box


def console_write(box, text: str, tag: str = "out", max_lines: int = 6000) -> None:
    tw = tk_text(box)
    box.configure(state="normal")
    tw.insert("end", text, tag)
    try:
        lines = int(tw.index("end-1c").split(".")[0])
        if lines > max_lines:
            tw.delete("1.0", f"{lines - max_lines + 1000}.0")
    except (tk.TclError, ValueError):
        pass
    tw.see("end")
    box.configure(state="disabled")


def terminal_chrome(parent, title: str):
    """macOS-style window bar (three dots + mono title) that makes a console look like a real terminal."""
    bar = ctk.CTkFrame(parent, fg_color=blend(COLOR_CONSOLE_BG, "#FFFFFF", 0.05), corner_radius=10, height=30)
    bar.pack_propagate(False)
    for c in ("#FF5F56", "#FFBD2E", "#27C93F"):
        ctk.CTkLabel(bar, text="●", text_color=c, font=ctk.CTkFont(size=13), width=16).pack(side="left", padx=(8 if c == "#FF5F56" else 0, 0))
    ctk.CTkLabel(bar, text=title, text_color="#5A6A8A", font=ctk.CTkFont(family=FONT_MONO, size=11)).pack(side="left", padx=10)
    return bar


def make_terminal(parent, title: str):
    """(wrapper, textbox) - a console box with terminal chrome on top."""
    wrap = ctk.CTkFrame(parent, fg_color="transparent")
    terminal_chrome(wrap, title).pack(fill="x", pady=(0, 4))
    box = make_console_box(wrap)
    box.pack(fill="both", expand=True)
    return wrap, box


class ConsolePanel(ctk.CTkFrame):
    ALL = "All output"

    def __init__(self, master, on_stop=None, **kw):
        super().__init__(master, fg_color="transparent", **kw)
        self.on_stop = on_stop
        self.tabs: dict = {}           # key -> (tab_name, textbox)
        self.order: list = []
        self._find_pos = "1.0"
        head = ctk.CTkFrame(self, fg_color="transparent")
        head.pack(fill="x", pady=(0, 6))
        make_label(head, "Live Output Console", 17, True, display=True).pack(side="left")
        make_button(head, "🗑 Clear", self.clear, "ghost", 80).pack(side="right")
        make_button(head, "💾 Export", self.export, "ghost", 84).pack(side="right", padx=6)
        make_button(head, "📋 Copy", self.copy, "ghost", 76).pack(side="right")
        make_button(head, "⏹ Stop", self._stop, "danger", 76).pack(side="right", padx=6)
        make_button(head, "▼", lambda: self.find(1), "ghost", 30).pack(side="right")
        make_button(head, "▲", lambda: self.find(-1), "ghost", 30).pack(side="right", padx=(6, 2))
        self.search = make_entry(head, "Search output…", width=180)
        self.search.pack(side="right", padx=4)
        self.search.bind("<Return>", lambda e: self.find(1))
        terminal_chrome(self, "finprime ~ live-output  —  streaming stdout / stderr").pack(fill="x", pady=(0, 4))
        self.tabview = ctk.CTkTabview(self, fg_color=COLOR_CONSOLE_BG, segmented_button_fg_color=COLOR_SIDEBAR,
                                      segmented_button_selected_color=COLOR_ACCENT,
                                      segmented_button_selected_hover_color=COLOR_ACCENT_HOVER,
                                      segmented_button_unselected_color=COLOR_SIDEBAR,
                                      text_color=COLOR_TEXT_MAIN, border_width=1, border_color=COLOR_BORDER, height=120)
        self.tabview.pack(fill="both", expand=True)
        tab = self.tabview.add(self.ALL)
        self.all_box = make_console_box(tab)
        self.all_box.pack(fill="both", expand=True)
        self.tabs[None] = (self.ALL, self.all_box)
        console_write(self.all_box, "[SYSTEM READY] Select a Python script or VBA macro module above…\n", "dim")

    def _tab_for(self, key: str):
        if key in self.tabs:
            return self.tabs[key][1]
        name = Path(key).stem if not key.startswith("pipe:") else "pipe · " + key[5:]
        base, n = name[:22], 2
        used = {t[0] for t in self.tabs.values()}
        while name in used:
            name, n = f"{base} ({n})", n + 1
        if len(self.order) >= 8:                       # keep the tab bar tidy
            oldest = self.order.pop(0)
            try:
                self.tabview.delete(self.tabs.pop(oldest)[0])
            except Exception:  # noqa: BLE001
                pass
        tab = self.tabview.add(name)
        box = make_console_box(tab)
        box.pack(fill="both", expand=True)
        self.tabs[key] = (name, box)
        self.order.append(key)
        return box

    def write(self, key: str, text: str, tag: str = "out") -> None:
        console_write(self.all_box, text, tag)
        if key and key != SANDBOX_KEY:
            console_write(self._tab_for(key), text, tag)

    def system(self, text: str, tag: str = "dim") -> None:
        console_write(self.all_box, text, tag)

    def focus_key(self, key: str) -> None:
        if key in self.tabs:
            try:
                self.tabview.set(self.tabs[key][0])
            except Exception:  # noqa: BLE001
                pass

    def current(self):
        name = self.tabview.get()
        for _k, (n, box) in self.tabs.items():
            if n == name:
                return box
        return self.all_box

    def current_key(self):
        name = self.tabview.get()
        for k, (n, _b) in self.tabs.items():
            if n == name:
                return k
        return None

    def text(self) -> str:
        return tk_text(self.current()).get("1.0", "end-1c")

    def clear(self) -> None:
        box = self.current()
        box.configure(state="normal")
        tk_text(box).delete("1.0", "end")
        box.configure(state="disabled")

    def copy(self) -> None:
        self.clipboard_clear()
        self.clipboard_append(self.text())

    def export(self) -> None:
        path = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Text", "*.txt"), ("Log", "*.log")],
                                            initialfile=f"console_{datetime.now():%Y%m%d_%H%M%S}.txt")
        if path:
            Path(path).write_text(self.text(), encoding="utf-8")

    def _stop(self) -> None:
        if self.on_stop:
            self.on_stop(self.current_key())

    def find(self, direction: int = 1) -> None:
        term = self.search.get()
        tw = tk_text(self.current())
        tw.tag_remove("find", "1.0", "end")
        if not term:
            return
        start = self._find_pos if direction > 0 else self._find_pos
        try:
            if direction > 0:
                idx = tw.search(term, start, stopindex="end", nocase=True) or tw.search(term, "1.0", stopindex="end", nocase=True)
            else:
                idx = tw.search(term, start, stopindex="1.0", nocase=True, backwards=True) or \
                      tw.search(term, "end", stopindex="1.0", nocase=True, backwards=True)
        except tk.TclError:
            return
        if not idx:
            return
        end = f"{idx}+{len(term)}c"
        tw.tag_add("find", idx, end)
        tw.see(idx)
        self._find_pos = end if direction > 0 else idx


# =============================================================================
#  MATRIX RAIN  (items are created once and moved - no per-frame rebuild)
# =============================================================================
class MatrixRainCanvas(tk.Canvas):
    CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    FS, TRAIL = 14, 12

    def __init__(self, master, **kw):
        super().__init__(master, bg=COLOR_BG_MAIN, highlightthickness=0, **kw)
        self.width = self.height = 0
        self.cols: list = []
        self.fade_sets: list = []
        self._job = None
        self.running = False
        self.bind("<Configure>", self._on_resize)

    def _on_resize(self, event):
        self.width, self.height = event.width, event.height
        self._build()

    def _build(self):
        self.delete("all")
        self.cols = []
        def ramp(c):
            return [blend(c, "#FFFFFF", 0.7)] + [blend(c, COLOR_BG_MAIN, min(0.93, j / self.TRAIL)) for j in range(1, self.TRAIL)]
        self.fade_sets = [ramp(COLOR_RAIN_HEAD), ramp(COLOR_ACCENT2)]
        row = self.FS + 4
        for c in range(0, max(1, self.width // self.FS), 2):
            items = [self.create_text(c * self.FS + self.FS // 2, -100, text="", anchor="n", state="hidden",
                                      font=(FONT_MONO, self.FS, "bold")) for _ in range(self.TRAIL)]
            self.cols.append({"fades": self.fade_sets[0 if random.random() < 0.7 else 1], "x": c * self.FS + self.FS // 2, "y": random.randint(-self.height, 0),
                              "speed": random.choice([1, 1, 2]), "items": items,
                              "chars": [random.choice(self.CHARS) for _ in range(self.TRAIL)]})
        self.row = row

    def start_animation(self):
        if not self.running:
            self.running = True
            self._animate()

    def stop_animation(self):
        self.running = False
        if self._job:
            try:
                self.after_cancel(self._job)
            except tk.TclError:
                pass
            self._job = None

    def _animate(self):
        if not self.running:
            return
        if self.cols:
            for col in self.cols:
                col["y"] += col["speed"] * (self.FS - 2)
                col["chars"].insert(0, random.choice(self.CHARS))
                col["chars"].pop()
                for j, item in enumerate(col["items"]):
                    y = col["y"] - j * self.row
                    if -self.row < y < self.height + self.row:
                        self.coords(item, col["x"], y)
                        self.itemconfigure(item, text=col["chars"][j], fill=col["fades"][j], state="normal")
                    else:
                        self.itemconfigure(item, state="hidden")
                if col["y"] - self.TRAIL * self.row > self.height and random.random() > 0.9:
                    col["y"] = random.randint(-self.height // 2, 0)
                    col["speed"] = random.choice([1, 1, 2])
        self._job = self.after(70, self._animate)


# =============================================================================
#  LOGIN / LOCK OVERLAY  (throttled attempts, first-launch confirm field)
# =============================================================================
class LoginOverlay(ctk.CTkFrame):
    def __init__(self, parent, on_success):
        super().__init__(parent, fg_color=COLOR_BG_MAIN, corner_radius=0)
        self.on_success = on_success
        self.fails = 0
        self.locked_until = 0.0
        self.first_launch = load_auth() is None
        self.rain = None
        self._cursor_on = True
        if CFG["matrix_rain"]:
            self.rain = MatrixRainCanvas(self)
            self.rain.place(x=0, y=0, relwidth=1, relheight=1)
            self.rain.start_animation()
        self.card = ctk.CTkFrame(self, fg_color=COLOR_SIDEBAR, corner_radius=20, border_width=2,
                                 border_color=blend(COLOR_BORDER, COLOR_ACCENT2, 0.55), width=860, height=660)
        self.card.place(relx=0.5, rely=0.5, anchor="center")
        self.card.pack_propagate(False)
        self.card.grid_propagate(False)
        self.card.grid_columnconfigure(0, weight=1)
        self.card.grid_columnconfigure(1, weight=1)
        self.card.grid_rowconfigure(0, weight=1)
        self.left = ctk.CTkFrame(self.card, fg_color="#080C16", corner_radius=14)
        self.left.grid(row=0, column=0, sticky="nsew", padx=12, pady=12)
        self._load_hero_image()
        right = ctk.CTkFrame(self.card, fg_color="transparent")
        right.grid(row=0, column=1, sticky="nsew", padx=40, pady=44)
        head = ctk.CTkFrame(right, fg_color="transparent")
        head.pack(anchor="w", fill="x")
        ctk.CTkLabel(head, text="", image=icon_tile("bolt", 44, (COLOR_ACCENT, COLOR_ACCENT2))).pack(side="left")
        wm = gradient_wordmark("FINPRIME'S ENGINE", 19, COLOR_ACCENT, COLOR_ACCENT2, spacing=1)
        if wm is not None:
            ctk.CTkLabel(head, text="", image=wm).pack(side="left", padx=12)
        else:
            make_label(head, "Finprime's Engine", 20, True, text_color=COLOR_ACCENT).pack(side="left", padx=12)
        self.terminal_line = ctk.CTkLabel(right, text="> secure_shell --unlock_", text_color=COLOR_ACCENT2,
                                          font=ctk.CTkFont(family=FONT_MONO, size=12))
        self.terminal_line.pack(anchor="w", pady=(26, 0))
        self.lbl_title = make_label(right, "Welcome back", 30, True)
        self.lbl_title.pack(anchor="w", pady=(2, 0))
        self.lbl_sub = make_label(right, "Sign in to unlock your automation hub", 12, muted=True)
        self.lbl_sub.pack(anchor="w", pady=(0, 24))
        make_label(right, "🔒 Master Password", 12, True).pack(anchor="w", pady=(0, 6))
        self.entry = make_entry(right, "Enter your master password", width=300, show="•", height=46)
        self.entry.pack(fill="x", pady=(0, 10))
        self.entry.bind("<Return>", lambda e: self.authenticate())
        self.confirm_lbl = make_label(right, "Confirm password", 12, True)
        self.confirm = make_entry(right, "Type it again", width=300, show="•", height=46)
        self.confirm.bind("<Return>", lambda e: self.authenticate())
        self.status = make_label(right, "", 11)
        self.status.configure(text_color="#FB7185")
        self.btn = make_button(right, "Unlock", self.authenticate, "primary", height=46, glyph="bolt",
                               font=ctk.CTkFont(family=FONT_DISPLAY, size=15, weight="bold"))
        if self.first_launch:
            self.lbl_title.configure(text="Create your vault")
            self.lbl_sub.configure(text="First launch: set a master password (6+ characters)")
            self.confirm_lbl.pack(anchor="w", pady=(4, 6))
            self.confirm.pack(fill="x", pady=(0, 10))
            self.btn.configure(text="Set master password", fg_color=COLOR_SUCCESS, hover_color=COLOR_SUCCESS_HOVER,
                               text_color="#FFFFFF", image=icon("bolt", 16, "#FFFFFF"))
        self.status.pack(anchor="w", pady=4)
        self.btn.pack(fill="x", pady=(14, 0))
        make_label(right, f"v{APP_VERSION}  ·  press Enter to unlock", 10, muted=True).pack(side="bottom", anchor="w")
        self.after(300, self._focus_entry)
        self.after(500, self._blink)

    def _blink(self):
        try:
            self._cursor_on = not self._cursor_on
            self.terminal_line.configure(text="> secure_shell --unlock" + ("_" if self._cursor_on else " "))
            self.after(520, self._blink)
        except tk.TclError:
            pass

    def _focus_entry(self):
        try:
            self.entry.focus_set()
        except tk.TclError:
            pass

    def _load_hero_image(self):
        target = None
        for name in ("login_bg.png", "login_bg.jpg", "login_bg.jpeg", "login_bg.webp", "login_bg"):
            p = workspace_dir() / name
            if p.exists():
                target = p
                break
        if target:
            try:
                img = Image.open(target).convert("RGBA")
                photo = ctk.CTkImage(light_image=img, dark_image=img, size=(390, 636))
                ctk.CTkLabel(self.left, image=photo, text="").pack(fill="both", expand=True)
                return
            except Exception as exc:  # noqa: BLE001
                print(f"[IMAGE DISPLAY ERROR]: {exc}")
        ctk.CTkLabel(self.left, text="", image=icon_tile("bolt", 120, (COLOR_ACCENT, COLOR_ACCENT2))).pack(expand=True, pady=(120, 0))
        make_label(self.left, "Enterprise Automation Engine", 15, True, display=True).pack()
        make_label(self.left, "Python  ·  VBA  ·  Pipelines  ·  Schedules", 11, muted=True).pack(pady=(4, 0))
        make_label(self.left, f"Secure workspace  v{APP_VERSION}", 11, muted=True).pack(pady=(0, 36), side="bottom")

    def stop(self):
        if self.rain:
            self.rain.stop_animation()

    def _tick_lock(self):
        remaining = int(self.locked_until - time.monotonic())
        if remaining > 0:
            self.status.configure(text=f"Too many attempts. Try again in {remaining}s.")
            self.btn.configure(state="disabled")
            self.after(500, self._tick_lock)
        else:
            self.status.configure(text="")
            self.btn.configure(state="normal")

    def authenticate(self):
        if time.monotonic() < self.locked_until:
            return
        pw = self.entry.get()
        if not pw:
            self.status.configure(text="Password cannot be empty.")
            return
        if self.first_launch:
            if len(pw) < 6:
                self.status.configure(text="Use at least 6 characters.")
                return
            if pw != self.confirm.get():
                self.status.configure(text="Passwords do not match.")
                return
            save_auth(pw)
            self.stop()
            self.on_success()
            return
        if check_password(pw):
            self.stop()
            self.on_success()
            return
        self.fails += 1
        self.entry.delete(0, "end")
        if self.fails >= 3:
            self.locked_until = time.monotonic() + min(300, 5 * 2 ** (self.fails - 3))
            self._tick_lock()
        else:
            self.status.configure(text=f"Access Denied: Incorrect Password ({3 - self.fails} tries left)")


# =============================================================================
#  TOASTS
# =============================================================================
class ToastManager:
    ICONS = {"info": "info", "success": "check", "error": "close", "warning": "warn"}

    def __init__(self, root):
        self.root = root
        self.frame = None
        self._job = None

    @staticmethod
    def palette(kind: str):
        return {"success": ("#4ADE80", "#16A34A"), "error": ("#FB7185", "#DC2626"),
                "warning": ("#FBBF24", "#F97316"), "info": (COLOR_ACCENT, COLOR_ACCENT2)}.get(kind, (COLOR_ACCENT, COLOR_ACCENT2))

    def show(self, message: str, kind: str = "info", ms: int = 3400) -> None:
        try:
            if self.frame is not None:
                self.frame.destroy()
            c1, c2 = self.palette(kind)
            self.frame = ctk.CTkFrame(self.root, fg_color="#0F1626", border_width=1,
                                      border_color=blend(COLOR_BORDER, c1, 0.6), corner_radius=14)
            ctk.CTkLabel(self.frame, text="", image=icon_tile(self.ICONS.get(kind, "info"), 30, (c1, c2))
                         ).pack(side="left", padx=(12, 4), pady=12)
            make_label(self.frame, message, 13, wraplength=340, justify="left").pack(side="left", padx=(6, 18), pady=12)
            self.frame.place(relx=1.0, x=60, y=78, anchor="ne")
            self.frame.lift()
            self._slide(60)
            if self._job:
                self.root.after_cancel(self._job)
            self._job = self.root.after(ms, self.hide)
        except tk.TclError:
            pass

    def _slide(self, x: int):
        try:
            if self.frame is None:
                return
            self.frame.place_configure(x=max(x, -24))
            if x > -24:
                self.root.after(14, lambda: self._slide(x - 9))
        except tk.TclError:
            pass

    def hide(self):
        try:
            if self.frame is not None:
                self.frame.destroy()
        except tk.TclError:
            pass
        self.frame = None


# =============================================================================
#  SNIPPET LIBRARY
# =============================================================================
SNIPPETS = {
    "python": {
        "Read Excel → DataFrame": '''import pandas as pd
from pathlib import Path

SRC = Path(r"C:\\path\\to\\input.xlsx")
df = pd.read_excel(SRC, sheet_name=0)
print(df.head())
print(f"Loaded {len(df):,} rows x {len(df.columns)} columns")
''',
        "Formatted Excel report (openpyxl)": '''import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rows = [("Region", "Revenue", "Profit"), ("North", 120000, 32000), ("South", 98000, 21000)]
wb = Workbook()
ws = wb.active
ws.title = "Summary"
for r in rows:
    ws.append(r)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center")
for col in range(2, 4):
    for c in ws[get_column_letter(col)][1:]:
        c.number_format = "#,##0"
for i in range(1, ws.max_column + 1):
    ws.column_dimensions[get_column_letter(i)].width = 16
ws.freeze_panes = "A2"
out = os.path.join(os.environ.get("FP_OUTPUT_DIR", "."), "summary_report.xlsx")
wb.save(out)
print(f"✅ Saved {out}")
''',
        "Monthly summary (groupby → pivot)": '''import pandas as pd

df = pd.read_excel(r"C:\\path\\to\\data.xlsx")
df["Date"] = pd.to_datetime(df["Date"])
df["Month"] = df["Date"].dt.to_period("M").astype(str)
summary = df.pivot_table(index="Month", columns="Category", values="Amount", aggfunc="sum", fill_value=0)
summary["Total"] = summary.sum(axis=1)
print(summary.round(2))
''',
        "Parameters (@param) template": '''# @param input_file: file -- Excel workbook to process
# @param report_month: str = January -- Month label for the report
# @param dry_run: bool = false -- Print only, do not write files
import os

input_file = os.environ["FP_PARAM_INPUT_FILE"]
report_month = os.environ.get("FP_PARAM_REPORT_MONTH", "January")
dry_run = os.environ.get("FP_PARAM_DRY_RUN", "false") == "true"
print(f"Processing {input_file} for {report_month} (dry run: {dry_run})")
''',
        "argparse template": '''import argparse

parser = argparse.ArgumentParser(description="Describe what this script does")
parser.add_argument("--input-file", help="Excel workbook to process")
parser.add_argument("--month", default="January", help="Month label")
parser.add_argument("--verbose", action="store_true", help="Chatty output")
args = parser.parse_args()
print(args)
''',
        "Logging + error handling": '''import logging
import sys

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", stream=sys.stdout)
log = logging.getLogger("job")


def main():
    log.info("Starting")
    # ... your work here ...
    log.info("✅ Done")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        log.exception("Job failed")
        sys.exit(1)
''',
    },
    "vba": {
        "Last used row helper": '''Option Explicit

Function LastRow(ws As Worksheet, Optional col As Long = 1) As Long
    LastRow = ws.Cells(ws.Rows.Count, col).End(xlUp).Row
End Function
''',
        "Loop over rows": '''Option Explicit

Sub ProcessRows()
    Dim ws As Worksheet, r As Long, lastR As Long
    Set ws = ActiveSheet
    lastR = ws.Cells(ws.Rows.Count, 1).End(xlUp).Row
    For r = 2 To lastR
        ' ws.Cells(r, 3).Value = ws.Cells(r, 1).Value * ws.Cells(r, 2).Value
    Next r
End Sub
''',
        "Fast mode wrapper": '''Option Explicit

Sub FastRun()
    Application.ScreenUpdating = False
    Application.Calculation = xlCalculationManual
    On Error GoTo CleanUp
    ' ... your work here ...
CleanUp:
    Application.Calculation = xlCalculationAutomatic
    Application.ScreenUpdating = True
    If Err.Number <> 0 Then MsgBox "Error: " & Err.Description
End Sub
''',
        "Export sheet to CSV": '''Option Explicit

Sub ExportSheetCsv()
    Dim p As String
    p = Environ("USERPROFILE") & "\\Documents\\export.csv"
    ActiveSheet.Copy
    ActiveWorkbook.SaveAs Filename:=p, FileFormat:=xlCSV
    ActiveWorkbook.Close SaveChanges:=False
End Sub
''',
    },
}


# =============================================================================
#  SIGNATURE WIDGETS  (sidebar nav item, hero banner, key caps, pills)
# =============================================================================
def bind_tree(widget, sequence, func):
    """Bind an event on a widget and every descendant (CTk composites are made of several widgets)."""
    widget.bind(sequence, func, add="+")
    for child in widget.winfo_children():
        bind_tree(child, sequence, func)


class NavItem(ctk.CTkFrame):
    """Sidebar row: colour icon tile + label + glowing accent bar when active."""

    def __init__(self, master, icon_name: str, text: str, command):
        super().__init__(master, fg_color="transparent", corner_radius=12, height=46)
        self.command = command
        self.active = False
        self.pack_propagate(False)
        self.bar = tk.Frame(self, width=4, bg=COLOR_SIDEBAR)
        self.bar.pack(side="left", fill="y", padx=(2, 10), pady=11)
        self.icon_lbl = ctk.CTkLabel(self, text="", image=icon_tile(icon_name, 30))
        self.icon_lbl.pack(side="left")
        self.text_lbl = ctk.CTkLabel(self, text=text, anchor="w", text_color=COLOR_TEXT_MUTED,
                                     font=ctk.CTkFont(family=FONT_FAMILY, size=14, weight="bold"))
        self.text_lbl.pack(side="left", padx=12, fill="x", expand=True)
        for w in (self, self.bar, self.icon_lbl, self.text_lbl):
            w.bind("<Button-1>", lambda e: self.command(), add="+")
            w.bind("<Enter>", lambda e: self._hover(True), add="+")
            w.bind("<Leave>", lambda e: self._hover(False), add="+")
            try:
                w.configure(cursor="hand2")
            except tk.TclError:
                pass

    def _paint(self, bg: str):
        self.configure(fg_color=bg)
        self.bar.configure(bg=COLOR_ACCENT if self.active else (bg if bg != "transparent" else COLOR_SIDEBAR))

    def _hover(self, on: bool):
        if self.active:
            return
        self._paint(COLOR_CARD_HOVER if on else "transparent")
        self.text_lbl.configure(text_color=COLOR_TEXT_MAIN if on else COLOR_TEXT_MUTED)

    def set_active(self, active: bool):
        self.active = active
        self._paint(blend(COLOR_SIDEBAR, COLOR_ACCENT, 0.16) if active else "transparent")
        self.text_lbl.configure(text_color=COLOR_TEXT_MAIN if active else COLOR_TEXT_MUTED)


def key_cap(parent, text: str):
    return ctk.CTkLabel(parent, text=text, text_color=COLOR_TEXT_MUTED, fg_color=blend(COLOR_BG_MAIN, "#FFFFFF", 0.06),
                        corner_radius=6, padx=7, pady=1, font=ctk.CTkFont(family=FONT_MONO, size=11, weight="bold"))


def pill(parent, glyph_name: str, text: str, color: str | None = None):
    f = ctk.CTkFrame(parent, fg_color=blend(COLOR_BG_MAIN, "#FFFFFF", 0.05), corner_radius=16)
    ctk.CTkLabel(f, text="", image=icon(glyph_name, 15, color), width=15).pack(side="left", padx=(12, 4), pady=6)
    lbl = make_label(f, text, 12, True)
    lbl.pack(side="left", padx=(0, 14))
    return f, lbl


class HeroBanner(tk.Canvas):
    """Gradient banner with a faint code watermark; re-renders when resized."""

    def __init__(self, master, height: int = 150):
        super().__init__(master, height=height, highlightthickness=0, bd=0, bg=COLOR_BG_MAIN)
        self._photo = None
        self._job = None
        self._size = (0, 0)
        self.title = ""
        self.sub = ""
        self.kicker = ""
        self.buttons: list = []
        self.bind("<Configure>", self._on_cfg)

    def set_text(self, kicker: str, title: str, sub: str):
        self.kicker, self.title, self.sub = kicker, title, sub
        self._draw_text()

    def add_button(self, btn):
        self.buttons.append(btn)
        self._draw_text()

    def _on_cfg(self, e):
        if self._job:
            self.after_cancel(self._job)
        self._job = self.after(90, self._render)

    def _render(self):
        self._job = None
        w, h = self.winfo_width(), self.winfo_height()
        if w < 60 or h < 40:
            return
        self._size = (w, h)
        try:
            self._photo = ImageTk.PhotoImage(self._make_bg(w, h), master=self)
        except Exception as exc:  # noqa: BLE001
            print(f"[hero] {exc}")
            return
        self._draw_text()

    def _make_bg(self, w: int, h: int) -> Image.Image:
        c1 = blend(COLOR_BG_MAIN, COLOR_ACCENT, 0.42)
        c2 = blend(COLOR_BG_MAIN, COLOR_ACCENT2, 0.30)
        mask = Image.linear_gradient("L").transpose(Image.Transpose.ROTATE_90).resize((w, h))
        grad = Image.composite(Image.new("RGB", (w, h), c2), Image.new("RGB", (w, h), c1), mask).convert("RGBA")
        glow = Image.new("L", (w, h), 0)
        gd = ImageDraw.Draw(glow)
        gd.ellipse((w * 0.62, -h * 0.9, w * 1.05, h * 0.9), fill=200)
        gd.ellipse((-w * 0.08, h * 0.5, w * 0.25, h * 1.6), fill=110)
        glow = glow.filter(ImageFilter.GaussianBlur(45))
        grad.paste(Image.new("RGBA", (w, h), hex_rgb(COLOR_ACCENT2) + (255,)), (0, 0), glow.point(lambda v: int(v * 0.28)))
        grid = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        dr = ImageDraw.Draw(grid)
        for x in range(0, w, 34):
            dr.line((x, 0, x, h), fill=(255, 255, 255, 10))
        for y in range(0, h, 34):
            dr.line((0, y, w, y), fill=(255, 255, 255, 10))
        grad = Image.alpha_composite(grad, grid)
        wm = render_glyph("code", h, "#FFFFFF")
        wm.putalpha(wm.getchannel("A").point(lambda v: int(v * 0.10)))
        wm = wm.resize((int(h * 1.15), int(h * 1.15)), Image.LANCZOS)
        grad.alpha_composite(wm, (max(0, w - wm.width - 210), (h - wm.height) // 2))
        out = Image.new("RGB", (w, h), COLOR_BG_MAIN)
        rr = Image.new("L", (w * 3, h * 3), 0)
        ImageDraw.Draw(rr).rounded_rectangle((0, 0, w * 3 - 1, h * 3 - 1), radius=54, fill=255)
        out.paste(grad.convert("RGB"), (0, 0), rr.resize((w, h), Image.LANCZOS))
        return out

    def _draw_text(self):
        self.delete("all")
        w, h = self._size
        if self._photo is not None:
            self.create_image(0, 0, image=self._photo, anchor="nw")
        x = 34
        self.create_text(x, h * 0.26, text=self.kicker.upper(), anchor="w", fill=blend(COLOR_ACCENT2, "#FFFFFF", 0.35),
                         font=(FONT_MONO, 10, "bold"))
        self.create_text(x, h * 0.50, text=self.title, anchor="w", fill="#FFFFFF", font=(FONT_DISPLAY, 26, "bold"))
        self.create_text(x, h * 0.76, text=self.sub, anchor="w", fill=blend("#FFFFFF", COLOR_ACCENT2, 0.25),
                         font=(FONT_FAMILY, 11))
        right = w - 26
        for btn in reversed(self.buttons):
            self.create_window(right, h * 0.5, window=btn, anchor="e")
            right -= btn.winfo_reqwidth() + 10


def chart_image(w: int, h: int, data: list, bg: str):
    """Anti-aliased stacked bar chart (ok / failed) as a PIL image + layout info for canvas labels."""
    S = 3
    W, H = w * S, h * S
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)
    pad_l, pad_r, pad_t, pad_b = 44 * S, 16 * S, 34 * S, 30 * S
    mx = max([ok + bad for _d, ok, bad in data] + [1])
    top = max(4, int(math.ceil(mx / 4.0)) * 4)
    base = H - pad_b
    plot_h = base - pad_t
    grid_ys = []
    for i in range(5):
        y = base - plot_h * i / 4
        grid_ys.append((y / S, round(top * i / 4)))
        x = pad_l
        while x < W - pad_r:
            d.line((x, y, x + 6 * S, y), fill=blend(bg, "#FFFFFF", 0.11), width=S)
            x += 12 * S
    slot = (W - pad_l - pad_r) / max(1, len(data))
    bw = slot * 0.46
    layout = []
    for i, (day, ok, bad) in enumerate(data):
        x0, x1 = int(pad_l + i * slot + (slot - bw) / 2), int(pad_l + i * slot + (slot + bw) / 2)
        total = ok + bad
        bh = max(int(plot_h * total / top), 5 * S)
        y0 = int(base - bh)
        seg_w, seg_h = x1 - x0, int(base - y0)
        if total:
            ok_h = int(seg_h * ok / total)
            green = Image.composite(Image.new("RGB", (seg_w, max(ok_h, 1)), "#16A34A"), Image.new("RGB", (seg_w, max(ok_h, 1)), "#6EE7B7"),
                                    Image.linear_gradient("L").resize((seg_w, max(ok_h, 1))))
            red_h = seg_h - ok_h
            col = Image.new("RGB", (seg_w, seg_h), "#16A34A")
            if red_h > 0:
                red = Image.composite(Image.new("RGB", (seg_w, red_h), "#E11D48"), Image.new("RGB", (seg_w, red_h), "#FDA4AF"),
                                      Image.linear_gradient("L").resize((seg_w, red_h)))
                col.paste(red, (0, 0))
            if ok_h > 0:
                col.paste(green, (0, red_h))
        else:
            col = Image.new("RGB", (seg_w, seg_h), blend(bg, "#FFFFFF", 0.09))
        m = Image.new("L", (seg_w, seg_h), 0)
        ImageDraw.Draw(m).rounded_rectangle((0, 0, seg_w - 1, seg_h - 1), radius=min(9 * S, seg_w // 2, seg_h // 2), fill=255)
        img.paste(col, (x0, y0), m)
        layout.append(((x0 + x1) / 2 / S, y0 / S, total, day))
    return img.resize((w, h), Image.LANCZOS), layout, grid_ys, base / S, pad_l / S


# =============================================================================
#  VIEWS  (mixed into the main window class)
# =============================================================================
class ViewsMixin:
    # ------------------------------------------------------------------ HOME
    def build_home(self, v):
        v.grid_columnconfigure((0, 1), weight=1)
        v.grid_rowconfigure(2, weight=3)
        v.grid_rowconfigure(3, weight=2)
        self.hero = HeroBanner(v, height=142)
        self.hero.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 16))
        bg_hint = blend(COLOR_BG_MAIN, COLOR_ACCENT2, 0.30)
        self.hero.add_button(make_button(self.hero, "＋ New script", self.new_script_from_cmd, "primary", 132, 38, bg_color=bg_hint))
        self.hero.add_button(make_button(self.hero, "🔍 Command palette", self.open_palette, "ghost", 162, 38, bg_color=bg_hint,
                                         fg_color=blend(COLOR_BG_MAIN, "#FFFFFF", 0.10), border_color=blend(COLOR_BG_MAIN, "#FFFFFF", 0.25)))

        stats = ctk.CTkFrame(v, fg_color="transparent")
        stats.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 16))
        stats.grid_columnconfigure(tuple(range(5)), weight=1, uniform="stat")
        self.stat_labels = {}
        specs = [("scripts", "Scripts", "folder"), ("today", "Runs today", "bolt"), ("rate", "Success rate · 7d", "check"),
                 ("avg", "Avg runtime · 7d", "timer"), ("fail", "Last failure", "warn")]
        for i, (key, title, ic) in enumerate(specs):
            card = ctk.CTkFrame(stats, fg_color=COLOR_CARD_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
            card.grid(row=0, column=i, sticky="nsew", padx=6)
            c1 = ICON_STYLES[ic][0]
            ctk.CTkFrame(card, height=3, fg_color=c1, corner_radius=2).pack(fill="x", padx=22, pady=(0, 0))
            body = ctk.CTkFrame(card, fg_color="transparent")
            body.pack(fill="both", expand=True, padx=16, pady=(10, 14))
            body.grid_columnconfigure(0, weight=1)
            make_label(body, title, 12, muted=True).grid(row=0, column=0, sticky="w")
            val = make_label(body, "—", 30, True, display=True)
            val.grid(row=1, column=0, sticky="w")
            sub = make_label(body, "", 11, muted=True)
            sub.grid(row=2, column=0, sticky="w")
            ctk.CTkLabel(body, text="", image=icon_tile(ic, 38)).grid(row=0, column=1, rowspan=2, sticky="ne")
            self.stat_labels[key] = (val, sub)

        def panel(row, col, title, ic):
            f = ctk.CTkFrame(v, fg_color=COLOR_CARD_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
            f.grid(row=row, column=col, sticky="nsew", padx=6, pady=(0, 14))
            head = ctk.CTkFrame(f, fg_color="transparent")
            head.pack(fill="x", padx=18, pady=(14, 4))
            ctk.CTkLabel(head, text="", image=icon(ic, 18)).pack(side="left")
            make_label(head, title, 15, True, display=True).pack(side="left", padx=8)
            return f

        chart_card = panel(2, 0, "Runs · last 7 days", "chart")
        self.chart = tk.Canvas(chart_card, bg=COLOR_CARD_SURFACE, highlightthickness=0, height=170)
        self.chart.pack(fill="both", expand=True, padx=12, pady=(0, 12))
        self._chart_data = []
        self._chart_job = None
        self._chart_photo = None
        self.chart.bind("<Configure>", lambda e: self._sched_chart())
        act = panel(2, 1, "Recent activity", "clock")
        self.activity_frame = ctk.CTkScrollableFrame(act, fg_color="transparent", height=120)
        self.activity_frame.pack(fill="both", expand=True, padx=6, pady=(0, 8))
        streak = panel(3, 0, "Streak & achievements", "trophy")
        self.streak_row = ctk.CTkFrame(streak, fg_color="transparent")
        self.streak_row.pack(fill="x", padx=16)
        self.ach_frame = ctk.CTkScrollableFrame(streak, fg_color="transparent", height=90)
        self.ach_frame.pack(fill="both", expand=True, padx=6, pady=(6, 8))
        up = panel(3, 1, "Upcoming schedules", "timer")
        self.upcoming_frame = ctk.CTkScrollableFrame(up, fg_color="transparent", height=90)
        self.upcoming_frame.pack(fill="both", expand=True, padx=6, pady=(0, 8))

    def _sched_chart(self):
        if self._chart_job:
            self.after_cancel(self._chart_job)
        self._chart_job = self.after(80, self.draw_chart)

    def draw_chart(self):
        self._chart_job = None
        c = self.chart
        w, h = c.winfo_width(), c.winfo_height()
        if w < 120 or h < 80 or not self._chart_data:
            return
        try:
            img, layout, grid_ys, base, pad_l = chart_image(w, h, self._chart_data, COLOR_CARD_SURFACE)
            self._chart_photo = ImageTk.PhotoImage(img, master=c)
        except Exception as exc:  # noqa: BLE001
            print(f"[chart] {exc}")
            return
        c.delete("all")
        c.create_image(0, 0, image=self._chart_photo, anchor="nw")
        for y, val in grid_ys:
            c.create_text(pad_l - 8, y, text=str(val), anchor="e", fill=COLOR_TEXT_MUTED, font=(FONT_MONO, 9))
        for cx, ytop, total, day in layout:
            if total:
                c.create_text(cx, ytop - 10, text=str(total), fill=COLOR_TEXT_MAIN, font=(FONT_MONO, 10, "bold"))
            c.create_text(cx, base + 14, text=datetime.strptime(day, "%Y-%m-%d").strftime("%a %d"),
                          fill=COLOR_TEXT_MUTED, font=(FONT_FAMILY, 9))
        c.create_rectangle(pad_l + 4, 8, pad_l + 14, 18, fill="#4ADE80", outline="")
        c.create_text(pad_l + 20, 13, text="succeeded", anchor="w", fill=COLOR_TEXT_MUTED, font=(FONT_FAMILY, 9))
        c.create_rectangle(pad_l + 112, 8, pad_l + 122, 18, fill="#FB7185", outline="")
        c.create_text(pad_l + 128, 13, text="failed / stopped", anchor="w", fill=COLOR_TEXT_MUTED, font=(FONT_FAMILY, 9))

    def refresh_home(self):
        if not self.ready:
            return
        hour = datetime.now().hour
        greet = "Good morning" if hour < 12 else "Good afternoon" if hour < 18 else "Good evening"
        ov = self.db.overview()
        n_scripts = len(self.list_script_files())
        self.hero.set_text("mission control", f"{greet}, commander",
                           f"{n_scripts} scripts ready  ·  {ov['runs_today']} runs today  ·  workspace {workspace_dir()}")
        self.stat_labels["scripts"][0].configure(text=str(n_scripts))
        self.stat_labels["scripts"][1].configure(text=f"{len(self.db.schedules())} schedules · {len(self.db.pipelines())} pipelines")
        self.stat_labels["today"][0].configure(text=str(ov["runs_today"]))
        self.stat_labels["today"][1].configure(text=f"{ov['week_runs']} in the last 7 days")
        rate = ov["success_rate"]
        self.stat_labels["rate"][0].configure(text="—" if rate is None else f"{rate:.0f}%",
                                              text_color=COLOR_TEXT_MAIN if rate is None or rate >= 80 else COLOR_WARN)
        self.stat_labels["rate"][1].configure(text="no runs yet" if rate is None else "of runs succeeded")
        self.stat_labels["avg"][0].configure(text=fmt_duration(ov["avg_duration"]))
        self.stat_labels["avg"][1].configure(text="per run")
        lf = ov["last_failure"]
        self.stat_labels["fail"][0].configure(text="none" if not lf else time_ago(lf["started"]),
                                              text_color="#4ADE80" if not lf else "#FB7185")
        self.stat_labels["fail"][1].configure(text="all clear" if not lf else Path(lf["script"]).stem[:22])
        self._chart_data = self.db.daily_counts(7)
        self.after(30, self.draw_chart)

        for w in self.activity_frame.winfo_children():
            w.destroy()
        rows = self.db.recent_runs(9)
        if not rows:
            make_label(self.activity_frame, "No runs yet - launch a script from Command Center.", 12, muted=True).pack(pady=16)
        glyph = {"success": ("check", "#4ADE80"), "failed": ("close", "#FB7185"), "timeout": ("warn", "#FBBF24"),
                 "stopped": ("stop", "#FBBF24")}
        for r in rows:
            gname, gcol = glyph.get(r["status"], ("info", COLOR_ACCENT))
            line = ctk.CTkFrame(self.activity_frame, fg_color="transparent", cursor="hand2", corner_radius=8)
            line.pack(fill="x", pady=1)
            ctk.CTkLabel(line, text="", image=icon(gname, 15, gcol), width=22).pack(side="left", padx=(6, 4), pady=5)
            make_label(line, Path(r["script"]).stem, 12, True).pack(side="left")
            make_label(line, f"   {time_ago(r['started'])}", 11, muted=True).pack(side="left")
            ctk.CTkLabel(line, text=fmt_duration(r["duration"]), text_color=COLOR_TEXT_MUTED,
                         font=ctk.CTkFont(family=FONT_MONO, size=11)).pack(side="right", padx=10)
            bind_tree(line, "<Button-1>", lambda e, rid=r["id"]: self.open_history_run(rid))
            bind_tree(line, "<Enter>", lambda e, ln=line: ln.configure(fg_color=COLOR_CARD_HOVER))
            bind_tree(line, "<Leave>", lambda e, ln=line: ln.configure(fg_color="transparent"))

        stats = self.db.achievement_stats()
        unlocked = set(json.loads(self.db.kv_get("achievements", "[]")))
        for w in self.streak_row.winfo_children():
            w.destroy()
        for gname, txt, col in (("bolt", f"{stats['streak']}-day streak", "#FB923C"), ("check", f"{stats['ok']} successful runs", "#4ADE80"),
                                ("trophy", f"{len(unlocked)}/{len(ACHIEVEMENTS)} achievements", "#FDE047")):
            pf, _lbl = pill(self.streak_row, gname, txt, col)
            pf.pack(side="left", padx=(0, 8))
        for w in self.ach_frame.winfo_children():
            w.destroy()
        for i, (aid, title, desc, _pred) in enumerate(ACHIEVEMENTS):
            got = aid in unlocked
            chip = ctk.CTkFrame(self.ach_frame, fg_color=blend(COLOR_CARD_SURFACE, "#FDE047", 0.07) if got else "transparent",
                                corner_radius=10, border_width=1, border_color=blend(COLOR_BORDER, "#FDE047", 0.55) if got else COLOR_BORDER)
            chip.grid(row=i // 3, column=i % 3, sticky="ew", padx=4, pady=3)
            ctk.CTkLabel(chip, text="", image=icon_tile("trophy", 28) if got else icon("lock", 16, "#3A4566"), width=30).pack(side="left", padx=(8, 4), pady=6)
            col = ctk.CTkFrame(chip, fg_color="transparent")
            col.pack(side="left", fill="x", expand=True, pady=4)
            make_label(col, title, 12, True, muted=not got).pack(anchor="w")
            make_label(col, desc, 10, muted=True, wraplength=150, justify="left").pack(anchor="w")
        self.ach_frame.grid_columnconfigure((0, 1, 2), weight=1)

        for w in self.upcoming_frame.winfo_children():
            w.destroy()
        upcoming = []
        for s in self.db.schedules():
            nr = next_run(s)
            if nr:
                upcoming.append((nr, s))
        upcoming.sort(key=lambda t: t[0])
        if not upcoming:
            make_label(self.upcoming_frame, "Nothing scheduled. Open Automation to add a schedule.", 12, muted=True).pack(pady=16)
        for nr, s in upcoming[:6]:
            row = ctk.CTkFrame(self.upcoming_frame, fg_color="transparent")
            row.pack(fill="x", pady=3)
            ctk.CTkLabel(row, text="", image=icon_tile("timer", 28), width=30).pack(side="left", padx=(6, 8))
            col = ctk.CTkFrame(row, fg_color="transparent")
            col.pack(side="left")
            make_label(col, self.target_label(s["target"]), 12, True).pack(anchor="w")
            make_label(col, f"{nr:%a %d %b · %H:%M}  ·  {describe_schedule(s)}", 10, muted=True).pack(anchor="w")

    def check_achievements(self):
        try:
            stats = self.db.achievement_stats()
            unlocked = set(json.loads(self.db.kv_get("achievements", "[]")))
            fresh = [a for a in ACHIEVEMENTS if a[0] not in unlocked and a[3](stats)]
            if fresh:
                self.db.kv_set("achievements", json.dumps(sorted(unlocked | {a[0] for a in fresh})))
                for a in fresh:
                    self.toast(f"🏆 Achievement unlocked: {a[1]}\n{a[2]}", "success", 5000)
        except Exception as exc:  # noqa: BLE001
            print(f"[achievements] {exc}")

    @staticmethod
    def target_label(target: str) -> str:
        if target.startswith("pipe:"):
            return f"{target[5:]} (pipeline)"
        return target[7:] if target.startswith("script:") else target

    # -------------------------------------------------------- COMMAND CENTER
    def build_cmd(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(3, weight=1)
        head = ctk.CTkFrame(v, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew")
        make_label(head, "⚡ Quick Launch", 17, True).pack(side="left")
        make_label(head, "   pin a script with the star to add a shortcut here", 11, muted=True).pack(side="left", pady=(4, 0))
        self.quick_frame = ctk.CTkScrollableFrame(v, orientation="horizontal", height=112, fg_color=COLOR_CARD_SURFACE,
                                                  corner_radius=16, border_width=1, border_color=COLOR_BORDER)
        self.quick_frame.grid(row=1, column=0, sticky="ew", pady=(8, 14))
        bar = ctk.CTkFrame(v, fg_color="transparent")
        bar.grid(row=2, column=0, sticky="ew", pady=(0, 8))
        make_label(bar, "📁 Saved Automations", 17, True).pack(side="left")
        make_button(bar, "＋ New", self.new_script_from_cmd, "outline", 74).pack(side="right")
        make_button(bar, "🔄", self.refresh_scripts_forced, "ghost", 40).pack(side="right", padx=6)
        self.sort_var = tk.StringVar(value="Name")
        ctk.CTkOptionMenu(bar, values=["Name", "Recently run", "Most used"], variable=self.sort_var, width=130,
                          fg_color=COLOR_SIDEBAR, button_color=COLOR_BORDER,
                          command=lambda _v: self.refresh_scripts()).pack(side="right")
        self.filter_var = tk.StringVar(value="All")
        ctk.CTkSegmentedButton(bar, values=["All", "★ Pinned", "Python", "VBA"], variable=self.filter_var,
                               selected_color=COLOR_ACCENT, selected_hover_color=COLOR_ACCENT_HOVER,
                               command=lambda _v: self.refresh_scripts()).pack(side="right", padx=10)
        self.search_var = tk.StringVar()
        self.search_entry = make_entry(bar, "", width=220, textvariable=self.search_var)
        self.search_entry.pack(side="right")
        make_label(bar, "🔍 name or #tag", 11, muted=True).pack(side="right", padx=(0, 6))
        self._search_job = None
        self.search_var.trace_add("write", self._on_search_typed)

        self.paned = tk.PanedWindow(v, orient="vertical", sashwidth=8, bg=COLOR_BG_MAIN, bd=0, sashrelief="flat",
                                    opaqueresize=True)
        self.paned.grid(row=3, column=0, sticky="nsew")
        list_pane = ctk.CTkFrame(self.paned, fg_color=COLOR_BG_MAIN, corner_radius=0)
        self.script_list = ctk.CTkScrollableFrame(list_pane, fg_color=COLOR_CARD_SURFACE, corner_radius=16,
                                                  border_width=1, border_color=COLOR_BORDER)
        self.script_list.pack(fill="both", expand=True)
        console_pane = ctk.CTkFrame(self.paned, fg_color=COLOR_BG_MAIN, corner_radius=0)
        self.console = ConsolePanel(console_pane, on_stop=self.stop_from_console)
        self.console.pack(fill="both", expand=True, pady=(8, 0))
        self.paned.add(list_pane, minsize=150, stretch="always")
        self.paned.add(console_pane, minsize=200, stretch="always")
        self.after(250, self._place_sash)

    def _place_sash(self):
        try:
            self.paned.sash_place(0, 1, max(160, int(self.paned.winfo_height() * 0.48)))
        except tk.TclError:
            pass

    def _on_search_typed(self, *_):
        if self._search_job:
            self.after_cancel(self._search_job)
        self._search_job = self.after(180, self.refresh_scripts)

    def list_script_files(self) -> list:
        me = Path(__file__).name.lower()
        try:
            names = os.listdir(workspace_dir())
        except OSError:
            names = []
        return sorted((f for f in names if f.lower().endswith((".py", ".bas")) and f.lower() not in (me, "desktop_app.py")
                       and not f.startswith(("_", "."))), key=str.lower)

    def script_params(self, filename: str) -> list:
        if not filename.lower().endswith(".py"):
            return []
        path = workspace_dir() / filename
        try:
            mtime = path.stat().st_mtime
        except OSError:
            return []
        cached = self.param_cache.get(filename)
        if cached and cached[0] == mtime:
            return cached[1]
        try:
            params = parse_params(path.read_text(encoding="utf-8", errors="replace"))
        except OSError:
            params = []
        self.param_cache[filename] = (mtime, params)
        return params

    def refresh_scripts_forced(self):
        self.refresh_scripts(force=True)
        self.toast("Script list refreshed", "info", 1500)

    def _filtered(self, files: list) -> list:
        mode = self.filter_var.get() if hasattr(self, "filter_var") else "All"
        terms = [t for t in (self.search_var.get().lower().split() if hasattr(self, "search_var") else [])]
        out = []
        for f in files:
            meta = self.meta.get(f, {})
            if mode == "★ Pinned" and not meta.get("favorite"):
                continue
            if mode == "Python" and not f.lower().endswith(".py"):
                continue
            if mode == "VBA" and not f.lower().endswith(".bas"):
                continue
            ok = True
            for t in terms:
                if t.startswith("#"):
                    ok = ok and any(t[1:] in tag.lower() for tag in meta.get("tags", []))
                else:
                    ok = ok and t in f.lower()
            if ok:
                out.append(f)
        sort = self.sort_var.get() if hasattr(self, "sort_var") else "Name"
        if sort == "Recently run":
            out.sort(key=lambda f: (self.last_runs.get(f, {}).get("started") or ""), reverse=True)
        elif sort == "Most used":
            out.sort(key=lambda f: self.last_runs.get(f, {}).get("n", 0), reverse=True)
        return out

    def refresh_scripts(self, force: bool = False):
        if not self.ready:
            return
        files = self.list_script_files()
        self.meta = self.db.meta_all()
        self.last_runs = self.db.last_runs_by_script()
        self.next_by_target = {}
        for s in self.db.schedules():
            nr = next_run(s)
            if nr and (s["target"] not in self.next_by_target or nr < self.next_by_target[s["target"]]):
                self.next_by_target[s["target"]] = nr
        visible = self._filtered(files)
        sig = (tuple(visible), self.filter_var.get(), self.sort_var.get(),
               tuple(sorted((k, bool(m.get("favorite")), tuple(m.get("tags", []))) for k, m in self.meta.items() if k in files)),
               tuple(f for f in files if self.meta.get(f, {}).get("favorite")))
        if force or sig != self._sig:
            self._sig = sig
            self._rebuild_cards(files, visible)
        else:
            for f in visible:
                self.update_card(f)

    def _rebuild_cards(self, files: list, visible: list):
        for w in self.script_list.winfo_children():
            w.destroy()
        self.cards = {}
        if not visible:
            msg = ("No scripts yet - open Script Builder to create your first automation."
                   if not files else "No scripts match the current filter / search.")
            make_label(self.script_list, msg, 13, muted=True).pack(pady=30)
        for f in visible:
            self._build_card(f)
            self.update_card(f)
        # quick launch = pinned scripts
        for w in self.quick_frame.winfo_children():
            w.destroy()
        pinned = [f for f in files if self.meta.get(f, {}).get("favorite")]
        if not pinned:
            make_label(self.quick_frame, "☆  Nothing pinned yet - click the star on a script to add a shortcut.", 12,
                       muted=True).pack(pady=32, padx=20)
        for f in pinned:
            col = ctk.CTkFrame(self.quick_frame, fg_color="transparent")
            col.pack(side="left", padx=8, pady=2)
            kind = "vba" if f.lower().endswith(".bas") else "py"
            ctk.CTkButton(col, text="", image=lang_tile(kind, 50), width=74, height=66, corner_radius=16,
                          fg_color="transparent", hover_color=COLOR_CARD_HOVER, border_width=1, border_color=COLOR_BORDER,
                          command=lambda f=f: self.launch(f)).pack()
            nm = Path(f).stem
            make_label(col, nm if len(nm) <= 13 else nm[:12] + "…", 11, True).pack(pady=(3, 0))

    def _build_card(self, f: str):
        is_vba = f.lower().endswith(".bas")
        meta = self.meta.get(f, {})
        fav = bool(meta.get("favorite"))
        card = ctk.CTkFrame(self.script_list, fg_color=COLOR_BG_MAIN, corner_radius=14, border_width=1, border_color=COLOR_BORDER)
        card.pack(fill="x", padx=8, pady=5)
        make_button(card, "", lambda f=f: self.toggle_favorite(f), "flat", 34, 34, glyph="star" if fav else "star_o").pack(side="left", padx=(8, 0))
        ctk.CTkLabel(card, text="", image=lang_tile("vba" if is_vba else "py", 40)).pack(side="left", padx=(4, 12), pady=10)
        mid = ctk.CTkFrame(card, fg_color="transparent")
        mid.pack(side="left", pady=8)
        top = ctk.CTkFrame(mid, fg_color="transparent")
        top.pack(anchor="w")
        make_label(top, Path(f).stem, 15, True).pack(side="left")
        make_label(top, "  Excel VBA" if is_vba else "  Python", 10, muted=True).pack(side="left", pady=(3, 0))
        for t in meta.get("tags", [])[:4]:
            ctk.CTkLabel(top, text=f"#{t}", text_color="#C4B5FD", fg_color=blend(COLOR_BG_MAIN, "#A78BFA", 0.16), corner_radius=8,
                         padx=7, pady=0, height=18, font=ctk.CTkFont(family=FONT_FAMILY, size=10, weight="bold")).pack(side="left", padx=(6, 0), pady=(2, 0))
        info = make_label(mid, "", 11, muted=True)
        info.pack(anchor="w")
        badge = ctk.CTkLabel(card, text="Idle", corner_radius=10, padx=12, pady=3, height=24,
                             font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold"))
        badge.pack(side="left", padx=16)
        make_button(card, "🗑 Delete", lambda f=f: self.delete_script(f), "danger", 92).pack(side="right", padx=(4, 12), pady=12)
        make_button(card, "✏ Edit", lambda f=f: self.edit_script(f), "ghost", 76).pack(side="right", padx=4)
        make_button(card, "", lambda f=f: self.edit_tags(f), "ghost", 36, glyph="tag").pack(side="right", padx=4)
        make_button(card, "", lambda f=f: self.open_history_for(f), "ghost", 36, glyph="clock").pack(side="right", padx=4)
        prev = make_button(card, "", lambda f=f: self.preview_outputs(f), "ghost", 36, glyph="chart")
        prev.pack(side="right", padx=4)
        run = make_button(card, "▶ Run", None, "success", 98)
        run.pack(side="right", padx=4)
        if self.dnd_ok and not is_vba:
            try:
                card.drop_target_register(DND_FILES)
                card.dnd_bind("<<Drop>>", lambda e, f=f: self.on_drop_file(f, e))
            except Exception:  # noqa: BLE001
                pass
        self.cards[f] = {"frame": card, "badge": badge, "info": info, "run": run, "prev": prev}

    def update_card(self, f: str):
        c = self.cards.get(f)
        if not c:
            return
        st = self.script_status.get(f, "idle")
        label, fg, bg = status_style(st)
        c["badge"].configure(text=label, text_color=fg, fg_color=bg)
        c["frame"].configure(border_color={"running": COLOR_ACCENT2, "failed": "#5B2333", "timeout": "#5B2333"}.get(st, COLOR_BORDER))
        if st == "running":
            c["run"].configure(text="Stop", image=icon("stop", 16, "#FFFFFF"), fg_color=COLOR_DANGER,
                               hover_color=COLOR_DANGER_HOVER, command=lambda f=f: self.stop_script(f))
        else:
            txt = "Run…" if self.script_params(f) else "Run"
            c["run"].configure(text=txt, image=icon("play", 16, "#FFFFFF"), fg_color=COLOR_SUCCESS,
                               hover_color=COLOR_SUCCESS_HOVER, command=lambda f=f: self.launch(f))
        last = self.last_runs.get(f)
        parts = ["never run"]
        outputs = []
        if last:
            parts = [time_ago(last["started"]), fmt_duration(last["duration"]), f"{last['n']} run{'s' if last['n'] != 1 else ''}"]
            try:
                outputs = json.loads(last.get("outputs") or "[]")
            except Exception:  # noqa: BLE001
                outputs = []
        nxt = self.next_by_target.get(f"script:{f}")
        if nxt:
            parts.append(f"next {nxt:%a %H:%M}")
        c["info"].configure(text="  ·  ".join(parts))
        c["prev"].configure(state="normal" if any(os.path.exists(o) for o in outputs) else "disabled")

    # ---- card actions ---------------------------------------------------------
    def toggle_favorite(self, f: str):
        self.db.set_favorite(f, not self.meta.get(f, {}).get("favorite"))
        self.refresh_scripts(force=True)

    def edit_tags(self, f: str):
        cur = ", ".join(self.meta.get(f, {}).get("tags", []))
        val = ask_text(self, "Tags", f"Tags for {f} (comma separated)", cur)
        if val is None:
            return
        tags = sorted({re.sub(r"[^\w\-]", "", t.strip().lstrip("#")) for t in val.split(",") if t.strip()} - {""})
        self.db.set_tags(f, tags)
        self.refresh_scripts(force=True)

    def delete_script(self, f: str):
        if not ask_confirm(self, "Delete script?", f"'{f}' will be moved to the Recycle Bin (or the workspace _trash folder "
                                                   "if send2trash is not installed).", "Delete", danger=True):
            return
        path = workspace_dir() / f
        try:
            try:
                from send2trash import send2trash
                send2trash(str(path))
                where = "Recycle Bin"
            except ImportError:
                trash = workspace_dir() / "_trash"
                trash.mkdir(exist_ok=True)
                shutil.move(str(path), str(trash / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}"))
                where = "workspace/_trash"
            self.toast(f"Deleted {f}  →  {where}", "warning")
        except Exception as exc:  # noqa: BLE001
            self.toast(f"Could not delete: {exc}", "error")
        self.refresh_scripts(force=True)

    def edit_script(self, f: str):
        if not self.confirm_discard_editor():
            return
        self.load_into_editor(f)
        self.select_view("builder")

    def open_history_for(self, f: str):
        self.hist_script.set(f)
        self.select_view("history")

    def preview_outputs(self, f: str):
        last = self.last_runs.get(f)
        try:
            outs = [o for o in json.loads((last or {}).get("outputs") or "[]") if os.path.exists(o)]
        except Exception:  # noqa: BLE001
            outs = []
        if outs:
            PreviewWindow(self, outs)
        else:
            self.toast("No output files were detected for the last run.", "info")

    def on_drop_file(self, f: str, event):
        try:
            files = self.tk.splitlist(event.data)
        except Exception:  # noqa: BLE001
            files = [event.data]
        if files:
            self.launch(f, trigger_file=files[0])

    def new_script_from_cmd(self):
        if not self.confirm_discard_editor():
            return
        self.new_script()
        self.select_view("builder")

    # ------------------------------------------------------------ SCRIPT BUILDER
    def build_builder(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(4, weight=1)
        bar = ctk.CTkFrame(v, fg_color=COLOR_CARD_SURFACE, corner_radius=14, border_width=1, border_color=COLOR_BORDER)
        bar.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        make_label(bar, "🛠  Engine mode", 13, True).pack(side="left", padx=(16, 10), pady=10)
        self.btn_mode_py = ctk.CTkButton(bar, text="Python  (.py)", corner_radius=10, height=34, width=150,
                                         image=lang_tile("py", 22), compound="left",
                                         font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"),
                                         command=lambda: self.set_editor_mode("python"))
        self.btn_mode_py.pack(side="left", padx=4)
        self.btn_mode_vba = ctk.CTkButton(bar, text="Excel VBA  (.bas)", corner_radius=10, height=34, width=160,
                                          image=lang_tile("vba", 22), compound="left",
                                          font=ctk.CTkFont(family=FONT_FAMILY, size=12, weight="bold"),
                                          command=lambda: self.set_editor_mode("vba"))
        self.btn_mode_vba.pack(side="left", padx=4)
        make_button(bar, "✨ AI Assist", self.open_ai, "outline", 110).pack(side="right", padx=12)
        self.snippet_menu = ctk.CTkOptionMenu(bar, values=list(SNIPPETS["python"]), command=self.insert_snippet, width=210,
                                              fg_color=COLOR_BG_MAIN, button_color=COLOR_BORDER)
        self.snippet_menu.set("Snippets")
        self.snippet_menu.pack(side="right")

        row = ctk.CTkFrame(v, fg_color="transparent")
        row.grid(row=1, column=0, sticky="ew")
        make_label(row, "Script / Macro Name:", 13, True).pack(side="left")
        self.dirty_label = make_label(row, "", 12)
        self.dirty_label.pack(side="left", padx=12)
        self.entry_name = make_entry(v, "e.g., EPM_Data_Transform", width=400, height=44)
        self.entry_name.configure(height=44)
        self.entry_name.grid(row=2, column=0, sticky="ew", pady=(6, 12))
        self.entry_name.bind("<KeyRelease>", lambda e: self.update_dirty())

        tools = ctk.CTkFrame(v, fg_color="transparent")
        tools.grid(row=3, column=0, sticky="ew", pady=(0, 6))
        self.lbl_editor_title = make_label(tools, "Python Code Editor:", 13, True)
        self.lbl_editor_title.pack(side="left")
        make_button(tools, "🧹 Clear", self.clear_editor, "danger", 76).pack(side="right")
        make_button(tools, "📋 Paste", self.paste_into_editor, "ghost", 80).pack(side="right", padx=6)
        make_button(tools, "💬 Comment", lambda: self.editor.toggle_comment(), "ghost", 96).pack(side="right")
        make_button(tools, "🕘 History", self.show_editor_history, "ghost", 90).pack(side="right", padx=6)
        make_button(tools, "🔎 Lint", self.lint_editor, "ghost", 70).pack(side="right")
        self.btn_format = make_button(tools, "✨ Format", self.format_editor, "ghost", 86)
        self.btn_format.pack(side="right", padx=6)

        self.editor = CodeEditor(v, on_change=self.on_editor_change, on_cursor=self.on_editor_cursor)
        self.editor.grid(row=4, column=0, sticky="nsew")
        status = ctk.CTkFrame(v, fg_color="transparent")
        status.grid(row=5, column=0, sticky="ew", pady=(4, 0))
        self.lbl_cursor = make_label(status, "Ln 1, Col 1", 11, muted=True)
        self.lbl_cursor.pack(side="left")
        self.lbl_editor_info = make_label(status, "Python · UTF-8 · Ctrl+S save · Ctrl+/ comment · Ctrl+K palette", 11, muted=True)
        self.lbl_editor_info.pack(side="right")

        actions = ctk.CTkFrame(v, fg_color="transparent")
        actions.grid(row=6, column=0, sticky="ew", pady=(10, 0))
        make_button(actions, "🔍 Check Syntax", self.check_syntax, "outline", 150, height=42).pack(side="left")
        make_button(actions, "▶ Run Test", self.builder_run_test, "ghost", 120, height=42).pack(side="left", padx=10)
        make_button(actions, "💾 Save to Workspace", self.save_script, "success", 190, height=42).pack(side="right")
        self.set_editor_mode(self.editor_mode, keep=True)

    def set_editor_mode(self, mode: str, keep: bool = False):
        self.editor_mode = mode
        py = mode == "python"
        self.btn_mode_py.configure(fg_color=blend(COLOR_CARD_SURFACE, COLOR_ACCENT, 0.30) if py else "transparent",
                                   text_color=COLOR_TEXT_MAIN if py else COLOR_TEXT_MUTED, hover_color=COLOR_CARD_HOVER)
        self.btn_mode_vba.configure(fg_color=blend(COLOR_CARD_SURFACE, COLOR_ACCENT, 0.30) if not py else "transparent",
                                    text_color=COLOR_TEXT_MAIN if not py else COLOR_TEXT_MUTED, hover_color=COLOR_CARD_HOVER)
        self.lbl_editor_title.configure(text="Python Code Editor:" if py else "VBA Macro Editor:")
        self.snippet_menu.configure(values=list(SNIPPETS[mode]))
        self.snippet_menu.set("Snippets")
        self.btn_format.configure(state="normal" if py else "disabled")
        self.lbl_editor_info.configure(text=("Python" if py else "VBA") + " · UTF-8 · Ctrl+S save · Ctrl+/ comment · Ctrl+K palette")
        self.editor.set_mode(mode)

    def on_editor_change(self):
        self.update_dirty()

    def on_editor_cursor(self, line, col):
        self.lbl_cursor.configure(text=f"Ln {line}, Col {col}")

    def is_dirty(self) -> bool:
        if not hasattr(self, "editor"):
            return False
        return (self.editor.get() != self.editor_original or safe_name(self.entry_name.get()) != self.editor_orig_name) \
            and bool(self.editor.get().strip() or self.entry_name.get().strip())

    def update_dirty(self):
        if self.is_dirty():
            self.dirty_label.configure(text="● unsaved changes", text_color=COLOR_WARN)
        else:
            self.dirty_label.configure(text="✓ saved" if self.editor_orig_name else "", text_color=COLOR_TEXT_MUTED)

    def confirm_discard_editor(self) -> bool:
        if not self.is_dirty():
            return True
        return ask_confirm(self, "Discard unsaved changes?", "The editor has changes that are not saved to the workspace.",
                           "Discard", danger=True)

    def load_into_editor(self, filename: str):
        path = workspace_dir() / filename
        try:
            code = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            self.toast(f"Could not open {filename}: {exc}", "error")
            return
        self.set_editor_mode("vba" if filename.lower().endswith(".bas") else "python")
        self.entry_name.delete(0, "end")
        self.entry_name.insert(0, path.stem)
        self.editor.set(code)
        self.editor.clear_marks()
        self.editor_original, self.editor_orig_name, self.editor_file = code, path.stem, filename
        self.update_dirty()

    def new_script(self):
        self.entry_name.delete(0, "end")
        self.editor.set("")
        self.editor_original, self.editor_orig_name, self.editor_file = "", "", None
        self.update_dirty()

    def clear_editor(self):
        if self.editor.get().strip() and self.is_dirty():
            if not ask_confirm(self, "Clear editor?", "This removes the code currently in the editor.", "Clear", danger=True):
                return
        self.editor.clear()
        self.update_dirty()

    def paste_into_editor(self):
        try:
            text = self.clipboard_get()
        except tk.TclError:
            self.toast("Clipboard is empty (or not text).", "warning")
            return
        self.editor.insert_at_cursor(text)

    def insert_snippet(self, label: str):
        code = SNIPPETS[self.editor_mode].get(label)
        if code:
            self.editor.insert_at_cursor(code)
        self.snippet_menu.set("Snippets")

    def show_editor_history(self):
        if not self.editor_file:
            self.toast("Save the script first - snapshots are stored on every overwrite.", "info")
            return
        show_versions(self, self.editor_file, self.editor.get(), self._restore_version)

    def _restore_version(self, text: str):
        self.editor.set(text)
        self.update_dirty()
        self.toast("Version restored into the editor (not saved yet).", "info")

    def open_ai(self):
        show_ai_assistant(self, "Python" if self.editor_mode == "python" else "VBA", self.editor.get,
                          lambda: self.last_error_text, self._ai_replace, self.editor.insert_at_cursor)

    def _ai_replace(self, code: str):
        self.editor.set(code)
        self.update_dirty()

    def format_editor(self):
        code = self.editor.get()
        if not code.strip():
            return
        self.btn_format.configure(state="disabled")

        def done(kind, val):
            self.btn_format.configure(state="normal" if self.editor_mode == "python" else "disabled")
            if kind == "ok":
                tool, text = val
                self.editor.set(text)
                self.update_dirty()
                self.toast(f"Formatted with {tool}", "success", 1800)
            else:
                self.toast(str(val), "error", 5000)

        run_bg(self, lambda: format_code(code), done)

    def lint_editor(self):
        self.select_view("test")
        self.run_lint_ui()

    def check_syntax(self):
        self.select_view("test")
        self.validate_syntax()

    def builder_run_test(self):
        self.select_view("test")
        self.run_sandbox_test()

    # ---- saving ---------------------------------------------------------------
    def save_script(self, event=None) -> bool:
        name = safe_name(self.entry_name.get())
        code = self.editor.get()
        if not name:
            self.toast("Enter a script name first.", "warning")
            return False
        if not code.strip():
            self.toast("The editor is empty - nothing to save.", "warning")
            return False
        vba = self.editor_mode == "vba"
        path = workspace_dir() / f"{name}{'.bas' if vba else '.py'}"
        if path.exists():
            old = path.read_text(encoding="utf-8", errors="replace")
            if old.replace("\r\n", "\n") == code.replace("\r\n", "\n"):
                self.editor_original, self.editor_orig_name, self.editor_file = code, name, path.name
                self.update_dirty()
                self.toast("No changes to save.", "info", 1500)
                return True
            if not ask_diff(self, f"Overwrite {path.name}?", old, code):
                return False
            save_version(path.name, old)
        try:
            with open(path, "w", encoding="utf-8", newline="\r\n" if vba else "\n") as fh:
                fh.write(code)
        except OSError as exc:
            self.toast(f"Could not save: {exc}", "error")
            return False
        self.editor_original, self.editor_orig_name, self.editor_file = code, name, path.name
        self.param_cache.pop(path.name, None)
        try:
            DRAFT_FILE.unlink()
        except OSError:
            pass
        self.update_dirty()
        self.toast(f"Saved {path.name} to workspace", "success")
        self.refresh_scripts(force=True)
        return True

    # ---- draft autosave -------------------------------------------------------
    def autosave_draft(self):
        try:
            if self.ready and self.is_dirty():
                DRAFT_FILE.write_text(json.dumps({"name": self.entry_name.get(), "mode": self.editor_mode,
                                                  "code": self.editor.get(), "file": self.editor_file, "saved": now_iso()}),
                                      encoding="utf-8")
        except Exception as exc:  # noqa: BLE001
            print(f"[draft] {exc}")
        finally:
            self.after(20000, self.autosave_draft)

    def restore_draft(self):
        try:
            if DRAFT_FILE.exists() and not self.editor.get().strip():
                d = json.loads(DRAFT_FILE.read_text(encoding="utf-8"))
                if d.get("code", "").strip():
                    self.set_editor_mode(d.get("mode", "python"))
                    self.entry_name.delete(0, "end")
                    self.entry_name.insert(0, d.get("name", ""))
                    self.editor.set(d["code"])
                    self.editor_file = d.get("file")
                    self.update_dirty()
                    self.toast(f"Unsaved draft restored ({time_ago(d.get('saved'))}). Open Script Builder to continue.", "info", 5000)
        except Exception as exc:  # noqa: BLE001
            print(f"[draft] {exc}")

    # ------------------------------------------------------------------ SANDBOX
    def build_test(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(2, weight=1)
        head = ctk.CTkFrame(v, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew")
        make_label(head, "Validate, lint and test-run without touching your workspace", 12, muted=True).pack(side="left")
        make_button(head, "🧹 Clean Sandbox", self.clean_sandbox, "danger", 140).pack(side="right")
        tools = ctk.CTkFrame(v, fg_color="transparent")
        tools.grid(row=1, column=0, sticky="ew", pady=10)
        make_button(tools, "🔍 Validate Syntax", self.validate_syntax, "outline", 150).pack(side="left")
        make_button(tools, "🔎 Lint", self.run_lint_ui, "ghost", 80).pack(side="left", padx=8)
        make_button(tools, "📦 Dependencies", self.check_deps_ui, "ghost", 130).pack(side="left")
        self.btn_install_missing = make_button(tools, "⬇ Install missing", self.install_sandbox_missing, "ghost", 150, state="disabled")
        self.btn_install_missing.pack(side="left", padx=8)
        self.btn_sb_stop = make_button(tools, "⏹ Stop", lambda: self.runner.stop(SANDBOX_KEY), "danger", 80, state="disabled")
        self.btn_sb_stop.pack(side="right")
        self.btn_sb_run = make_button(tools, "▶ Run Test", self.run_sandbox_test, "success", 120)
        self.btn_sb_run.pack(side="right", padx=8)
        wrap, self.test_console = make_terminal(v, "finprime ~ sandbox  —  nothing here is saved to history")
        wrap.grid(row=2, column=0, sticky="nsew")
        self.sandbox_missing: list = []

    def sb(self, text: str, tag: str = "out"):
        console_write(self.test_console, text, tag)

    def sb_header(self, what: str):
        name = self.entry_name.get().strip() or "(unsaved)"
        self.sb(f"\n[{datetime.now():%H:%M:%S}] {what} · {name} ({'Python' if self.editor_mode == 'python' else 'VBA'})\n", "info")

    def clean_sandbox(self):
        self.test_console.configure(state="normal")
        tk_text(self.test_console).delete("1.0", "end")
        self.test_console.configure(state="disabled")
        self.editor.clear_marks()
        try:
            shutil.rmtree(Path(tempfile.gettempdir()) / "finprime_sandbox", ignore_errors=True)
        except Exception:  # noqa: BLE001
            pass

    def validate_syntax(self):
        code = self.editor.get()
        self.editor.clear_marks()
        self.sb_header("Validating")
        if not code.strip():
            self.sb("⚠ Editor is empty - nothing to validate.\n", "warning")
            return
        if self.editor_mode == "vba":
            issues = [i for i in check_vba(code) if not i.split(":", 1)[1].strip().startswith("tip")]
            if issues:
                self.sb("❌ VBA structure problems:\n", "error")
                for i in issues:
                    self.sb(f"   line {i}\n", "error")
                    m = re.match(r"(\d+):", i)
                    if m and int(m.group(1)) > 0:
                        self.editor.mark_line(int(m.group(1)))
            else:
                self.sb("✅ VBA block structure looks valid. (Full compilation happens inside Excel.)\n", "success")
            return
        try:
            tree = ast.parse(code)
        except SyntaxError as exc:
            self.sb(f"❌ SyntaxError on line {exc.lineno}: {exc.msg}\n", "error")
            if exc.text:
                self.sb(f"   {exc.text.rstrip()}\n   {' ' * max(0, (exc.offset or 1) - 1)}^\n", "error")
            if exc.lineno:
                self.editor.mark_line(exc.lineno)
                self.editor.goto_line(exc.lineno)
            return
        funcs = sum(isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) for n in ast.walk(tree))
        imps = len(scan_imports(code, workspace_dir()))
        self.sb("✅ Syntax valid - Poisonous Py is ready to be run.\n", "success")
        self.sb(f"   {len(code.splitlines())} lines · {funcs} functions · {imps} third-party imports · "
                f"{len(parse_params(code))} parameters detected\n", "dim")

    def run_lint_ui(self):
        code = self.editor.get()
        self.editor.clear_marks()
        self.sb_header("Linting")
        if not code.strip():
            self.sb("⚠ Editor is empty.\n", "warning")
            return
        if self.editor_mode == "vba":
            issues = check_vba(code)
            self.sb("VBA checker: " + ("no issues" if not issues else f"{len(issues)} notes") + "\n", "success" if not issues else "warning")
            for i in issues:
                self.sb(f"   {i}\n", "warning")
            return
        self.sb("Running linter…\n", "dim")

        def done(kind, val):
            if kind == "err":
                self.sb(f"❌ Lint failed: {val}\n", "error")
                return
            tool, issues = val
            if not issues:
                self.sb(f"✅ {tool}: no issues found.\n", "success")
                return
            self.sb(f"⚠ {tool}: {len(issues)} finding(s)\n", "warning")
            for i in issues[:200]:
                self.sb(f"   {i}\n", "error" if "SyntaxError" in i or " E9" in i else "warning")
                m = re.search(r"(?:script\.py|<stdin>|stdin)?:?(\d+):\d+", i) or re.match(r"(\d+):", i)
                if m:
                    self.editor.mark_line(int(m.group(1)))
            if tool == "built-in checker":
                self.sb("   tip: install ruff or pyflakes from Env Health Check for deeper analysis.\n", "dim")

        run_bg(self, lambda: run_lint(code), done)

    def check_deps_ui(self):
        code = self.editor.get()
        self.sb_header("Checking dependencies")
        if self.editor_mode == "vba":
            self.sb("VBA modules have no Python dependencies. Excel + pywin32 are required to run them.\n", "dim")
            return
        mods = sorted(scan_imports(code, workspace_dir()))
        if not mods:
            self.sb("✅ No third-party imports - only the standard library.\n", "success")
            return

        def done(kind, val):
            if kind == "err":
                self.sb(f"❌ {val}\n", "error")
                return
            missing = []
            for m in mods:
                ok = val.get(m, False)
                self.sb(f"   {'✓' if ok else '✗'} {m}" + ("" if ok else f"   →  pip install {pip_name(m)}") + "\n",
                        "success" if ok else "error")
                if not ok:
                    missing.append(pip_name(m))
            self.sandbox_missing = sorted(set(missing))
            self.btn_install_missing.configure(state="normal" if missing else "disabled")
            if not missing:
                self.sb("✅ All dependencies are installed.\n", "success")

        run_bg(self, lambda: check_modules(mods), done)

    def install_sandbox_missing(self):
        pkgs = list(self.sandbox_missing)
        if pkgs:
            self.btn_install_missing.configure(state="disabled")
            self.install_packages(pkgs, self.test_console, lambda: self.toast("Install finished - re-run Dependencies to verify.", "info"))

    def run_sandbox_test(self):
        code = self.editor.get()
        if not code.strip():
            self.toast("Nothing to test - the editor is empty.", "warning")
            return
        if self.runner.is_running(SANDBOX_KEY):
            self.toast("A sandbox test is already running.", "info")
            return
        vba = self.editor_mode == "vba"
        if not vba:
            try:
                ast.parse(code)
            except SyntaxError as exc:
                self.sb_header("Test run")
                self.sb(f"❌ Fix the syntax error first - line {exc.lineno}: {exc.msg}\n", "error")
                self.editor.mark_line(exc.lineno or 1)
                return
        folder = Path(tempfile.gettempdir()) / "finprime_sandbox"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / ("_sandbox.bas" if vba else "_sandbox.py")
        with open(path, "w", encoding="utf-8", newline="\r\n" if vba else "\n") as fh:
            fh.write(code)
        params = [] if vba else parse_params(code)
        values = None
        if params:
            values = ask_params(self, "sandbox test", params, self.db.get_last_params(self.editor_file or "") if self.editor_file else {})
            if values is None:
                return
        vals = {p.name: p.default for p in params}
        vals.update(values or {})
        env = build_param_env(params, vals)
        self.start_run(SANDBOX_KEY, self.entry_name.get().strip() or "sandbox", path, "vba" if vba else "python",
                       build_cli(params, vals), env, "sandbox")

    def update_sandbox_buttons(self):
        running = self.script_status.get(SANDBOX_KEY) == "running"
        self.btn_sb_run.configure(state="disabled" if running else "normal")
        self.btn_sb_stop.configure(state="normal" if running else "disabled")

    # ------------------------------------------------------------------ HISTORY
    def build_history(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(1, weight=1)
        head = ctk.CTkFrame(v, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        make_label(head, "Click a run to inspect its full output", 12, muted=True).pack(side="left")
        make_button(head, "🗑 Clear history", self.clear_history, "danger", 130).pack(side="right")
        make_button(head, "🔄", self.refresh_history, "ghost", 40).pack(side="right", padx=8)
        self.hist_status = tk.StringVar(value="Any status")
        ctk.CTkOptionMenu(head, values=["Any status", "success", "failed", "stopped", "timeout"], variable=self.hist_status,
                          width=130, fg_color=COLOR_SIDEBAR, button_color=COLOR_BORDER,
                          command=lambda _v: self.refresh_history()).pack(side="right")
        self.hist_script = tk.StringVar(value="All scripts")
        self.hist_menu = ctk.CTkOptionMenu(head, values=["All scripts"], variable=self.hist_script, width=240,
                                           fg_color=COLOR_SIDEBAR, button_color=COLOR_BORDER,
                                           command=lambda _v: self.refresh_history())
        self.hist_menu.pack(side="right", padx=8)
        paned = tk.PanedWindow(v, orient="vertical", sashwidth=8, bg=COLOR_BG_MAIN, bd=0, sashrelief="flat")
        paned.grid(row=1, column=0, sticky="nsew")
        top = ctk.CTkFrame(paned, fg_color=COLOR_BG_MAIN, corner_radius=0)
        cols = [("id", "#", 60), ("when", "Started", 170), ("script", "Script", 300), ("status", "Status", 110),
                ("dur", "Duration", 100), ("exit", "Exit", 60), ("src", "Source", 130)]
        self.hist_tree = self.make_tree(top, cols)
        self.hist_tree.bind("<<TreeviewSelect>>", self.on_history_select)
        bottom = ctk.CTkFrame(paned, fg_color=COLOR_BG_MAIN, corner_radius=0)
        bar = ctk.CTkFrame(bottom, fg_color="transparent")
        bar.pack(fill="x", pady=(8, 6))
        make_label(bar, "Run details", 14, True).pack(side="left")
        make_button(bar, "📋 Copy log", self.copy_run_log, "ghost", 100).pack(side="right")
        make_button(bar, "📊 Preview output", self.preview_run_outputs, "ghost", 140).pack(side="right", padx=8)
        make_button(bar, "▶ Re-run", self.rerun_selected, "success", 100).pack(side="right")
        self.hist_detail = make_console_box(bottom)
        self.hist_detail.pack(fill="both", expand=True)
        paned.add(top, minsize=140, stretch="always")
        paned.add(bottom, minsize=160, stretch="always")
        self.hist_paned = paned
        self.after(300, self._place_hist_sash)
        self.hist_rows: dict = {}
        self.selected_run = None

    def _place_hist_sash(self):
        try:
            self.hist_paned.sash_place(0, 1, max(150, int(self.hist_paned.winfo_height() * 0.5)))
        except tk.TclError:
            pass

    def make_tree(self, parent, cols, height=None):
        wrap = ctk.CTkFrame(parent, fg_color=COLOR_BG_MAIN, border_width=1, border_color=COLOR_BORDER, corner_radius=12)
        wrap.pack(fill="both", expand=True)
        wrap.grid_rowconfigure(0, weight=1)
        wrap.grid_columnconfigure(0, weight=1)
        tree = ttk.Treeview(wrap, columns=[c[0] for c in cols], show="headings", selectmode="browse")
        for cid, title, width in cols:
            tree.heading(cid, text=title)
            tree.column(cid, width=width, anchor="w", stretch=True)
        sb = ttk.Scrollbar(wrap, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=sb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        tree.tag_configure("success", foreground="#3fb950")
        tree.tag_configure("failed", foreground=COLOR_DANGER_HOVER)
        tree.tag_configure("timeout", foreground=COLOR_WARN)
        tree.tag_configure("stopped", foreground=COLOR_WARN)
        tree.tag_configure("off", foreground=COLOR_TEXT_MUTED)
        tree.tag_configure("odd", background=blend(COLOR_BG_MAIN, "#FFFFFF", 0.03))
        return tree

    def refresh_history(self):
        if not self.ready:
            return
        scripts = [r["script"] for r in self.db.q("SELECT DISTINCT script FROM runs ORDER BY script")]
        self.hist_menu.configure(values=["All scripts"] + scripts)
        if self.hist_script.get() not in ["All scripts"] + scripts:
            self.hist_script.set("All scripts")
        self.hist_tree.delete(*self.hist_tree.get_children())
        self.hist_rows = {}
        for r in self.db.list_runs(self.hist_script.get(), self.hist_status.get()):
            self.hist_rows[str(r["id"])] = r
            self.hist_tree.insert("", "end", iid=str(r["id"]), tags=(r["status"], "odd" if len(self.hist_rows) % 2 else "even"),
                                  values=(r["id"], r["started"], r["script"], r["status"], fmt_duration(r["duration"]),
                                          r["exit_code"], r["source"] or "manual"))
        self.hist_detail.configure(state="normal")
        tk_text(self.hist_detail).delete("1.0", "end")
        self.hist_detail.configure(state="disabled")
        self.selected_run = None

    def on_history_select(self, _e=None):
        sel = self.hist_tree.selection()
        if not sel:
            return
        run = self.db.get_run(int(sel[0]))
        if not run:
            return
        self.selected_run = run
        box = self.hist_detail
        box.configure(state="normal")
        tk_text(box).delete("1.0", "end")
        box.configure(state="disabled")
        tag = "success" if run["status"] == "success" else "error" if run["status"] in ("failed", "timeout") else "warning"
        console_write(box, f"{run['script']}  ·  {run['status'].upper()}  ·  exit {run['exit_code']}  ·  "
                           f"{fmt_duration(run['duration'])}  ·  {run['started']}  ·  source: {run['source']}\n", tag)
        args = json.loads(run["args"] or "[]")
        if args:
            console_write(box, "args: " + " ".join(args) + "\n", "dim")
        if run["stdout"]:
            console_write(box, "\n── stdout ──\n", "info")
            console_write(box, run["stdout"], "out")
        if run["stderr"]:
            console_write(box, "\n── stderr ──\n", "info")
            console_write(box, run["stderr"], "stderr")
        outs = json.loads(run["outputs"] or "[]")
        if outs:
            console_write(box, "\n── output files ──\n" + "\n".join(outs) + "\n", "info")
        tk_text(box).see("1.0")

    def open_history_run(self, run_id: int):
        self.hist_script.set("All scripts")
        self.hist_status.set("Any status")
        self.select_view("history")
        iid = str(run_id)
        if self.hist_tree.exists(iid):
            self.hist_tree.selection_set(iid)
            self.hist_tree.see(iid)

    def copy_run_log(self):
        if self.selected_run:
            self.clipboard_clear()
            self.clipboard_append(f"{self.selected_run['stdout']}\n{self.selected_run['stderr']}")
            self.toast("Log copied to clipboard", "info", 1400)

    def preview_run_outputs(self):
        if not self.selected_run:
            return
        outs = [o for o in json.loads(self.selected_run["outputs"] or "[]") if os.path.exists(o)]
        if outs:
            PreviewWindow(self, outs)
        else:
            self.toast("No output files recorded for this run (or they were moved).", "info")

    def rerun_selected(self):
        if self.selected_run:
            self.launch(self.selected_run["script"])

    def clear_history(self):
        scope = self.hist_script.get()
        label = "all runs" if scope == "All scripts" else f"all runs of {scope}"
        if ask_confirm(self, "Clear history?", f"This permanently deletes {label} from the run history.", "Clear", danger=True):
            self.db.clear_runs(None if scope == "All scripts" else scope)
            self.refresh_history()
            self.refresh_scripts(force=True)

    # --------------------------------------------------------------- AUTOMATION
    def build_auto(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(1, weight=1)
        head = ctk.CTkFrame(v, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew")
        make_label(head, "Schedules, pipelines and folder watchers run while Finprime's Engine is open (tray keeps it alive)", 12, muted=True).pack(side="left")
        tabs = ctk.CTkTabview(v, fg_color=COLOR_CARD_SURFACE, border_width=1, border_color=COLOR_BORDER,
                              segmented_button_selected_color=COLOR_ACCENT, segmented_button_selected_hover_color=COLOR_ACCENT_HOVER,
                              segmented_button_unselected_color=COLOR_SIDEBAR, segmented_button_fg_color=COLOR_SIDEBAR,
                              text_color=COLOR_TEXT_MAIN)
        tabs.grid(row=1, column=0, sticky="nsew", pady=10)
        t_s, t_p, t_w = tabs.add("Schedules"), tabs.add("Pipelines"), tabs.add("Folder watchers")

        bar = ctk.CTkFrame(t_s, fg_color="transparent")
        bar.pack(fill="x", pady=(0, 8))
        make_button(bar, "＋ New schedule", self.add_schedule, "success", 140).pack(side="left")
        make_button(bar, "✎ Edit", self.edit_schedule, "ghost", 70).pack(side="left", padx=6)
        make_button(bar, "⏯ Enable / Disable", self.toggle_schedule, "ghost", 150).pack(side="left")
        make_button(bar, "▶ Run now", self.run_schedule_now, "ghost", 100).pack(side="left", padx=6)
        make_button(bar, "🗑 Delete", self.delete_schedule, "danger", 90).pack(side="right")
        self.sched_tree = self.make_tree(t_s, [("target", "Runs", 300), ("when", "When", 200), ("next", "Next run", 180),
                                                ("on", "Enabled", 80), ("last", "Last fired", 180)])

        bar = ctk.CTkFrame(t_p, fg_color="transparent")
        bar.pack(fill="x", pady=(0, 8))
        make_button(bar, "＋ New pipeline", self.add_pipeline, "success", 140).pack(side="left")
        make_button(bar, "✎ Edit", self.edit_pipeline, "ghost", 70).pack(side="left", padx=6)
        make_button(bar, "▶ Run pipeline", self.run_selected_pipeline, "ghost", 130).pack(side="left")
        make_button(bar, "🗑 Delete", self.delete_pipeline, "danger", 90).pack(side="right")
        self.pipe_tree = self.make_tree(t_p, [("name", "Pipeline", 240), ("steps", "Steps", 620), ("stop", "Stops on failure", 130)])

        bar = ctk.CTkFrame(t_w, fg_color="transparent")
        bar.pack(fill="x", pady=(0, 8))
        make_button(bar, "＋ Watch a folder", self.add_watcher, "success", 150).pack(side="left")
        make_button(bar, "⏯ Enable / Disable", self.toggle_watcher, "ghost", 150).pack(side="left", padx=6)
        make_button(bar, "🗑 Delete", self.delete_watcher, "danger", 90).pack(side="right")
        self.watch_tree = self.make_tree(t_w, [("folder", "Folder", 380), ("pattern", "Pattern", 100), ("target", "Runs", 240),
                                                ("on", "Enabled", 80), ("fired", "Triggered", 90)])

    def schedule_targets(self) -> list:
        t = [(f, f"script:{f}") for f in self.list_script_files()]
        t += [(f"{p['name']}  (pipeline)", f"pipe:{p['name']}") for p in self.db.pipelines()]
        return t

    def refresh_auto(self):
        if not self.ready:
            return
        self.sched_tree.delete(*self.sched_tree.get_children())
        for s in self.db.schedules():
            nr = next_run(s)
            self.sched_tree.insert("", "end", iid=str(s["id"]), tags=() if s["enabled"] else ("off",),
                                   values=(self.target_label(s["target"]), describe_schedule(s),
                                           nr.strftime("%a %d %b %H:%M") if nr else "—", "on" if s["enabled"] else "off",
                                           s["last_fire"] or "—"))
        self.pipe_tree.delete(*self.pipe_tree.get_children())
        for p in self.db.pipelines():
            self.pipe_tree.insert("", "end", iid=p["name"], values=(p["name"], "  →  ".join(p["steps"]),
                                                                     "yes" if p["stop_on_fail"] else "no"))
        self.watch_tree.delete(*self.watch_tree.get_children())
        for w in self.db.watchers():
            self.watch_tree.insert("", "end", iid=str(w["id"]), tags=() if w["enabled"] else ("off",),
                                   values=(w["folder"], w["pattern"], self.target_label(w["target"]),
                                           "on" if w["enabled"] else "off", w["fired"]))

    @staticmethod
    def _selected(tree):
        sel = tree.selection()
        return sel[0] if sel else None

    def add_schedule(self):
        res = ask_schedule(self, self.schedule_targets())
        if res:
            self.db.add_schedule(res)
            self.refresh_auto()
            self.refresh_scripts(force=True)
            self.toast("Schedule saved", "success", 1600)
        elif not self.schedule_targets():
            self.toast("Create a script first - there is nothing to schedule yet.", "info")

    def edit_schedule(self):
        sid = self._selected(self.sched_tree)
        if not sid:
            self.toast("Select a schedule first.", "info", 1500)
            return
        existing = next((s for s in self.db.schedules() if str(s["id"]) == sid), None)
        res = ask_schedule(self, self.schedule_targets(), existing)
        if res:
            self.db.update_schedule(int(sid), res)
            self.refresh_auto()
            self.refresh_scripts(force=True)

    def toggle_schedule(self):
        sid = self._selected(self.sched_tree)
        if sid:
            cur = next((s for s in self.db.schedules() if str(s["id"]) == sid), None)
            if cur:
                self.db.toggle_schedule(int(sid), not cur["enabled"])
                self.refresh_auto()
                self.refresh_scripts(force=True)

    def run_schedule_now(self):
        sid = self._selected(self.sched_tree)
        cur = next((s for s in self.db.schedules() if str(s["id"]) == sid), None) if sid else None
        if cur:
            self.fire_target(cur["target"], "manual")

    def delete_schedule(self):
        sid = self._selected(self.sched_tree)
        if sid and ask_confirm(self, "Delete schedule?", "This schedule will stop running.", "Delete", danger=True):
            self.db.delete_schedule(int(sid))
            self.refresh_auto()
            self.refresh_scripts(force=True)

    def add_pipeline(self):
        scripts = self.list_script_files()
        if not scripts:
            self.toast("Create some scripts first.", "info")
            return
        res = ask_pipeline(self, scripts)
        if res:
            self.db.save_pipeline(res["name"], res["steps"], res["stop_on_fail"])
            self.refresh_auto()

    def edit_pipeline(self):
        name = self._selected(self.pipe_tree)
        p = self.db.get_pipeline(name) if name else None
        if not p:
            self.toast("Select a pipeline first.", "info", 1500)
            return
        res = ask_pipeline(self, self.list_script_files(), p)
        if res:
            self.db.save_pipeline(res["name"], res["steps"], res["stop_on_fail"], old_name=name)
            self.refresh_auto()

    def run_selected_pipeline(self):
        name = self._selected(self.pipe_tree)
        if name:
            self.run_pipeline(name)
        else:
            self.toast("Select a pipeline first.", "info", 1500)

    def delete_pipeline(self):
        name = self._selected(self.pipe_tree)
        if name and ask_confirm(self, "Delete pipeline?", f"'{name}' will be removed. Schedules that use it stop working.",
                                "Delete", danger=True):
            self.db.delete_pipeline(name)
            self.refresh_auto()

    def add_watcher(self):
        res = ask_watcher(self, self.schedule_targets())
        if res:
            self.db.add_watcher(res)
            self.refresh_auto()
            self.toast("Watching. Files already in the folder are ignored - only new ones trigger a run.", "success", 4500)

    def toggle_watcher(self):
        wid = self._selected(self.watch_tree)
        cur = next((w for w in self.db.watchers() if str(w["id"]) == wid), None) if wid else None
        if cur:
            self.db.toggle_watcher(cur["id"], not cur["enabled"])
            self.watcher.seen.pop(cur["id"], None)          # re-baseline when re-enabled
            self.refresh_auto()

    def delete_watcher(self):
        wid = self._selected(self.watch_tree)
        if wid and ask_confirm(self, "Delete watcher?", "This folder will no longer be watched.", "Delete", danger=True):
            self.db.delete_watcher(int(wid))
            self.refresh_auto()

    # --------------------------------------------------------------- ENV HEALTH
    PACKAGES = [
        ("customtkinter", "customtkinter", "UI framework", True),
        ("pandas", "pandas", "data processing + output preview", True),
        ("openpyxl", "openpyxl", "read / write Excel files", True),
        ("Pillow", "PIL", "images (login artwork, tray icon)", True),
        ("pywin32", "win32com", "run VBA macros through Excel", False),
        ("send2trash", "send2trash", "deletes go to the Recycle Bin", False),
        ("keyring", "keyring", "saves your AI API key securely", False),
        ("pystray", "pystray", "minimize to system tray", False),
        ("plyer", "plyer", "desktop notifications when long runs finish", False),
        ("tkinterdnd2", "tkinterdnd2", "drag & drop files onto scripts", False),
        ("ruff", "ruff", "fast linter + formatter", False),
        ("black", "black", "code formatter", False),
        ("pyflakes", "pyflakes", "linter", False),
    ]

    def build_health(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(3, weight=1)
        head = ctk.CTkFrame(v, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew")
        make_label(head, "🩺 Environment & Dependency Health Checker", 17, True).pack(side="left")
        make_button(head, "🔄 Recheck Packages", self.load_health_checks, "ghost", 160).pack(side="right")
        self.health_info = make_label(v, "", 11, muted=True, justify="left")
        self.health_info.grid(row=1, column=0, sticky="w", pady=(4, 8))
        self.health_rows = ctk.CTkScrollableFrame(v, fg_color=COLOR_CARD_SURFACE, corner_radius=12, border_width=1,
                                                  border_color=COLOR_BORDER, height=300)
        self.health_rows.grid(row=2, column=0, sticky="ew")
        card = ctk.CTkFrame(v, fg_color=COLOR_CARD_SURFACE, corner_radius=12, border_width=1, border_color=COLOR_BORDER)
        card.grid(row=3, column=0, sticky="nsew", pady=(12, 0))
        bar = ctk.CTkFrame(card, fg_color="transparent")
        bar.pack(fill="x", padx=14, pady=(12, 6))
        make_label(bar, "📦 Script dependency scanner", 14, True).pack(side="left")
        self.btn_install_all = make_button(bar, "⬇ Install all missing", self.install_all_missing, "primary", 170, state="disabled")
        self.btn_install_all.pack(side="right")
        make_button(bar, "🔍 Scan all scripts", self.scan_all_scripts, "outline", 150).pack(side="right", padx=8)
        self.scan_console = make_console_box(card)
        self.scan_console.pack(fill="both", expand=True, padx=14, pady=(0, 14))
        self.missing_pkgs: list = []

    def load_health_checks(self):
        if not self.ready:
            return
        mode = "" if own_interpreter() else "  (custom interpreter - package versions are only shown for the dashboard's own Python)"
        self.health_info.configure(text=f"Dashboard: Python {sys.version.split()[0]} · {sys.executable}\n"
                                        f"Scripts run with: {get_python()}{mode}")
        for w in self.health_rows.winfo_children():
            w.destroy()
        make_label(self.health_rows, "Checking packages…", 12, muted=True).pack(pady=20)
        imports = [p[1] for p in self.PACKAGES]

        def done(kind, res):
            for w in self.health_rows.winfo_children():
                w.destroy()
            if kind == "err":
                make_label(self.health_rows, f"Check failed: {res}", 12).pack(pady=20)
                return
            for dist, imp, desc, required in self.PACKAGES:
                ok = res.get(imp, False)
                row = ctk.CTkFrame(self.health_rows, fg_color=COLOR_BG_MAIN, corner_radius=8, border_width=1, border_color=COLOR_BORDER)
                row.pack(fill="x", padx=6, pady=4)
                ctk.CTkLabel(row, text="", image=icon_tile("chip", 30)).pack(side="left", padx=(12, 10), pady=8)
                make_label(row, dist, 13, True).pack(side="left")
                make_label(row, f"   {'required' if required else 'optional'} · {desc}", 11, muted=True).pack(side="left")
                if ok:
                    ver = package_version(dist) if own_interpreter() else None
                    ctk.CTkLabel(row, text=f"Installed{f' ({ver})' if ver else ''}", text_color="#3fb950", fg_color="#113c23",
                                 corner_radius=6, padx=10, pady=3, font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold")
                                 ).pack(side="right", padx=12)
                else:
                    ctk.CTkLabel(row, text="Missing" if required else "Not installed", text_color=COLOR_DANGER_HOVER if required else COLOR_WARN,
                                 fg_color="#4a1b18" if required else "#3a2c0b", corner_radius=6, padx=10, pady=3,
                                 font=ctk.CTkFont(family=FONT_FAMILY, size=11, weight="bold")).pack(side="right", padx=12)
                    make_button(row, "⬇ Install", lambda d=dist: self.install_packages([d], self.scan_console, self.load_health_checks),
                                "outline", 90, height=28).pack(side="right")

        run_bg(self, lambda: check_modules(imports), done)

    def install_packages(self, pkgs: list, box, after=None):
        console_write(box, f"\nInstalling: {', '.join(pkgs)}\n", "info")

        def emit(line):
            self.events.put(("call", lambda l=line: console_write(box, l, "dim")))

        def done(kind, val):
            ok = kind == "ok" and val == 0
            console_write(box, "✅ Install complete.\n" if ok else "❌ pip reported errors (see above).\n", "success" if ok else "error")
            if after:
                after()

        run_bg(self, lambda: pip_install(pkgs, emit), done)

    def scan_all_scripts(self):
        console_write(self.scan_console, f"\n[{datetime.now():%H:%M:%S}] Scanning workspace scripts…\n", "info")

        def work():
            ws, per = workspace_dir(), {}
            for f in self.list_script_files():
                if f.lower().endswith(".py"):
                    try:
                        mods = scan_imports((ws / f).read_text(encoding="utf-8", errors="replace"), ws)
                    except OSError:
                        continue
                    if mods:
                        per[f] = sorted(mods)
            allm = sorted({m for ms in per.values() for m in ms})
            return per, check_modules(allm)

        def done(kind, val):
            if kind == "err":
                console_write(self.scan_console, f"❌ {val}\n", "error")
                return
            per, status = val
            missing = set()
            if not per:
                console_write(self.scan_console, "✅ No third-party imports found in any script.\n", "success")
            for f, mods in per.items():
                bad = [m for m in mods if not status.get(m)]
                if bad:
                    console_write(self.scan_console, f"✗ {f}: missing {', '.join(pip_name(m) for m in bad)}\n", "error")
                    missing.update(pip_name(m) for m in bad)
                else:
                    console_write(self.scan_console, f"✓ {f}: all {len(mods)} imports available\n", "success")
            self.missing_pkgs = sorted(missing)
            self.btn_install_all.configure(state="normal" if missing else "disabled")
            if missing:
                console_write(self.scan_console, f"→ pip install {' '.join(self.missing_pkgs)}\n", "warning")

        run_bg(self, work, done)

    def install_all_missing(self):
        if self.missing_pkgs:
            self.btn_install_all.configure(state="disabled")
            self.install_packages(list(self.missing_pkgs), self.scan_console, self.scan_all_scripts)

    # ----------------------------------------------------------------- SETTINGS
    def build_settings(self, v):
        v.grid_columnconfigure(0, weight=1)
        v.grid_rowconfigure(1, weight=1)
        s = ctk.CTkScrollableFrame(v, fg_color="transparent")
        s.grid(row=1, column=0, sticky="nsew", pady=(0, 8))

        def section(title, note=""):
            f = ctk.CTkFrame(s, fg_color=COLOR_CARD_SURFACE, corner_radius=16, border_width=1, border_color=COLOR_BORDER)
            f.pack(fill="x", pady=8, padx=4)
            ic_name, plain = split_icon(title)
            hd = ctk.CTkFrame(f, fg_color="transparent")
            hd.pack(fill="x", padx=18, pady=(14, 2))
            if ic_name:
                ctk.CTkLabel(hd, text="", image=icon_tile(ic_name, 30)).pack(side="left")
            make_label(hd, plain, 16, True, display=True).pack(side="left", padx=10)
            if note:
                make_label(f, note, 11, muted=True, wraplength=900, justify="left").pack(anchor="w", padx=18)
            body = ctk.CTkFrame(f, fg_color="transparent")
            body.pack(fill="x", padx=18, pady=(8, 14))
            return body

        def row(parent):
            r = ctk.CTkFrame(parent, fg_color="transparent")
            r.pack(fill="x", pady=4)
            return r

        # workspace
        b = section("📁 Workspace", "Scripts, macros, outputs and the login artwork (login_bg.png) live here.")
        r = row(b)
        self.set_ws = make_entry(r, width=520)
        self.set_ws.pack(side="left")
        self.set_ws.insert(0, str(CFG["workspace"]))
        make_button(r, "Browse…", self.browse_workspace, "ghost", 90).pack(side="left", padx=8)
        make_button(r, "Apply", self.apply_workspace, "primary", 90).pack(side="left")

        # appearance
        b = section("🎨 Appearance")
        r = row(b)
        make_label(r, "Theme", 12, True).pack(side="left", padx=(0, 12))
        ctk.CTkOptionMenu(r, values=list(THEMES), width=180, fg_color=COLOR_BG_MAIN, button_color=COLOR_BORDER,
                          command=self.change_theme).pack(side="left")
        self.theme_menu_value = CFG["theme"]
        for child in r.winfo_children():
            if isinstance(child, ctk.CTkOptionMenu):
                child.set(CFG["theme"])
        var = tk.BooleanVar(value=bool(CFG["glitch"]))
        r = row(b)
        ctk.CTkSwitch(r, text="Glitch page transitions", variable=var, command=lambda: self.set_cfg("glitch", var.get())).pack(side="left")
        seg = ctk.CTkSegmentedButton(r, values=list(GLITCH_LEVELS), selected_color=COLOR_ACCENT, selected_hover_color=COLOR_ACCENT_HOVER,
                                     unselected_color=COLOR_BG_MAIN, text_color="#0A0E17" if False else COLOR_TEXT_MAIN,
                                     command=lambda v_: self.set_cfg("glitch_level", v_))
        seg.set(CFG["glitch_level"])
        seg.pack(side="left", padx=18)
        make_button(r, "▶ Preview", self.preview_glitch, "outline", 110, 32).pack(side="left")
        for key, text in (("matrix_rain", "Matrix rain on the login screen"),
                          ("themed_text", "Warhammer-flavoured status text (Litany complete / Rite failed…)"),
                          ("sound", "Play a sound when a run finishes")):
            var = tk.BooleanVar(value=bool(CFG[key]))
            ctk.CTkSwitch(b, text=text, variable=var, command=lambda k=key, vv=var: self.set_cfg(k, vv.get())).pack(anchor="w", pady=4)

        # runtime
        b = section("🚀 Runtime", "Leave the interpreter blank to use the one running this app. Set it when you package the app "
                                   "with PyInstaller or want scripts to use another environment.")
        r = row(b)
        make_label(r, "Python interpreter", 12, True).pack(side="left", padx=(0, 12))
        self.set_py = make_entry(r, "automatic", width=420)
        self.set_py.pack(side="left")
        if CFG["python_path"]:
            self.set_py.insert(0, CFG["python_path"])
        make_button(r, "Browse…", lambda: (lambda p: (self.set_py.delete(0, "end"), self.set_py.insert(0, p)) if p else None)(
            filedialog.askopenfilename(title="Pick python.exe")), "ghost", 90).pack(side="left", padx=8)
        r = row(b)
        make_label(r, "Kill scripts after (seconds, 0 = never)", 12, True).pack(side="left", padx=(0, 12))
        self.set_timeout = make_entry(r, width=90)
        self.set_timeout.pack(side="left")
        self.set_timeout.insert(0, str(CFG["run_timeout_s"]))
        make_label(r, "Notify when a run takes longer than (s)", 12, True).pack(side="left", padx=(28, 12))
        self.set_notify = make_entry(r, width=90)
        self.set_notify.pack(side="left")
        self.set_notify.insert(0, str(CFG["notify_after_s"]))
        make_button(r, "Save", self.save_runtime, "primary", 80).pack(side="left", padx=16)

        # security
        b = section("🔒 Security", "Auto-lock re-shows the login screen after inactivity. Running jobs and schedules keep going.")
        r = row(b)
        make_label(r, "Auto-lock after (minutes, 0 = off)", 12, True).pack(side="left", padx=(0, 12))
        self.set_lock = make_entry(r, width=90)
        self.set_lock.pack(side="left")
        self.set_lock.insert(0, str(CFG["auto_lock_minutes"]))
        make_button(r, "Save", self.save_lock, "primary", 80).pack(side="left", padx=12)
        var = tk.BooleanVar(value=bool(CFG["minimize_to_tray"]))
        ctk.CTkSwitch(b, text="Closing the window minimizes to the system tray (needs pystray)", variable=var,
                      command=lambda: self.set_tray(var.get())).pack(anchor="w", pady=6)
        make_label(b, "Change master password", 12, True).pack(anchor="w", pady=(8, 2))
        r = row(b)
        self.pw_old = make_entry(r, "current", width=180, show="•")
        self.pw_old.pack(side="left")
        self.pw_new = make_entry(r, "new (6+ chars)", width=180, show="•")
        self.pw_new.pack(side="left", padx=8)
        self.pw_new2 = make_entry(r, "repeat new", width=180, show="•")
        self.pw_new2.pack(side="left")
        make_button(r, "Change", self.change_password, "primary", 90).pack(side="left", padx=8)

        # AI
        b = section("✨ AI assistant (Claude API)", "Used by the ✨ AI Assist button in Script Builder. Code you send is processed by "
                                                    "Anthropic's API.")
        r = row(b)
        make_label(r, "Model", 12, True).pack(side="left", padx=(0, 12))
        self.set_model = make_entry(r, width=240)
        self.set_model.pack(side="left")
        self.set_model.insert(0, CFG["ai_model"])
        make_button(r, "Save model", self.save_model, "ghost", 110).pack(side="left", padx=8)
        r = row(b)
        make_label(r, "API key", 12, True).pack(side="left", padx=(0, 20))
        self.set_key = make_entry(r, "sk-ant-…", width=340, show="•")
        self.set_key.pack(side="left")
        make_button(r, "Save securely", lambda: self.save_key(True), "primary", 120).pack(side="left", padx=8)
        make_button(r, "This session only", lambda: self.save_key(False), "ghost", 140).pack(side="left")
        self.key_status = make_label(b, "", 11, muted=True)
        self.key_status.pack(anchor="w")

        # data
        b = section("🧰 Data & maintenance")
        r = row(b)
        make_button(r, "📂 Open config folder", lambda: open_path(CONFIG_DIR), "ghost", 170).pack(side="left")
        make_button(r, "📂 Open workspace", lambda: open_path(workspace_dir()), "ghost", 160).pack(side="left", padx=8)
        make_button(r, "🗑 Clear all run history", self._clear_all_history, "danger", 190).pack(side="left")
        make_label(b, f"{APP_NAME} v{APP_VERSION}  ·  data: {CONFIG_DIR}", 11, muted=True).pack(anchor="w", pady=(8, 0))
        self.refresh_settings()

    def refresh_settings(self):
        if not self.ready:
            return
        has = bool(get_api_key())
        self.key_status.configure(text="✔ An API key is available." if has else "No API key set yet.",
                                  text_color="#3fb950" if has else COLOR_TEXT_MUTED)

    # ---- settings actions ------------------------------------------------------
    def set_cfg(self, key, value):
        CFG[key] = value
        CFG.save()
        if key == "themed_text":
            self.refresh_scripts(force=True)

    def preview_glitch(self):
        other = "cmd" if self.current_view != "cmd" else "home"
        cur = self.current_view
        self.select_view(other)
        self.after(1200, lambda: self.select_view(cur))

    def browse_workspace(self):
        p = filedialog.askdirectory(title="Choose workspace folder")
        if p:
            self.set_ws.delete(0, "end")
            self.set_ws.insert(0, p)

    def apply_workspace(self):
        p = Path(self.set_ws.get().strip())
        try:
            p.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            self.toast(f"Cannot use that folder: {exc}", "error")
            return
        CFG["workspace"] = str(p)
        CFG.save()
        self.param_cache.clear()
        self.refresh_scripts(force=True)
        self.toast(f"Workspace is now {p}", "success")

    def change_theme(self, name):
        CFG["theme"] = name
        CFG.save()
        apply_theme(name)
        self.refresh_application()

    def save_runtime(self):
        try:
            t, n = int(self.set_timeout.get() or 0), int(self.set_notify.get() or 0)
            assert t >= 0 and n >= 0
        except Exception:  # noqa: BLE001
            self.toast("Timeout and notify values must be whole numbers ≥ 0.", "warning")
            return
        py = self.set_py.get().strip()
        if py and not Path(py).exists():
            self.toast("That interpreter path does not exist.", "warning")
            return
        CFG["python_path"], CFG["run_timeout_s"], CFG["notify_after_s"] = py, t, n
        CFG.save()
        self.toast("Runtime settings saved", "success", 1600)

    def save_lock(self):
        try:
            m = int(self.set_lock.get() or 0)
            assert m >= 0
        except Exception:  # noqa: BLE001
            self.toast("Enter minutes as a whole number (0 disables auto-lock).", "warning")
            return
        CFG["auto_lock_minutes"] = m
        CFG.save()
        self.toast("Auto-lock saved", "success", 1600)

    def set_tray(self, on: bool):
        CFG["minimize_to_tray"] = on
        CFG.save()
        if on:
            if not self.setup_tray():
                self.toast("System tray needs 'pystray' + 'Pillow' - install pystray from Env Health Check.", "warning", 5000)

    def change_password(self):
        old, new, new2 = self.pw_old.get(), self.pw_new.get(), self.pw_new2.get()
        if not check_password(old):
            self.toast("Current password is incorrect.", "error")
            return
        if len(new) < 6 or new != new2:
            self.toast("New password must match and be at least 6 characters.", "warning")
            return
        save_auth(new)
        for e in (self.pw_old, self.pw_new, self.pw_new2):
            e.delete(0, "end")
        self.toast("Master password changed", "success")

    def save_model(self):
        CFG["ai_model"] = self.set_model.get().strip() or CFG.DEFAULTS["ai_model"]
        CFG.save()
        self.toast("Model saved", "success", 1400)

    def save_key(self, persist: bool):
        key = self.set_key.get().strip()
        if not key:
            self.toast("Paste an API key first.", "warning")
            return
        try:
            set_api_key(key, persist)
            self.toast("API key saved securely" if persist and try_import("keyring") else "API key kept for this session", "success")
        except Exception as exc:  # noqa: BLE001
            set_api_key(key, False)
            self.toast(f"Could not store it securely ({exc}); using it for this session only.", "warning", 5000)
        self.set_key.delete(0, "end")
        self.refresh_settings()

    def _clear_all_history(self):
        if ask_confirm(self, "Clear ALL history?", "Every recorded run will be deleted.", "Clear", danger=True):
            self.db.clear_runs()
            self.refresh_scripts(force=True)
            self.toast("History cleared", "info")


# =============================================================================
#  MAIN WINDOW
# =============================================================================
if HAS_DND:
    class BaseRoot(ctk.CTk, TkinterDnD.DnDWrapper):          # type: ignore[misc,name-defined]
        def __init__(self):
            super().__init__()
            self.dnd_ok = False
            try:
                self.TkdndVersion = TkinterDnD._require(self)
                self.dnd_ok = True
            except Exception as exc:  # noqa: BLE001
                print(f"[DND] drag & drop unavailable: {exc}")
else:
    class BaseRoot(ctk.CTk):                                  # type: ignore[no-redef]
        dnd_ok = False


class FinprimesEngineDashboard(ViewsMixin, BaseRoot):
    NAV_GROUPS = [
        ("MAIN", [("home", "home", "Overview"), ("cmd", "rocket", "Command Center")]),
        ("BUILD", [("builder", "code", "Script Builder"), ("test", "flask", "Test Run Sandbox")]),
        ("OPERATE", [("history", "clock", "Run History"), ("auto", "timer", "Automation")]),
        ("SYSTEM", [("health", "pulse", "Env Health Check"), ("settings", "gear", "Settings")]),
    ]
    NAV = [(k, title) for _g, items in NAV_GROUPS for k, _i, title in items]
    PAGE_INFO = {
        "home": ("home", "Overview", "Live stats across all of your automations"),
        "cmd": ("rocket", "Command Center", "Launch, monitor and manage your saved scripts"),
        "builder": ("code", "Script Builder", "Write Python and VBA in a full-featured editor"),
        "test": ("flask", "Test Run Sandbox", "Validate, lint and test-run safely before you save"),
        "history": ("clock", "Run History", "Every run with its complete log and output files"),
        "auto": ("timer", "Automation", "Schedules, pipelines and folder watchers"),
        "health": ("pulse", "Env Health Check", "Interpreter and dependency diagnostics"),
        "settings": ("gear", "Settings", "Workspace, appearance, security and AI"),
    }

    def __init__(self):
        super().__init__()
        init_fonts()
        self.title(APP_NAME)
        self.configure(fg_color=COLOR_BG_MAIN)
        self.geometry("1440x860")
        self.minsize(1180, 780)
        self.set_window_icon()
        self.after(0, self._maximize)
        self.after(350, self.style_titlebar)
        migrate_legacy_auth()
        setup_ttk_style()

        self.db = Database(DB_FILE)
        self.events: queue.Queue = queue.Queue()
        self.runner = ScriptRunner(self.events.put, self.db)
        self.watcher = FolderWatcher()
        self.toasts = ToastManager(self)

        self.ready = False
        self.locked = True
        self.overlay = None
        self.tray = None
        self.glitch = None
        self.current_view = "home"
        self.nav_items: dict = {}
        self.script_status: dict = {}
        self.pipes_running: set = set()
        self.meta: dict = {}
        self.last_runs: dict = {}
        self.next_by_target: dict = {}
        self.cards: dict = {}
        self._sig = None
        self.param_cache: dict = {}
        self.editor_mode = "python"
        self.editor_original = ""
        self.editor_orig_name = ""
        self.editor_file = None
        self.last_error_text = ""
        self.last_activity = time.time()
        self._focus_next = None
        self._timers_started = False
        self._pulse_on = False

        for seq in ("<Any-KeyPress>", "<Any-ButtonPress>", "<Motion>"):
            self.bind_all(seq, self._touch, add="+")
        self.bind_all("<Control-k>", self.open_palette)
        self.bind_all("<Control-s>", self.on_ctrl_s)
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        self.show_lock(animate=False)
        self.after(60, self.pump_events)

    # ---------------------------------------------------------------- basics
    def _maximize(self):
        try:
            self.state("zoomed")
        except tk.TclError:
            pass

    def set_window_icon(self):
        try:
            if IS_WIN:
                ico = CONFIG_DIR / f"app_{CFG['theme'].replace(' ', '_')}.ico"
                if not ico.exists():
                    make_app_icon_file(ico)
                self.iconbitmap(str(ico))
            elif ImageTk is not None:
                self._icon_ref = ImageTk.PhotoImage(render_tile("bolt", 64, COLOR_ACCENT, COLOR_ACCENT2), master=self)
                self.iconphoto(True, self._icon_ref)
        except Exception as exc:  # noqa: BLE001
            print(f"[icon] {exc}")

    def style_titlebar(self):
        """Windows 10/11: dark title bar tinted to match the sidebar (silently ignored elsewhere)."""
        if not IS_WIN:
            return
        try:
            import ctypes
            hwnd = ctypes.windll.user32.GetParent(self.winfo_id())
            dwm = ctypes.windll.dwmapi

            def put(attr, value):
                v = ctypes.c_int(value)
                dwm.DwmSetWindowAttribute(hwnd, attr, ctypes.byref(v), ctypes.sizeof(v))

            def bgr(hexcol):
                r, g, b = hex_rgb(hexcol)
                return r | (g << 8) | (b << 16)

            put(20, 1)
            put(35, bgr(COLOR_SIDEBAR))
            put(36, bgr(COLOR_TEXT_MAIN))
            put(34, bgr(COLOR_BORDER))
        except Exception as exc:  # noqa: BLE001
            print(f"[titlebar] {exc}")

    def _touch(self, _event=None):
        self.last_activity = time.time()

    def toast(self, message: str, kind: str = "info", ms: int = 3400):
        if not self.locked:
            self.toasts.show(message, kind, ms)

    # ------------------------------------------------------ glitch plumbing
    def can_glitch(self) -> bool:
        try:
            if not (CFG["glitch"] and GlitchTransition.available() and self.winfo_viewable()):
                return False
            if self.state() == "iconic":
                return False
            if any(isinstance(w, ctk.CTkToplevel) for w in self.winfo_children()):
                return False
            if IS_WIN and self.focus_displayof() is None:
                return False
            return True
        except tk.TclError:
            return False

    def run_glitch(self, host, switch_fn, caption: str) -> bool:
        if not self.can_glitch():
            return False
        if self.glitch is not None and self.glitch.active:
            self.glitch.abort()
        tr = GlitchTransition(host, CFG["glitch_level"], caption)
        if tr.start(switch_fn):
            self.glitch = tr
            return True
        return False

    # ------------------------------------------------------------- lock / login
    def show_lock(self, animate: bool = True):
        if self.overlay is not None:
            return

        def build():
            self.locked = True
            for w in self.winfo_children():                  # close dialogs so nothing leaks past the lock
                if isinstance(w, ctk.CTkToplevel):
                    w.destroy()
            self.overlay = LoginOverlay(self, self.on_login_success)
            self.overlay.place(x=0, y=0, relwidth=1, relheight=1)
            self.overlay.tkraise()

        if animate and self.ready and self.run_glitch(self, build, "> session.lock --now"):
            return
        build()

    def on_login_success(self):
        def unlock():
            ov, self.overlay = self.overlay, None
            self.locked = False
            self.last_activity = time.time()
            if ov is not None:
                ov.destroy()
            if not self.ready:
                self.build_dashboard()

        if self.run_glitch_unlock(unlock):
            return
        unlock()

    def run_glitch_unlock(self, unlock) -> bool:
        if not (CFG["glitch"] and GlitchTransition.available() and self.winfo_viewable()):
            return False
        if self.glitch is not None and self.glitch.active:
            self.glitch.abort()
        tr = GlitchTransition(self, CFG["glitch_level"], "> session.unlock --ok")
        if tr.start(unlock):
            self.glitch = tr
            return True
        return False

    def lock_now(self):
        if self.ready and not self.locked:
            self.show_lock()

    def tick_idle(self):
        try:
            minutes = int(CFG["auto_lock_minutes"] or 0)
            if minutes > 0 and self.ready and not self.locked and time.time() - self.last_activity > minutes * 60:
                self.show_lock()
        finally:
            self.after(15000, self.tick_idle)

    # -------------------------------------------------------------- dashboard
    def build_dashboard(self):
        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self.setup_sidebar()
        self.setup_content()
        self.ready = True
        self.refresh_scripts(force=True)
        self._switch_view("home")
        self.restore_draft()
        if not self._timers_started:
            self._timers_started = True
            self.after(20000, self.tick_schedules)
            self.after(6000, self.tick_watchers)
            self.after(15000, self.tick_idle)
            self.after(20000, self.autosave_draft)
            self.after(1000, self.tick_clock)
            self.after(600, self.pulse_jobs)
        if CFG["minimize_to_tray"]:
            self.setup_tray()

    def setup_sidebar(self):
        self.sidebar = ctk.CTkFrame(self, width=252, corner_radius=0, fg_color=COLOR_SIDEBAR)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.pack_propagate(False)
        tk.Frame(self.sidebar, width=1, bg=COLOR_BORDER).pack(side="right", fill="y")
        brand = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        brand.pack(fill="x", padx=16, pady=(22, 14))
        ctk.CTkLabel(brand, text="", image=icon_tile("bolt", 44, (COLOR_ACCENT, COLOR_ACCENT2))).pack(side="left")
        txt = ctk.CTkFrame(brand, fg_color="transparent")
        txt.pack(side="left", padx=12)
        wm = gradient_wordmark("FINPRIME'S", 21, COLOR_ACCENT, COLOR_ACCENT2, spacing=1)
        if wm is not None:
            ctk.CTkLabel(txt, text="", image=wm, height=26).pack(anchor="w")
        else:
            make_label(txt, "FINPRIME'S", 18, True, text_color=COLOR_ACCENT).pack(anchor="w")
        ctk.CTkLabel(txt, text=f"AUTOMATION ENGINE  ·  v{APP_VERSION}", text_color=COLOR_TEXT_MUTED,
                     font=ctk.CTkFont(family=FONT_MONO, size=9, weight="bold")).pack(anchor="w")
        self.nav_items = {}
        for group, items in self.NAV_GROUPS:
            ctk.CTkLabel(self.sidebar, text=group, anchor="w", text_color="#4A5680",
                         font=ctk.CTkFont(family=FONT_MONO, size=10, weight="bold")).pack(fill="x", padx=26, pady=(12, 2))
            for key, ic, label in items:
                item = NavItem(self.sidebar, ic, label, lambda k=key: self.select_view(k))
                item.pack(fill="x", padx=10, pady=1)
                self.nav_items[key] = item
        make_button(self.sidebar, "📁  Workspace Folder", self.open_workspace, "ghost", height=38).pack(
            side="bottom", fill="x", padx=14, pady=(6, 16))
        make_button(self.sidebar, "🔄  Refresh Engine", self.refresh_application, "outline", height=38).pack(
            side="bottom", fill="x", padx=14, pady=6)
        sysc = ctk.CTkFrame(self.sidebar, fg_color=blend(COLOR_SIDEBAR, "#FFFFFF", 0.04), corner_radius=12)
        sysc.pack(side="bottom", fill="x", padx=14, pady=(6, 4))
        row = ctk.CTkFrame(sysc, fg_color="transparent")
        row.pack(fill="x", padx=12, pady=(10, 0))
        self.sys_dot = ctk.CTkLabel(row, text="●", text_color="#4ADE80", font=ctk.CTkFont(size=12), width=14)
        self.sys_dot.pack(side="left")
        self.sys_state = make_label(row, "Engine idle", 12, True)
        self.sys_state.pack(side="left", padx=4)
        self.sys_info = ctk.CTkLabel(sysc, text="", text_color=COLOR_TEXT_MUTED, justify="left", anchor="w",
                                     font=ctk.CTkFont(family=FONT_MONO, size=10))
        self.sys_info.pack(fill="x", padx=14, pady=(2, 10))

    def setup_content(self):
        self.main = ctk.CTkFrame(self, corner_radius=0, fg_color=COLOR_BG_MAIN)
        self.main.grid(row=0, column=1, sticky="nsew")
        self.main.grid_rowconfigure(1, weight=1)
        self.main.grid_columnconfigure(0, weight=1)
        self.build_topbar()
        self.stage = ctk.CTkFrame(self.main, corner_radius=0, fg_color=COLOR_BG_MAIN)
        self.stage.grid(row=1, column=0, sticky="nsew")
        self.stage.grid_rowconfigure(0, weight=1)
        self.stage.grid_columnconfigure(0, weight=1)
        self.views = {}
        builders = {"home": self.build_home, "cmd": self.build_cmd, "builder": self.build_builder,
                    "test": self.build_test, "history": self.build_history, "auto": self.build_auto,
                    "health": self.build_health, "settings": self.build_settings}
        for key, fn in builders.items():
            frame = ctk.CTkFrame(self.stage, fg_color="transparent")
            frame.grid(row=0, column=0, sticky="nsew", padx=30, pady=(14, 20))
            self.views[key] = frame
            fn(frame)

    def build_topbar(self):
        bar = ctk.CTkFrame(self.main, fg_color=COLOR_BG_MAIN, corner_radius=0, height=72)
        bar.grid(row=0, column=0, sticky="ew")
        bar.pack_propagate(False)
        self.topbar = bar
        ctk.CTkFrame(bar, height=1, fg_color=COLOR_BORDER, corner_radius=0).place(relx=0, rely=1.0, relwidth=1, anchor="sw")
        left = ctk.CTkFrame(bar, fg_color="transparent")
        left.pack(side="left", padx=(30, 0), fill="y")
        self.tb_icon = ctk.CTkLabel(left, text="", image=icon_tile("home", 44))
        self.tb_icon.pack(side="left")
        titles = ctk.CTkFrame(left, fg_color="transparent")
        titles.pack(side="left", padx=14)
        self.tb_title = make_label(titles, "Overview", 21, True, display=True)
        self.tb_title.pack(anchor="w", pady=(14, 0))
        self.tb_sub = make_label(titles, "", 11, muted=True)
        self.tb_sub.pack(anchor="w")
        right = ctk.CTkFrame(bar, fg_color="transparent")
        right.pack(side="right", padx=(0, 26), fill="y")
        make_button(right, "", self.lock_now, "ghost", 38, 38, glyph="lock").pack(side="right", pady=17)
        clock = ctk.CTkFrame(right, fg_color="transparent")
        clock.pack(side="right", padx=16)
        self.clock_lbl = ctk.CTkLabel(clock, text="", text_color=COLOR_TEXT_MAIN, font=ctk.CTkFont(family=FONT_MONO, size=15, weight="bold"))
        self.clock_lbl.pack(anchor="e", pady=(15, 0))
        self.date_lbl = ctk.CTkLabel(clock, text="", text_color=COLOR_TEXT_MUTED, font=ctk.CTkFont(family=FONT_FAMILY, size=10))
        self.date_lbl.pack(anchor="e")
        jobs = ctk.CTkFrame(right, fg_color=blend(COLOR_BG_MAIN, "#FFFFFF", 0.05), corner_radius=18, height=36)
        jobs.pack(side="right", pady=18, padx=(0, 4))
        self.jobs_dot = ctk.CTkLabel(jobs, text="●", text_color="#4ADE80", width=14, font=ctk.CTkFont(size=12))
        self.jobs_dot.pack(side="left", padx=(12, 2), pady=6)
        self.jobs_lbl = make_label(jobs, "All quiet", 12, True)
        self.jobs_lbl.pack(side="left", padx=(0, 14))
        cmd = ctk.CTkFrame(right, fg_color=COLOR_CARD_SURFACE, corner_radius=18, border_width=1, border_color=COLOR_BORDER,
                           height=36, width=330)
        cmd.pack(side="right", padx=18, pady=18)
        cmd.pack_propagate(False)
        ctk.CTkLabel(cmd, text="", image=icon("search", 15)).pack(side="left", padx=(14, 6))
        make_label(cmd, "Search or run a command…", 12, muted=True).pack(side="left")
        key_cap(cmd, "Ctrl K").pack(side="right", padx=10)
        for w in [cmd] + list(cmd.winfo_children()):
            w.bind("<Button-1>", lambda e: self.open_palette(), add="+")
            try:
                w.configure(cursor="hand2")
            except tk.TclError:
                pass
        self.tick_clock_once()

    def tick_clock_once(self):
        now = datetime.now()
        self.clock_lbl.configure(text=now.strftime("%H:%M:%S"))
        self.date_lbl.configure(text=now.strftime("%a %d %b %Y"))

    def tick_clock(self):
        try:
            if self.ready:
                self.tick_clock_once()
        except tk.TclError:
            pass
        finally:
            self.after(1000, self.tick_clock)

    def running_count(self) -> int:
        return len(self.runner.running_keys()) + len(self.pipes_running)

    def update_jobs(self):
        if not self.ready:
            return
        n = self.running_count()
        self.jobs_lbl.configure(text=f"{n} running" if n else "All quiet")
        self.sys_state.configure(text=f"{n} job{'s' if n != 1 else ''} running" if n else "Engine idle")
        if not n:
            self.jobs_dot.configure(text_color="#4ADE80")
            self.sys_dot.configure(text_color="#4ADE80")
        self.update_sidebar_info()

    def update_sidebar_info(self):
        try:
            n = len(self.list_script_files())
            self.sys_info.configure(text=f"Python {sys.version.split()[0]}\n{n} scripts · {len(self.db.schedules())} schedules")
        except tk.TclError:
            pass

    def pulse_jobs(self):
        try:
            if self.ready and self.running_count():
                self._pulse_on = not self._pulse_on
                col = COLOR_ACCENT2 if self._pulse_on else blend(COLOR_ACCENT2, COLOR_BG_MAIN, 0.65)
                self.jobs_dot.configure(text_color=col)
                self.sys_dot.configure(text_color=col)
        except tk.TclError:
            pass
        finally:
            self.after(520, self.pulse_jobs)

    # ---------------------------------------------------------------- views
    def select_view(self, name: str, animate: bool = True):
        if not self.ready or name not in self.views:
            return
        if name != self.current_view and animate:
            caption = "> navigate --to " + self.PAGE_INFO[name][1].lower().replace(" ", "_")
            if self.run_glitch(self.stage, lambda: self._switch_view(name), caption):
                return
        self._switch_view(name)

    def _switch_view(self, name: str):
        self.current_view = name
        self.views[name].tkraise()
        for k, item in self.nav_items.items():
            item.set_active(k == name)
        ic, title, sub = self.PAGE_INFO[name]
        self.tb_icon.configure(image=icon_tile(ic, 44))
        self.tb_title.configure(text=title)
        self.tb_sub.configure(text=sub)
        hook = {"home": self.refresh_home, "cmd": self.refresh_scripts, "history": self.refresh_history,
                "auto": self.refresh_auto, "health": self.load_health_checks, "settings": self.refresh_settings}.get(name)
        if hook:
            hook()

    def refresh_application(self):
        """Rebuild the UI (used for theme changes / manual refresh) without losing the editor or console."""
        state = {"view": self.current_view, "mode": self.editor_mode, "name": self.entry_name.get(),
                 "code": self.editor.get(), "orig": self.editor_original, "orig_name": self.editor_orig_name,
                 "file": self.editor_file, "console": self.console.text()}
        self.sidebar.destroy()
        self.main.destroy()
        self.cards, self._sig = {}, None
        self.param_cache.clear()
        apply_theme(CFG["theme"])
        setup_ttk_style()
        self.configure(fg_color=COLOR_BG_MAIN)
        self.setup_sidebar()
        self.setup_content()
        self.set_editor_mode(state["mode"])
        self.entry_name.insert(0, state["name"])
        self.editor.set(state["code"])
        self.editor_original, self.editor_orig_name, self.editor_file = state["orig"], state["orig_name"], state["file"]
        self.update_dirty()
        if state["console"].strip():
            self.console.system(state["console"], "out")
        self.refresh_scripts(force=True)
        self._switch_view(state["view"])
        self.update_jobs()
        self.set_window_icon()
        self.style_titlebar()
        self.toast("Engine refreshed", "info", 1400)

    def open_workspace(self):
        open_path(workspace_dir())

    # ------------------------------------------------------------ shortcuts
    def on_ctrl_s(self, _event=None):
        if self.ready and not self.locked and self.current_view == "builder":
            self.save_script()
            return "break"

    def open_palette(self, _event=None):
        if not self.ready or self.locked:
            return "break"
        acts = [(f"Go to · {label.strip()}", lambda k=key: self.select_view(k)) for key, label in self.NAV]
        acts += [("New script", self.new_script_from_cmd), ("Lock now", self.lock_now),
                 ("Refresh engine", self.refresh_application), ("Open workspace folder", self.open_workspace),
                 ("Stop all running scripts", self.runner.stop_all),
                 ("Cycle theme", lambda: self.change_theme(list(THEMES)[(list(THEMES).index(CFG["theme"]) + 1) % len(THEMES)])),
                 ("Check environment", lambda: self.select_view("health"))]
        for f in self.list_script_files():
            acts.append((f"Run · {f}", lambda f=f: self.launch(f)))
            acts.append((f"Edit · {f}", lambda f=f: self.edit_script(f)))
            acts.append((f"History · {f}", lambda f=f: self.open_history_for(f)))
        for p in self.db.pipelines():
            acts.append((f"Run pipeline · {p['name']}", lambda n=p["name"]: self.run_pipeline(n)))
        show_palette(self, acts)
        return "break"

    # ------------------------------------------------------------ event pump
    def pump_events(self):
        n, t0, batch = 0, time.monotonic(), []
        try:
            while n < 600 and time.monotonic() - t0 < 0.03:
                ev = self.events.get_nowait()
                n += 1
                if ev[0] == "log":
                    if batch and batch[-1][0] == ev[1] and batch[-1][2] == ev[3]:
                        batch[-1][1] += ev[2]
                    else:
                        batch.append([ev[1], ev[2], ev[3]])
                else:
                    self.flush_logs(batch)
                    batch = []
                    try:
                        self.handle_event(ev)
                    except Exception as exc:  # noqa: BLE001
                        print(f"[event error] {type(exc).__name__}: {exc}")
        except queue.Empty:
            pass
        try:
            self.flush_logs(batch)
        finally:
            self.after(35 if n else 90, self.pump_events)

    def flush_logs(self, batch: list):
        if not self.ready:
            return
        for key, text, tag in batch:
            try:
                if key == SANDBOX_KEY:
                    console_write(self.test_console, text, tag)
                else:
                    self.console.write(key, text, tag)
                    if self._focus_next == key:
                        self.console.focus_key(key)
                        self._focus_next = None
            except tk.TclError:
                pass

    def handle_event(self, ev):
        kind = ev[0]
        if kind == "call":
            ev[1]()
        elif kind == "status" and self.ready:
            _, key, status = ev
            self.script_status[key] = status
            if key == SANDBOX_KEY:
                self.update_sandbox_buttons()
            elif not key.startswith("pipe:"):
                self.update_card(key)
            self.update_jobs()
        elif kind == "finished" and self.ready:
            self.on_run_finished(ev[2])
            self.update_jobs()

    def on_run_finished(self, res: RunResult):
        if res.status in ("failed", "timeout"):
            self.last_error_text = (res.stderr or res.stdout or "")[-4000:]
        if res.key == SANDBOX_KEY:
            return
        self.last_runs = self.db.last_runs_by_script()
        self.update_card(res.key)
        if self.current_view == "home":
            self.refresh_home()
        elif self.current_view == "history":
            self.refresh_history()
        self.check_achievements()
        name, dur, themed = Path(res.script).stem, fmt_duration(res.duration), CFG["themed_text"]
        if res.status == "success":
            self.toast(f"Litany complete: {name} ({dur})" if themed else f"{name} finished in {dur}", "success")
        elif res.status == "stopped":
            self.toast(f"Halted: {name}", "warning")
        else:
            self.toast(f"Rite failed: {name} - see the console / history" if themed else f"{name} failed", "error", 5000)
        if res.duration >= int(CFG["notify_after_s"] or 0) > 0:
            self.notify(APP_NAME, f"{name}: {res.status} ({dur})")
        self.beep(res.status == "success")

    def notify(self, title: str, message: str):
        try:
            plyer = try_import("plyer")
            if plyer:
                plyer.notification.notify(title=title, message=message, timeout=6)
            elif self.tray:
                self.tray.notify(message, title)
        except Exception:  # noqa: BLE001
            pass

    def beep(self, ok: bool):
        if not CFG["sound"]:
            return
        try:
            import winsound  # type: ignore
            winsound.MessageBeep(winsound.MB_OK if ok else winsound.MB_ICONHAND)
        except Exception:  # noqa: BLE001
            self.bell()

    # ------------------------------------------------------------------ runs
    def prepare_run(self, filename: str, values: dict | None, trigger_file: str | None = None):
        path = workspace_dir() / filename
        kind = "vba" if filename.lower().endswith(".bas") else "python"
        params = parse_params(path.read_text(encoding="utf-8", errors="replace")) if kind == "python" else []
        vals = {p.name: p.default for p in params}
        vals.update(self.db.get_last_params(filename))
        vals.update(values or {})
        if trigger_file:
            for p in params:
                if p.type == "file":
                    vals[p.name] = trigger_file
                    break
        env = build_param_env(params, vals)
        if trigger_file:
            env["FP_TRIGGER_FILE"] = trigger_file
        return path, kind, build_cli(params, vals), env

    def launch(self, filename: str, *, values=None, trigger_file=None, source="manual", interactive=True):
        path = workspace_dir() / filename
        if not path.exists():
            self.toast(f"{filename} no longer exists.", "error")
            self.refresh_scripts(force=True)
            return
        if self.runner.is_running(filename):
            self.toast(f"{Path(filename).stem} is already running.", "info", 1800)
            return
        params = self.script_params(filename)
        if interactive and params and values is None:
            vals = ask_params(self, filename, params, self.db.get_last_params(filename), trigger_file)
            if vals is None:
                return
            self.db.set_last_params(filename, vals)
            values = vals
        try:
            path, kind, cli, env = self.prepare_run(filename, values, trigger_file)
        except OSError as exc:
            self.toast(f"Could not read {filename}: {exc}", "error")
            return
        if source == "manual":
            self._focus_next = filename
        self.start_run(filename, filename, path, kind, cli, env, source)

    def start_run(self, key, script, path, kind, cli, env, source):
        timeout = int(CFG["run_timeout_s"] or 0)
        threading.Thread(target=self.runner.execute, daemon=True, kwargs=dict(
            key=key, script=script, path=path, kind=kind, args=cli, env_extra=env, timeout=timeout, source=source)).start()

    def stop_script(self, key: str):
        if not self.runner.stop(key):
            self.toast("This run cannot be interrupted (VBA macros must be stopped inside Excel).", "info", 4000)

    def stop_from_console(self, key):
        if key is None or key == SANDBOX_KEY:
            if not self.runner.running_keys():
                self.toast("Nothing is running.", "info", 1400)
            self.runner.stop_all()
        else:
            self.stop_script(key)

    # ------------------------------------------------- schedules & pipelines
    def fire_target(self, target: str, source: str, trigger_file: str | None = None):
        if target.startswith("pipe:"):
            self.run_pipeline(target[5:], source, trigger_file)
        elif target.startswith("script:"):
            self.launch(target[7:], interactive=False, source=source, trigger_file=trigger_file)

    def run_pipeline(self, name: str, source: str = "manual", trigger_file: str | None = None):
        p = self.db.get_pipeline(name)
        key = f"pipe:{name}"
        if not p:
            self.toast(f"Pipeline '{name}' no longer exists.", "error")
            return
        if key in self.pipes_running:
            self.toast(f"Pipeline '{name}' is already running.", "info", 1800)
            return
        self.pipes_running.add(key)
        if source == "manual":
            self._focus_next = key
        threading.Thread(target=self._pipeline_worker, args=(p, key, source, trigger_file), daemon=True).start()

    def _pipeline_worker(self, p: dict, key: str, source: str, trigger_file):
        emit = self.events.put
        steps, ok = p["steps"], True
        emit(("status", key, "running"))
        emit(("log", key, f"\n[{datetime.now():%H:%M:%S}] 🔗 PIPELINE {p['name']} - {len(steps)} step(s)\n", "info"))
        try:
            for i, step in enumerate(steps, 1):
                emit(("log", key, f"── step {i}/{len(steps)}: {step}\n", "info"))
                if not (workspace_dir() / step).exists():
                    emit(("log", key, f"   ✗ {step} does not exist any more\n", "error"))
                    ok = False
                else:
                    try:
                        path, kind, cli, env = self.prepare_run(step, None, trigger_file)
                        res = self.runner.execute(step, step, path, kind, cli, env, int(CFG["run_timeout_s"] or 0),
                                                  source=f"pipeline:{p['name']}")
                    except Exception as exc:  # noqa: BLE001
                        emit(("log", key, f"   ✗ could not start: {exc}\n", "error"))
                        res = None
                    if res is None:
                        emit(("log", key, f"   ✗ {step} could not run (already running?)\n", "error"))
                        ok = False
                    elif res.status != "success":
                        ok = False
                if not ok and p["stop_on_fail"]:
                    emit(("log", key, "⛔ pipeline stopped because a step failed\n", "error"))
                    break
            try:
                self.db.kv_set("pipeline_runs", str(int(self.db.kv_get("pipeline_runs", "0")) + 1))
            except Exception:  # noqa: BLE001
                pass
            emit(("log", key, ("✅ pipeline finished OK\n" if ok else "❌ pipeline finished with failures\n"),
                  "success" if ok else "error"))
            emit(("status", key, "success" if ok else "failed"))
            emit(("call", lambda: self.toast(f"Pipeline {p['name']}: " + ("all steps succeeded" if ok else "finished with failures"),
                                             "success" if ok else "error", 5000)))
            emit(("call", self.check_achievements))
        finally:
            self.pipes_running.discard(key)

    def tick_schedules(self):
        try:
            now = datetime.now()
            fired = False
            for s in self.db.schedules():
                if is_due(s, now):
                    self.db.mark_fired(s["id"])
                    fired = True
                    self.toast(f"Scheduled run: {self.target_label(s['target'])}", "info", 3000)
                    self.fire_target(s["target"], "schedule")
            if fired and self.current_view in ("auto", "home"):
                (self.refresh_auto if self.current_view == "auto" else self.refresh_home)()
        except Exception as exc:  # noqa: BLE001
            print(f"[scheduler] {exc}")
        finally:
            self.after(20000, self.tick_schedules)

    def tick_watchers(self):
        try:
            for w, path in self.watcher.poll(self.db.watchers()):
                self.db.bump_watcher(w["id"])
                self.toast(f"New file {Path(path).name} → {self.target_label(w['target'])}", "info", 4000)
                self.fire_target(w["target"], "watcher", path)
                if self.current_view == "auto":
                    self.refresh_auto()
        except Exception as exc:  # noqa: BLE001
            print(f"[watcher] {exc}")
        finally:
            self.after(5000, self.tick_watchers)

    # ------------------------------------------------------------ tray / quit
    def setup_tray(self) -> bool:
        if self.tray:
            return True
        pystray, pil = try_import("pystray"), try_import("PIL.Image")
        if not pystray or not pil:
            return False
        try:
            from PIL import Image, ImageDraw
            img = Image.new("RGBA", (64, 64), (0, 0, 0, 0))
            d = ImageDraw.Draw(img)
            d.ellipse((2, 2, 62, 62), fill=COLOR_BG_MAIN)
            d.polygon([(36, 6), (16, 36), (30, 36), (24, 58), (48, 26), (33, 26)], fill=COLOR_ACCENT)
            menu = pystray.Menu(
                pystray.MenuItem("Open Finprime's Engine", lambda: self.after(0, self.show_window), default=True),
                pystray.MenuItem("Lock", lambda: self.after(0, self.lock_now)),
                pystray.MenuItem("Quit", lambda: self.after(0, self.quit_app)))
            self.tray = pystray.Icon("finprime_engine", img, APP_NAME, menu)
            self.tray.run_detached()
            return True
        except Exception as exc:  # noqa: BLE001
            print(f"[tray] {exc}")
            self.tray = None
            return False

    def show_window(self):
        self.deiconify()
        self.lift()
        self.focus_force()

    def on_close(self):
        if self.ready and CFG["minimize_to_tray"] and self.setup_tray():
            self.withdraw()
            self.notify(APP_NAME, "Still running in the tray - schedules and watchers stay active.")
            return
        busy = self.runner.running_keys() + list(self.pipes_running)
        if busy and self.ready and not ask_confirm(self, "Quit while scripts are running?",
                                                   f"{len(busy)} run(s) still active. Quitting will stop them.", "Quit", danger=True):
            return
        self.quit_app()

    def quit_app(self):
        try:
            if self.tray:
                self.tray.stop()
        except Exception:  # noqa: BLE001
            pass
        try:
            self.runner.stop_all()
        except Exception:  # noqa: BLE001
            pass
        self.destroy()


def main():
    app = FinprimesEngineDashboard()
    app.mainloop()


if __name__ == "__main__":
    main()
